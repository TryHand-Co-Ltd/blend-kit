"""Populate the versioned editable report template without overwriting a run."""
from __future__ import annotations

import io
import xml.etree.ElementTree as ET
from copy import copy
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.packaging.core import DocumentProperties
from openpyxl.workbook.properties import CalcProperties

from report_model import (REPORT_LAYOUTS, VERSION, FAMILY, ReportData, SUMMARY_FIELDS,
                          INPUT_FIELDS, CASE_HEADER_ROW, CASE_START_ROW,
                          summary_formulas, detail_backlink_formula, display_lines as _lines)

INK = '172033'
LINK = '1D4ED8'
INPUT_FILL = 'FFF3CD'


def _text(cell, value):
    """All source text stays literal, including leading equals signs."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'
    cell.font = Font(name=cell.parent['A1'].font.name, size=11, color=INK)
    cell.alignment = Alignment(vertical='top', wrap_text=True)


def _chunks(value, width=95, max_lines=17):
    """Split only display rows; concatenating chunks preserves every character."""
    value = str(value)
    if _lines(value, width) <= max_lines:
        return [value]
    result, chunk = [], ''
    for char in value:
        if chunk and _lines(chunk + char, width) > max_lines:
            result.append(chunk)
            chunk = ''
        chunk += char
    if chunk:
        result.append(chunk)
    return result


def _height(sheet, row, values, widths, *, line_height=17, padding=10):
    lines = max(_lines(value, width) for value, width in zip(values, widths))
    sheet.row_dimensions[row].height = min(409, max(30, lines * line_height + padding))


def _jump(cell, sheet, row, caption=None):
    if caption is not None:
        _text(cell, caption)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet.replace(chr(39), chr(39)*2)}'!A{row}", display=str(cell.value))
    cell.font = Font(name=cell.parent['A1'].font.name, size=11, color=LINK, underline='single')


def _finish(sheet, end_column, end_row, header=None):
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = False
    sheet.print_area = f'A1:{end_column}{end_row}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    if header:
        sheet.freeze_panes = f'B{header + 1}'
        sheet.print_title_rows = f'{header}:{header}'


def build_workbook(data: ReportData):
    """Return the pristine projection for generation and read-only comparison."""
    locale = REPORT_LAYOUTS[data.language]
    asset = Path(__file__).resolve().parents[1] / 'assets' / locale['template']
    book = load_workbook(asset)
    if {p.name: p.value for p in book.custom_doc_props} != {
            'TemplateFamily': FAMILY, 'TemplateVersion': VERSION, 'Language': data.language}:
        raise ValueError('Report template family/version/language mismatch')
    if book.sheetnames != list(locale['sheets'][:2]):
        raise ValueError('Report template sheet schema mismatch')
    summary, results = (book[name] for name in locale['sheets'][:2])
    if tuple(results.cell(locale['header_row'], col).value for col in range(1, 8)) != locale['headers']:
        raise ValueError('Report template results headers mismatch')
    if any(summary.cell(SUMMARY_FIELDS[key], 2).value not in (None, '') for key in INPUT_FIELDS):
        raise ValueError('Report template input fields must be empty')
    if any(cell.value not in (None, '') for cells in results.iter_rows(min_row=locale['start_row']) for cell in cells):
        raise ValueError('Report template result rows must be empty')
    book.properties = DocumentProperties(creator='BLEND', lastModifiedBy='BLEND', title=locale['title'])
    book.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True)
    for key, row in SUMMARY_FIELDS.items():
        value = '' if key in INPUT_FIELDS else data.summary.get(key, '')
        # Fixed summary anchors must never silently clip source text.
        if _lines(value, 100) > 23:
            raise ValueError(f'Summary text exceeds readable row capacity: {key}')
        _text(summary.cell(row, 2), value)
        _height(summary, row, (summary.cell(row, 1).value or '', value), (38, 100))
        if key in INPUT_FIELDS:
            summary.cell(row, 2).fill = PatternFill('solid', fgColor=INPUT_FILL)
            summary.cell(row, 2).protection = Protection(locked=False)
    for row, case in enumerate(dict.fromkeys(record.case_id for record in data.rows), CASE_START_ROW):
        _text(summary.cell(row, 1), case)
        _height(summary, row, (case,), (38,))
    for address, formula in summary_formulas(data).items():
        summary[address] = formula
        summary[address].font = Font(name=locale['font'], size=11, color=INK)
        summary[address].fill = PatternFill('solid', fgColor='F1F5F9')
        summary[address].protection = Protection(locked=True)
        summary[address].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    _finish(summary, 'B', max(CASE_HEADER_ROW, CASE_START_ROW + len(set(r.case_id for r in data.rows)) - 1))

    details = None
    detail_row = locale['detail_start_row']
    for row, record in enumerate(data.rows, locale['start_row']):
        values = [record.identity, record.screen_function, record.conditions, record.expected,
                  '', locale['statuses']['NOT RUN'], '']
        for col, value in enumerate(values, 1):
            _text(results.cell(row, col), value)
        for col in (2, 3, 4):
            width = 22 if col == 2 else 32
            if _lines(values[col - 1], width) <= 12:
                continue
            if details is None:
                details = book.create_sheet(locale['sheets'][2])
                details['A1'].font = Font(name=locale['font'], size=11)
                _text(details['A1'], locale['sheets'][2])
                details['A1'].font = Font(name=locale['font'], size=15, bold=True, color=INK)
                details.row_dimensions[1].height = 32
                details.row_dimensions[2].height = 12
                for column, label in enumerate(locale['detail_headers'], 1):
                    target = details.cell(locale['detail_header_row'], column)
                    _text(target, label)
                    target._style = copy(results.cell(locale['header_row'], column)._style)
                for column, column_width in [('A', 24), ('B', 30), ('C', 104)]:
                    details.column_dimensions[column].width = column_width
                details.row_dimensions[locale['detail_header_row']].height = 32
            first = detail_row
            for chunk in _chunks(values[col - 1], max_lines=15 if data.language == 'ja' else 17):
                for column, text in enumerate((record.identity, locale['headers'][col - 1], chunk), 1):
                    _text(details.cell(detail_row, column), text)
                details.cell(detail_row, 1).value = detail_backlink_formula(record.identity, data)
                details.cell(detail_row, 1).font = Font(name=locale['font'], size=11, color=LINK, underline='single')
                # Yu Gothic needs print headroom beyond Excel's native AutoFit.
                _height(details, detail_row, (record.identity, locale['headers'][col - 1], chunk), (20, 28, 95),
                        line_height=22 if data.language == 'ja' else 17,
                        padding=18 if data.language == 'ja' else 10)
                detail_row += 1
            values[col - 1] = locale['detail_link']
            _jump(results.cell(row, col), details.title, first, locale['detail_link'])
        for col in (5, 6, 7):
            results.cell(row, col).fill = PatternFill('solid', fgColor=INPUT_FILL)
            results.cell(row, col).protection = Protection(locked=False)
        _height(results, row, values, (16, 22, 32, 32, 32, 10, 21))
    last = max(locale['start_row'], locale['start_row'] + len(data.rows) - 1)
    results.auto_filter.ref = f'A{locale["header_row"]}:G{last}'
    _finish(results, 'G', last, locale['header_row'])
    results.data_validations.dataValidation = []
    validation = DataValidation(type='list', formula1='"' + ','.join(locale['statuses'].values()) + '"', allow_blank=False)
    validation.showErrorMessage = True
    validation.errorTitle = 'Trạng thái không hợp lệ' if data.language == 'vi' else '無効な状態'
    validation.error = 'Chọn trạng thái trong danh sách.' if data.language == 'vi' else '一覧から状態を選択してください。'
    results.add_data_validation(validation)
    validation.add(f'F{locale["start_row"]}:F{last}')
    results.conditional_formatting._cf_rules.clear()
    for token, fill, color in [('FAIL', 'FEE2E2', '991B1B'), ('BLOCKED', 'FEF3C7', '92400E'),
                               ('SKIPPED', 'F1F5F9', '475569'), ('NOT RUN', 'F1F5F9', '475569')]:
        results.conditional_formatting.add(f'F{locale["start_row"]}:F{last}', CellIsRule(
            operator='equal', formula=[f'"{locale["statuses"][token]}"'],
            fill=PatternFill('solid', fgColor=fill), font=Font(name=locale['font'], size=11, color=color, bold=True)))
    if details is not None:
        _finish(details, 'C', detail_row - 1, locale['detail_header_row'])
    book.active = 0
    return book


def serialize_report(book) -> bytes:
    """Keep Excel's privacy setting on newly serialized report bytes."""
    buffer = io.BytesIO()
    book.save(buffer)
    output = io.BytesIO()
    with ZipFile(buffer) as source, ZipFile(output, 'w') as target:
        for entry in source.infolist():
            raw = source.read(entry.filename)
            if entry.filename == 'xl/workbook.xml':
                root = ET.fromstring(raw)
                namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
                properties = root.find(f'{{{namespace}}}workbookPr')
                if properties is None:
                    properties = ET.Element(f'{{{namespace}}}workbookPr')
                    root.insert(1 if root.find(f'{{{namespace}}}fileVersion') is not None else 0, properties)
                properties.set('filterPrivacy', '1')
                raw = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            target.writestr(entry, raw)
    return output.getvalue()


def render_report(data: ReportData, output: Path) -> dict:
    output = Path(output)
    if output.exists():
        raise FileExistsError('Existing output is preserved')
    book = build_workbook(data)
    raw = serialize_report(book)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('xb') as handle:
        handle.write(raw)
    return {'sheets': book.sheetnames, 'rows': len(data.rows),
            'detail_cases': len({book.worksheets[2].cell(row, 1).value for row in range(5, book.worksheets[2].max_row + 1)}) if len(book.worksheets) > 2 else 0,
            'images': 0}
