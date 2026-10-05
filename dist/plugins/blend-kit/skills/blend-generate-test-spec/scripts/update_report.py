"""Identity-keyed same-file writeback. Never executes tests or certifies pixels."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing

import block_report as blocks
import report_model as model
from render_report import _text, serialize_report


def report_bindings(source_dir, report, language='vi', *, feature_id=None):
    """Read/verify saved report first; callers never guess row positions."""
    data, _ = model.prepare_saved_report(Path(source_dir), report, language)
    blocks.check(report,source_dir,language)
    if feature_id and data.feature_id and feature_id != data.feature_id:
        raise ValueError('Feature ID differs from frozen design')
    feature = data.feature_id or feature_id
    if feature:
        model.working.identifier(feature)
    _, inputs = blocks.layout(data)
    return {'design_revision':data.summary['revision'],'feature_id':feature,
        'schema_version':data.report_version,'sheet':model.report_locale(data)['sheets'][1],
        'variants':{identity:{'case_id':identity.split(' / ')[0], 'variant_id':identity.split(' / ')[1],
            'status_cell':fields['status'],'actual_cell':fields['actual'],
            'checkpoints':fields.get('checkpoints',{}),'evidence_slots':fields.get('evidence_slots',{}),'picture_row':fields['picture'],
            'evidence_capacity':{'rows_per_checkpoint':blocks.READABLE_EVIDENCE_ROWS,'max_row_height_points':409,
                  'expansion':('owned reserved rows only; native collapse requires separate verification')}}
            for identity,fields in inputs.items()}}


def _validator():
    # Built profiles bundle the single shared helper alongside this writer.
    try:
        from run_artifacts import validate_evidence
    except ImportError:
        source_helper=Path(__file__).resolve().parents[2]/'blend-automation-test'/'scripts'/'run_artifacts.py'
        if not source_helper.is_file():
            raise ValueError('Shared run_artifacts helper is unavailable')
        import importlib.util
        spec=importlib.util.spec_from_file_location('blend_run_artifacts',source_helper)
        module=importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        validate_evidence=module.validate_evidence
    return validate_evidence


def _image_bytes(image):
    raw=image._data()
    image.ref=io.BytesIO(raw)  # openpyxl consumes its stream while reading.
    return raw


def _preserved(book, editable, mutable_rows=()):
    """Capture values/formulas/navigation/validation and all saved image bindings."""
    import xml.etree.ElementTree as ET
    cells={(sheet.title,c.coordinate):(c.value,c.data_type,c.hyperlink.location if c.hyperlink else None)
        for sheet in book for row in sheet for c in row if c.value is not None and (sheet.title,c.coordinate) not in editable}
    validations={s.title:ET.tostring(s.data_validations.to_tree()) for s in book}
    images=[]
    for sheet in book:
        for image in sheet._images:
            images.append((sheet.title,hashlib.sha256(_image_bytes(image)).hexdigest(),
                image.anchor._from.row,image.anchor._from.col,image.anchor._from.rowOff,
                image.anchor.pic.nvPicPr.cNvPr.name,image.anchor.pic.nvPicPr.cNvPr.descr))
    dimensions={(sheet.title,row):(d.height,d.hidden,d.outlineLevel,d.collapsed) for sheet in book
        for row,d in sheet.row_dimensions.items() if (sheet.title,row) not in mutable_rows}
    return cells,validations,images,dimensions




def _add_readable_images(sheet, checkpoint, records):
    """Caption row then proportionate image rows; unused slots stay collapsed."""
    rows = checkpoint['picture_rows']
    width = blocks.evidence_width(sheet)
    existing = [image for image in sheet._images if image.anchor._from.row+1 in rows]
    names = {image.anchor.pic.nvPicPr.cNvPr.name for image in existing}
    members = [(image, image.anchor.pic.nvPicPr.cNvPr.name, image.anchor.pic.nvPicPr.cNvPr.descr or checkpoint.get('title','')) for image in existing]
    for record, path in records:
        name = 'BLEND|'+record['case_id']+'|'+record['variant_id']+'|'+record['checkpoint_id']+'|'+record['annotated_sha256']
        if name not in names:
            caption = record.get('title') if model.descriptive_image_title(record.get('title')) else record['observed_note']
            if not model.descriptive_image_title(caption):
                caption = checkpoint.get('title','')+' — Evidence '+str(len(members)+1)
            model.public_text(caption, 'image title')
            image = Image(io.BytesIO(path.read_bytes()))
            members.append((image,name,caption))
            names.add(name)
    plan, cursor = [], 0
    for image, name, caption in members:
        raw = _image_bytes(image)
        pixels = Image(io.BytesIO(raw))
        ratio = min(1, width/pixels.width)
        display_width, height = pixels.width*ratio, pixels.height*ratio
        required_rows = max(1, int((height*.75+408)//409))
        if cursor+1+required_rows > len(rows):
            raise ValueError('Evidence capacity exceeded; original workbook preserved, request a new layout')
        title_row = rows[cursor]
        area = rows[cursor+1:cursor+1+required_rows]
        plan.append((image, name, caption, title_row, area, display_width, height))
        cursor += 1+required_rows
    for row in rows:
        _text(sheet.cell(row,1),'')
        sheet.row_dimensions[row].height=1
        sheet.row_dimensions[row].hidden=True
        sheet.row_dimensions[row].outlineLevel=0
    for image, name, caption, title_row, area, display_width, height in plan:
        _text(sheet.cell(title_row,1),caption)
        from openpyxl.styles import Alignment
        sheet.cell(title_row,1).alignment=Alignment(horizontal='left',vertical='top',indent=1,wrap_text=True)
        sheet.row_dimensions[title_row].height = max(20,model.display_lines(caption,109)*14+6)
        sheet.row_dimensions[title_row].outlineLevel=1
        if sheet.row_dimensions[title_row].height > 409:
            raise ValueError('Image title exceeds row capacity')
        for row in area:
            sheet.row_dimensions[row].height=max(1,height*.75/len(area))
            sheet.row_dimensions[row].outlineLevel=1
        picture = image.anchor.pic if not isinstance(image.anchor,str) else None
        if picture is None:
            picture = SpreadsheetDrawing()._picture_frame(len(sheet._images)+1)
        picture.nvPicPr.cNvPr.name=name
        picture.nvPicPr.cNvPr.descr=caption
        image.anchor=blocks.customer_image_anchor(sheet,area,0,display_width,height,picture)
        if image not in sheet._images:
            sheet.add_image(image)
    sheet.row_dimensions[checkpoint['label_row']].collapsed=True


def _refresh_evidence_headers(sheet, data, inputs):
    from render_report import _text
    for case in data.cases:
        members = [fields for identity,fields in inputs.items() if identity.split(' / ')[0] == case['id']]
        rows = set()
        for fields in members:
            rows.update(fields.get('picture_rows', []))
            for checkpoint in fields['checkpoints'].values():
                rows.update(checkpoint['picture_rows'])
        count = sum(image.anchor._from.row+1 in rows for image in sheet._images)
        if count:
            _text(sheet.cell(members[0]['evidence_row'],1), ('Ảnh minh chứng — ' if data.language=='vi' else '証拠画像 — ')+str(count)+(' ảnh' if data.language=='vi' else '枚'))


def update_report(source_dir, report, payload, *, language='vi', run_dir=None, feature_id=None):
    report=Path(report)
    data,captures=model.prepare_saved_report(Path(source_dir),report,language)
    binding=report_bindings(source_dir,report,language,feature_id=feature_id)
    identity=payload.get('identity',{})
    for key,value in (('design_revision',binding['design_revision']),('feature_id',binding['feature_id'])):
        if not value or identity.get(key)!=value:
            raise ValueError('Writeback identity mismatch: '+key)
    for key in ('case_id','variant_id'):
        model.working.identifier(identity.get(key,''))
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}-\d{3}',identity.get('run_id','')):
        raise ValueError('Run ID must use YYYY-MM-DD-NNN')
    target=identity['case_id']+' / '+identity['variant_id']
    if target not in binding['variants']:
        raise ValueError('Unknown report target identity')
    fields=binding['variants'][target]
    evidence_slots=blocks.evidence_bindings(fields)
    actual=model.required(payload.get('actual'),'actual')
    model.public_text(actual,'actual')
    if model.markdown_leak(actual):
        raise ValueError('Actual must be literal spreadsheet text, without Markdown')
    status=payload.get('status')
    if status not in model.REPORT_LAYOUTS[language]['statuses']:
        raise ValueError('Use canonical PASS/FAIL/BLOCKED/SKIPPED/NOT RUN status')
    evidence=payload.get('evidence',[])
    if not isinstance(evidence,list):
        raise ValueError('Evidence must be an array')
    if status=='PASS' and fields['checkpoints']:
        for cp in fields['checkpoints']:
            if not re.search(r'(?m)^'+re.escape(cp)+r':\s*\S',actual):
                raise ValueError('PASS requires observation for every checkpoint: '+cp)
        required={cp for cp,v in fields['checkpoints'].items() if v['artifact']=='screenshot'}
        if required-{e.get('checkpoint_id') for e in evidence}:
            raise ValueError('PASS requires reviewed evidence for every screenshot checkpoint')
    validated=[]
    if evidence:
        if run_dir is None:
            raise ValueError('Evidence writeback requires run directory')
        validate=_validator()
        for record in evidence:
            cp=record.get('checkpoint_id')
            if cp not in evidence_slots or evidence_slots[cp]['artifact']!='screenshot':
                raise ValueError('Unknown/non-screenshot evidence checkpoint')
            path=validate(record,Path(run_dir),{**identity,'checkpoint_id':cp,'report':str(report.resolve())})
            validated.append((record,Path(path)))
    captures[report]=model._capture(report)
    book=load_workbook(io.BytesIO(captures[report][0]))
    summary,tests=book.worksheets
    run_cell=summary[f'B{model.SUMMARY_FIELDS["run_id"]}']
    if run_cell.value and run_cell.value!=identity['run_id']:
        raise ValueError('Report belongs to a different run; preserve history and choose an explicitly requested new report')
    metadata=payload.get('run_metadata',{})
    if not isinstance(metadata,dict) or set(metadata)-set(model.INPUT_FIELDS):
        raise ValueError('Unsupported run metadata field')
    editable={(tests.title,fields['status_cell']),(tests.title,fields['actual_cell'])}
    editable|={(summary.title,f'B{model.SUMMARY_FIELDS[k]}') for k in metadata}
    editable.add((summary.title,run_cell.coordinate))
    mutable_rows={(tests.title,tests[fields['actual_cell']].row)}
    for cp in evidence_slots.values():
        if any(record['checkpoint_id']==cp['id'] for record,_ in validated):
            mutable_rows.update((tests.title,row) for row in cp['picture_rows'])
            if 'label_row' in cp:
                mutable_rows.add((tests.title,cp['label_row']))
            if cp.get('readable'):
                editable.update((tests.title, f'A{row}') for row in cp['picture_rows'])
                editable.add((tests.title, f'A{cp["label_row"]}'))
    before=_preserved(book,editable,mutable_rows)
    _text(run_cell,identity['run_id'])
    for key,value in metadata.items():
        if key=='run_id' and value!=identity['run_id']:
            raise ValueError('Run metadata identity mismatch')
        model.public_text(value,key)
        if key=='period':
            model.timestamp(value,key)
        _text(summary[f'B{model.SUMMARY_FIELDS[key]}'],value)
    _text(tests[fields['status_cell']],model.REPORT_LAYOUTS[language]['statuses'][status])
    _text(tests[fields['actual_cell']],actual)
    actual_width = 37
    from openpyxl.styles import Alignment
    for address in (fields['actual_cell'],fields['status_cell']):
        tests[address].alignment=Alignment(horizontal='left',vertical='top',indent=1,wrap_text=True)
    actual_height = model.display_lines(actual,actual_width)*14+6
    current_height=tests.row_dimensions[tests[fields['actual_cell']].row].height or 20
    tests.row_dimensions[tests[fields['actual_cell']].row].height=min(409,max(current_height,actual_height,20))
    if actual_height>409:
        raise ValueError('Actual exceeds reserved row capacity; shorten to observations without dropping assertions')
    for cp in evidence_slots:
        _add_readable_images(tests,evidence_slots[cp],[(r,p) for r,p in validated if r['checkpoint_id']==cp]) if any(r['checkpoint_id']==cp for r,p in validated) else None
    if data.report_version == model.VERSION and validated:
        _, all_inputs = blocks.layout(data)
        same_case = [value for name,value in all_inputs.items() if name.split(' / ')[0] == identity['case_id']]
        area = {row for value in same_case for cp in blocks.evidence_bindings(value).values() for row in cp['picture_rows']}
        count = sum(image.anchor._from.row+1 in area for image in tests._images)
        header = same_case[0]['evidence_row']
        _text(tests.cell(header,1), ('Ảnh minh chứng — ' if language=='vi' else '証拠画像 — ')+str(count)+(' ảnh' if language=='vi' else '枚'))
    raw=serialize_report(book)
    descriptor,name=tempfile.mkstemp(prefix='.blend-writeback-',suffix='.xlsx',dir=report.parent)
    candidate=Path(name)
    try:
        with os.fdopen(descriptor,'wb') as handle:
            handle.write(raw)
            handle.flush()
            os.fsync(handle.fileno())
        blocks.check(candidate,source_dir,language)
        reread=load_workbook(candidate)
        after=_preserved(reread,editable,mutable_rows)
        target_names={image[5] for image in before[2] if image[5].startswith('BLEND|'+identity['case_id']+'|'+identity['variant_id']+'|')}
        same_target=lambda image: (image[0],image[1],image[5],image[6])
        if before[:2]!=after[:2] or before[3]!=after[3] or any((same_target(image) not in [same_target(i) for i in after[2]] if image[5] in target_names else image not in after[2]) for image in before[2]):
            raise ValueError('Non-target report history changed during writeback')
        if reread.worksheets[1][fields['actual_cell']].value!=actual or reread.worksheets[1][fields['status_cell']].value!=model.REPORT_LAYOUTS[language]['statuses'][status]:
            raise ValueError('Writeback readback failed')
        model._unchanged(captures)
        os.replace(candidate,report)
        result=blocks.check(report,source_dir,language)
    finally:
        if candidate.exists():
            candidate.unlink()
    return {'identity':identity,'status':status,'evidence_added':len(after[2])-len(before[2]),
        'readback_verified':True,'check':result,'native_visual_review':'required separately'}


def main():
    for stream in (sys.stdin,sys.stdout,sys.stderr):
        if hasattr(stream,'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir',type=Path,required=True)
    parser.add_argument('--report',type=Path,required=True)
    parser.add_argument('--language',choices=('vi','ja'),default='vi')
    parser.add_argument('--feature-id',help='Verified ID for legacy sources which have no Feature ID field')
    parser.add_argument('--run-dir',type=Path)
    parser.add_argument('--bindings',action='store_true',help='Read-only binding inventory; no stdin/write')
    args=parser.parse_args()
    try:
        result=report_bindings(args.source_dir,args.report,args.language,feature_id=args.feature_id) if args.bindings else update_report(
            args.source_dir,args.report,json.load(sys.stdin),language=args.language,run_dir=args.run_dir,feature_id=args.feature_id)
    except (ValueError,TypeError,OSError,ImportError,KeyError) as error:
        parser.exit(1,f'Report writeback blocked: {error}\n')
    print(json.dumps(result,ensure_ascii=False))


if __name__=='__main__':
    main()
