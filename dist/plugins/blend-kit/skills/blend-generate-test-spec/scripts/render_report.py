"""Populate the versioned editable report template without overwriting a run."""
from __future__ import annotations

import io
import unicodedata
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.hyperlink import Hyperlink

from report_model import ReportData, display_lines as _lines

INK = '172033'
LINK = '1D4ED8'
INPUT_FILL = 'FFF3CD'


def _text(cell, value):
    """All source text stays literal, including leading equals signs."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'
    cell.font = Font(name=cell.parent['A1'].font.name, size=11, color=INK)
    cell.alignment = Alignment(vertical='top', wrap_text=True)


def _chunks(value, width=95, max_lines=17):
    """Split only display rows; concatenating chunks preserves every character."""
    value = str(value)
    if _lines(value, width) <= max_lines:
        return [value]
    result, chunk = [], []
    completed, units = 0, 0
    for char in value:
        char_units = 2 if unicodedata.east_asian_width(char) in ('W','F') else 1
        prospective = (completed + max(1,(units+width-1)//width) + 1 if char == '\n' else
            completed + max(1,(units+char_units+width-1)//width))
        if chunk and prospective > max_lines:
            result.append(''.join(chunk))
            chunk = []
            completed, units = 0, 0
        if char == '\n':
            completed += max(1,(units+width-1)//width)
            units = 0
        else:
            units += char_units
        chunk.append(char)
    if chunk:
        result.append(''.join(chunk))
    return result


def _height(sheet, row, values, widths, *, line_height=17, padding=10):
    lines = max(_lines(value, width) for value, width in zip(values, widths))
    sheet.row_dimensions[row].height = min(409, max(30, lines * line_height + padding))


def _jump(cell, sheet, row, caption=None):
    if caption is not None:
        _text(cell, caption)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet.replace(chr(39), chr(39)*2)}'!A{row}", display=str(cell.value))
    cell.font = Font(name=cell.parent['A1'].font.name, size=11,
                     color=INK if '\n' in str(cell.value) else LINK,
                     underline=None if '\n' in str(cell.value) else 'single')


def _finish(sheet, end_column, end_row, header=None):
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = False
    sheet.print_area = f'A1:{end_column}{end_row}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    if header:
        sheet.freeze_panes = f'B{header + 1}'
        sheet.print_title_rows = f'{header}:{header}'


def build_workbook(data: ReportData):
    from block_report import build
    return build(data)


def serialize_report(book) -> bytes:
    """Keep Excel's privacy setting on newly serialized report bytes."""
    buffer = io.BytesIO()
    book.save(buffer)
    output = io.BytesIO()
    with ZipFile(buffer) as source, ZipFile(output, 'w') as target:
        for entry in source.infolist():
            raw = source.read(entry.filename)
            if entry.filename == 'xl/workbook.xml':
                root = ET.fromstring(raw)
                namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
                properties = root.find(f'{{{namespace}}}workbookPr')
                if properties is None:
                    properties = ET.Element(f'{{{namespace}}}workbookPr')
                    root.insert(1 if root.find(f'{{{namespace}}}fileVersion') is not None else 0, properties)
                properties.set('filterPrivacy', '1')
                raw = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            target.writestr(entry, raw)
    return output.getvalue()


def render_report(data: ReportData, output: Path) -> dict:
    output = Path(output)
    if output.exists():
        raise FileExistsError('Existing output is preserved')
    book = build_workbook(data)
    raw = serialize_report(book)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('xb') as handle:
        handle.write(raw)
    return {'sheets': book.sheetnames, 'rows': len(data.rows),
            'detail_cases': len({book.worksheets[2].cell(row, 1).value for row in range(5, book.worksheets[2].max_row + 1)}) if len(book.worksheets) > 2 else 0,
            'images': 0}
