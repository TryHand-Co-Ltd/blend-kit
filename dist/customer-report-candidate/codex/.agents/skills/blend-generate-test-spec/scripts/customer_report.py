"""Validate a closed internal run and export a separate customer workbook.

No tests, network requests, uploads, translation, or changes to the input master.
Delivery confirmations are transient caller evidence, never customer metadata.
"""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
import hashlib
import io
import ipaddress
import json
import math
from pathlib import Path
import re
import sys
import tempfile
import unicodedata
from urllib.parse import unquote, urlsplit

import export_report as working

VERSION = "1.0.0"
CUSTOMER_LAYOUTS = {
    "vi": {
        "template": "customer-test-report-template.vi.xlsx",
        "sheets": ("Tổng quan", "Kết quả", "Chi tiết"),
        "title": "Báo cáo kết quả kiểm thử",
        "headers": ("TC / Variant", "Màn hình / chức năng", "Điều kiện / thao tác", "Kết quả mong đợi", "Kết quả thực tế", "Trạng thái", "Bằng chứng / lỗi"),
        "detail_headers": ("TC / Variant", "Nội dung", "Chi tiết"),
        "detail_link": "Xem chi tiết", "evidence_link": "Bằng chứng", "bug_link": "Lỗi",
        "no_data": "Chưa có dữ liệu", "partial": "Đã đóng lượt chạy; còn lỗi hoặc phạm vi chưa được xác nhận đạt.",
        "evaluated": "Các biến thể trong thiết kế này đã được xác nhận đạt; không khẳng định toàn bộ chức năng đã được bao phủ.",
        "observation": "Chỉ là quan sát; kỳ vọng hoặc điều kiện kiểm chưa được xác nhận.",
        "out_of_scope": "Ngoài phạm vi lượt chạy", "reason": "Lý do / tồn đọng", "executed": "Thực hiện",
        "labels": {"run_id": "Lượt chạy", "feature": "Chức năng", "build": "Build", "environment": "Môi trường", "period": "Thời gian", "scope": "Phạm vi", "excluded": "Ngoài phạm vi", "limitations": "Giới hạn", "conclusion": "Kết luận", "open_items": "Tồn đọng / bước tiếp theo", "cases": "Số testcase", "variants": "Số biến thể", "selected_variants": "Biến thể trong lượt chạy", "evaluated_variants": "Biến thể có kết luận hợp lệ", "pass_rate": "Tỷ lệ PASS / biến thể có kết luận", "completion_rate": "Tỷ lệ có kết luận / biến thể xác nhận trong scope", "case_statuses": "Kết quả testcase", "variant_statuses": "Trạng thái biến thể"},
        "parts": {"configuration": "Cấu hình", "trigger": "Kích hoạt", "observation": "Vị trí quan sát", "actor": "Vai trò / quyền", "inputs": "Dữ liệu", "action": "Thao tác", "expected": "Kết quả mong đợi", "preservation": "Bảo toàn", "steps": "Bước", "actual": "Thực tế", "evidence": "Bằng chứng"},
    },
    "ja": {
        "template": "customer-test-report-template.ja.xlsx",
        "sheets": ("概要", "結果一覧", "テスト詳細"),
        "title": "テスト結果報告書",
        "headers": ("TC / Variant", "画面・機能", "条件・操作", "期待結果", "実際の結果", "状態", "証拠・不具合"),
        "detail_headers": ("TC / Variant", "項目", "詳細"),
        "detail_link": "詳細を確認", "evidence_link": "証拠", "bug_link": "不具合",
        "no_data": "データなし", "partial": "実行を終了しました。未解決の不具合、または合格を確認できていない範囲があります。",
        "evaluated": "本設計の全バリエーションで合格を確認しました。機能全体の網羅を保証するものではありません。",
        "observation": "観測記録のみ。期待結果または実行条件が未確定です。",
        "out_of_scope": "今回の実行対象外", "reason": "理由・残存事項", "executed": "実施",
        "labels": {"run_id": "実行ID", "feature": "機能", "build": "ビルド", "environment": "環境", "period": "実施期間", "scope": "対象範囲", "excluded": "対象外", "limitations": "制限事項", "conclusion": "結論", "open_items": "残存事項・次の対応", "cases": "ケース数", "variants": "バリエーション数", "selected_variants": "実行対象のバリエーション数", "evaluated_variants": "有効な判定のあるバリエーション数", "pass_rate": "合格率 / 有効な判定", "completion_rate": "判定完了率 / 対象内の確定バリエーション", "case_statuses": "ケース別状態", "variant_statuses": "バリエーション別状態"},
        "parts": {"configuration": "設定", "trigger": "起動", "observation": "確認箇所", "actor": "役割・権限", "inputs": "入力データ", "action": "操作", "expected": "期待結果", "preservation": "維持状態", "steps": "手順", "actual": "実際の結果", "evidence": "証拠"},
    },
}
for _layout in CUSTOMER_LAYOUTS.values():
    _layout.update(header_row=4, start_row=5, detail_header_row=4, detail_start_row=5, summary_start_row=4)
CUSTOMER_LAYOUTS["vi"]["labels"].update(passed_variants="Biến thể được xác nhận PASS", confirmed_variants="Biến thể xác nhận trong scope")
CUSTOMER_LAYOUTS["ja"]["labels"].update(passed_variants="有効な合格バリエーション数", confirmed_variants="対象内の確定バリエーション数")
CUSTOMER_LAYOUTS["ja"]["font"] = "Yu Gothic"
CUSTOMER_LAYOUTS["vi"]["font"] = "Arial"


@dataclass(frozen=True)
class ClosureConfirmation:
    closed_by: str
    closed_at: str
    source: str
    # Exact intended recipients, not the agent's authenticated account.
    audience: str


@dataclass(frozen=True)
class EvidenceAccess:
    url: str
    audience: str
    verified_by: str
    verified_at: str
    method: str  # human-attestation or recipient-session
    source: str


@dataclass(frozen=True)
class ScreenshotInput:
    case_id: str
    variant: str
    path: Path
    caption: str
    reviewed_by: str
    reviewed_at: str
    source: str


@dataclass(frozen=True)
class CustomerImage:
    content: bytes  # re-encoded PNG, no original metadata or source locator
    caption: str


@dataclass(frozen=True)
class CustomerResult:
    case_id: str
    variant: str
    screen_function: str
    conditions: str
    expected: str
    actual: str
    status: str  # original recorded token; never silently rewritten
    evidence: tuple[str, ...]
    bug: str
    needs_detail: bool

    @property
    def identity(self) -> str:
        return f"{self.case_id} / {self.variant}"


@dataclass(frozen=True)
class CustomerDetail:
    case_id: str
    variant: str
    blocks: tuple[tuple[str, str], ...]
    images: tuple[CustomerImage, ...] = ()

    @property
    def identity(self) -> str:
        return f"{self.case_id} / {self.variant}"


@dataclass(frozen=True)
class CustomerReportData:
    language: str
    summary: tuple[tuple[str, str], ...]
    metrics: dict
    rows: tuple[CustomerResult, ...]
    details: tuple[CustomerDetail, ...]


def required(value, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"Missing text: {field}")
    return working.cell_value(value.strip())


def timestamp(value, field: str) -> datetime:
    parsed = datetime.fromisoformat(required(value, field).replace("Z", "+00:00"))
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError(f"Timestamp requires timezone: {field}")
    return parsed


def shared_url(value: str, field: str) -> str:
    url = urlsplit(value)
    host = (url.hostname or "").lower()
    if url.scheme != "https" or not host or url.username or url.password or "." not in host or host in ("localhost",) or host.endswith((".local", ".localhost", ".internal")):
        raise ValueError(f"Customer evidence requires a shared HTTPS URL: {field}")
    try:
        address = ipaddress.ip_address(host)
    except ValueError:
        pass
    else:
        if not address.is_global:
            raise ValueError(f"Private evidence address: {field}")
    if re.search(r"(?:token|secret|password|signature|credential|api[_-]?key)=", url.query, re.I):
        raise ValueError(f"Credential-bearing URL: {field}")
    return value


PRIVATE = re.compile(r"(?<![\w])(?:[A-Za-z]:[\\/]|\\\\|file:|localhost\b|127\.0\.0\.1|~/|/(?:Users|home|tmp|private|mnt|var|opt|etc|root|workspace|Volumes)/)|(?:\.agents|\.codex|blend-context|docker-codeigniter)[\\/]|(?:\b(?:research|sources|application|tests|docs)/[^\s]+)|(?:\b[^\s/]+\.(?:md|sql|php|py)(?:\b|#))", re.I)


def public_text(value: str, field: str, *, dummy_input: bool = False) -> str:
    value = required(value, field)
    # Explicit source-authorized dummy inputs retain their meaning. They are not
    # evidence locators and cannot exempt actual results, notes, URLs or metadata.
    scan = value
    if dummy_input:
        scan = re.sub(r"\[dummy-input: [^\]\n]+\]", "(intentional test input)", scan)
    if PRIVATE.search(unquote(scan)) or re.search(r"@(?:Steps|CTX-)|sha256:|<!--|^\s*\|.*\|\s*$", scan, re.M):
        raise ValueError(f"Private/internal reference in published field: {field}")
    if re.search(r"\b(?:password|secret|access[_-]?token|api[_-]?key)\s*[:=]\s*\S+", scan, re.I):
        raise ValueError(f"Potential credential in published field: {field}")
    for url in re.findall(r"(?:https?|file)://[^\s<>\])]+", scan):
        shared_url(url, field)
    return value


def display_lines(value: str, width: int) -> int:
    """Shared display contract: JA full-width glyphs occupy two Latin units."""
    return sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1
                                    for char in line) / width)) for line in value.split("\n"))


def needs_detail(conditions: str, expected: str, actual: str, images: tuple = (), *,
                 screen_function: str = "", link_count: int = 0, bug: str = "") -> bool:
    """Material text above twelve readable lines uses conditional detail blocks."""
    return (bool(images) or link_count > 1 or display_lines(screen_function, 22) > 12 or
            display_lines(bug, 21) > 10 or
            any(display_lines(value, 32) > 12 for value in (conditions, expected, actual)))


def _records(value, cls):
    if isinstance(value, cls):
        return value
    if not isinstance(value, dict):
        raise ValueError(f"Expected {cls.__name__} object")
    try:
        record = cls(**value)
    except TypeError as error:
        raise ValueError(f"Invalid {cls.__name__} fields") from error
    return record


def _capture(path: Path) -> tuple[bytes, tuple]:
    stat = path.stat()
    raw = path.read_bytes()
    after = path.stat()
    identity = (stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns)
    if identity != (after.st_dev, after.st_ino, after.st_size, after.st_mtime_ns):
        raise ValueError("Input changed while captured")
    return raw, identity


def _unchanged(captures: dict[Path, tuple[bytes, tuple]]) -> None:
    for path, (raw, identity) in captures.items():
        current, current_identity = _capture(path)
        if current != raw or current_identity != identity:
            raise ValueError("Input identity/content changed during export")


def _read_text(cell, field: str, *, optional=False) -> str:
    if cell.data_type == "f":
        raise ValueError(f"Untrusted formula in recorded field: {field}")
    if optional and cell.value is None:
        return ""
    if isinstance(cell.value, datetime) and field == "executed at":
        return cell.value.isoformat()
    return required(cell.value, field)


def prepare_customer_report(run_file: Path, design_dir: Path, language: str,
                            closure_confirmation, evidence_access, screenshots):
    """Return customer-only data and private capture receipts for drift checks."""
    from openpyxl import load_workbook
    if language not in CUSTOMER_LAYOUTS:
        raise ValueError("Customer language must be ja or vi")
    locale = CUSTOMER_LAYOUTS[language]
    closure = _records(closure_confirmation, ClosureConfirmation)
    for key in ("closed_by", "source", "audience"):
        required(getattr(closure, key), f"closure.{key}")
    closed_at = timestamp(closure.closed_at, "closure.closed_at")
    accesses = {}
    for item in evidence_access:
        access = _records(item, EvidenceAccess)
        shared_url(access.url, "evidence_access.url")
        for key in ("audience", "verified_by", "source"):
            required(getattr(access, key), f"evidence_access.{key}")
        timestamp(access.verified_at, "evidence_access.verified_at")
        if access.audience != closure.audience or access.method not in ("human-attestation", "recipient-session") or access.url in accesses:
            raise ValueError("Evidence audience/method/uniqueness not verified")
        accesses[access.url] = access
    captures = {run_file: _capture(run_file)}
    book = load_workbook(io.BytesIO(captures[run_file][0]), data_only=False)
    props = {p.name: p.value for p in book.custom_doc_props}
    master_language = props.get("Language")
    if props.get("TemplateFamily") != "test-case-report" or props.get("TemplateVersion") != VERSION or master_language not in CUSTOMER_LAYOUTS or book.sheetnames != ["Run", "Cases", "Data"]:
        raise ValueError("Run must use the JA/VI working schema v1")
    run = book["Run"]
    if [run.cell(24, i).value for i in range(1, 20)] != working.WORKBOOK_LOCALES[master_language]["headers"]["Run"]:
        raise ValueError("Working run header/schema mismatch")
    fingerprint_text = _read_text(run["B11"], "source fingerprints")
    fingerprints = {}
    for line in fingerprint_text.splitlines():
        match = re.fullmatch(r"((?:scope-and-approach|test-cases|test-data)\.(?:ja|vi)\.md): sha256:([0-9a-f]{64})", line)
        if not match or match[1] in fingerprints:
            raise ValueError("Invalid/duplicate design fingerprints")
        fingerprints[match[1]] = match[2]
    required_names = {f"{family}.{lang}.md" for lang in {language, master_language} for family in working.HEADINGS[lang]}
    if not required_names <= fingerprints.keys():
        raise ValueError("Requested-language design is not bound to run")
    captured_design = {}
    for name in required_names:
        path = design_dir / name
        captures[path] = _capture(path)
        captured_design[name] = captures[path][0]
        if hashlib.sha256(captured_design[name]).hexdigest() != fingerprints[name]:
            raise ValueError("Design fingerprint mismatch; supply frozen run revision")
    design = working.parse_sources(design_dir, language, captured=captured_design)
    master_design = working.parse_sources(design_dir, master_language, captured=captured_design)
    def identities(item):
        return (item["revision"], [(case["id"], [variant[0] for variant in case["variants"]], case["priority"], case["basis"], case["readiness"], case["gap"],
                                  "local:" if case["fixture"].startswith("local: ") else case["fixture"]) for case in item["cases"]],
                [fixture["id"] for fixture in item["fixtures"]], [(row[0], row[1]) for row in item["gaps"]], [row[2] for row in item["coverage"]])
    if identities(design) != identities(master_design):
        raise ValueError("Requested-language identity/classification differs from executed design")
    original_cases = {case["id"]: case for case in master_design["cases"]}
    if _read_text(run["B3"], "design revision") != design["revision"]:
        raise ValueError("Run/design revision mismatch")
    if not re.fullmatch(r"sha256:[0-9a-f]{64}", _read_text(run["B12"], "template fingerprint")):
        raise ValueError("Missing working template identity")
    # Fingerprints alone do not protect an edited Cases/Data sheet. Compare the
    # frozen projection exactly; display strings are never parsed back into steps.
    for offset, case in enumerate(master_design["cases"], working.CASE_START):
        values = [case["id"], case["group"], case["title"], *[case[k] for k in working.CASE_KEYS],
                  "\n".join(" | ".join(row) for row in case["steps"]),
                  "\n".join(" | ".join(row) for row in case["variants"])]
        if [book["Cases"].cell(offset, c).value or "" for c in range(1, 22)] != values:
            raise ValueError("Working case design differs from frozen source")
    if any(cell.value is not None for row in book["Cases"].iter_rows(min_row=working.CASE_START + len(design["cases"])) for cell in row):
        raise ValueError("Working case inventory differs from frozen source")
    for offset, fixture in enumerate(master_design["fixtures"], working.DATA_START):
        if [book["Data"].cell(offset, c).value for c in range(1, 10)] != [fixture["id"], *fixture["values"]]:
            raise ValueError("Working fixture differs from frozen source")
    meta = {key: public_text(_read_text(run[cell], key), key) for key, cell in
            (("run_id", "B5"), ("build", "B6"), ("environment", "B7"), ("period", "B9"))}
    _read_text(run["B8"], "run tester")
    expected_pairs = [(case, variant) for case in design["cases"] for variant in case["variants"]]
    if not expected_pairs:
        raise ValueError("Empty design cannot be a final report")
    records = {}
    for row in run.iter_rows(min_row=working.RUN_START, max_col=19):
        if not any(cell.value is not None for cell in row):
            continue
        key = (_read_text(row[1], "Case ID"), _read_text(row[2], "Variant"))
        if key in records:
            raise ValueError("Duplicate Case/Variant in run")
        if row[0].value not in ("=$B$5", meta["run_id"]):
            raise ValueError("Row Run ID differs from run metadata")
        records[key] = row
    if set(records) != {(case["id"], variant[0]) for case, variant in expected_pairs}:
        raise ValueError("Run/design Case/Variant inventory mismatch")
    images = {}
    for item in screenshots:
        shot = _records(item, ScreenshotInput)
        key = (shot.case_id, shot.variant)
        if key not in records:
            raise ValueError("Screenshot has unknown Case/Variant")
        required(shot.reviewed_by, "screenshot.reviewed_by")
        required(shot.source, "screenshot.source")
        timestamp(shot.reviewed_at, "screenshot.reviewed_at")
        caption = public_text(shot.caption, "screenshot.caption")
        path = Path(shot.path)
        captures[path] = _capture(path)
        from PIL import Image, ImageOps
        with Image.open(io.BytesIO(captures[path][0])) as image:
            if image.format not in ("PNG", "JPEG") or getattr(image, "n_frames", 1) != 1:
                raise ValueError("Screenshot must be a reviewed static PNG/JPEG")
            clean = ImageOps.exif_transpose(image).convert("RGB")
            clean.info.clear()
            buffer = io.BytesIO()
            clean.save(buffer, format="PNG")
        images.setdefault(key, []).append(CustomerImage(buffer.getvalue(), caption))
    fixtures = {item["id"]: item["values"][:5] for item in design["fixtures"]}
    rows, details, states, open_items = [], [], [], []
    for case, variant in expected_pairs:
        key = (case["id"], variant[0])
        record = records[key]
        for column, value in ((3, original_cases[case["id"]]["group"]), (4, case["priority"]), (5, case["basis"]), (6, case["readiness"])):
            if _read_text(record[column], "run classification") != value:
                raise ValueError("Run classification differs from frozen design")
        selected = _read_text(record[7], "run scope")
        status = _read_text(record[8], "status")
        if selected not in ("Yes", "No") or status not in ("PASS", "FAIL", "BLOCKED", "SKIPPED", "NOT RUN"):
            raise ValueError("Invalid recorded scope/status")
        actual, evidence, bug, tester, at, zone, reason = [_read_text(record[i], label, optional=True) for i, label in
            ((9, "actual"), (10, "evidence"), (11, "bug"), (12, "tester"), (13, "executed at"), (14, "timezone"), (15, "reason"))]
        if isinstance(record[13].value, datetime):
            at = record[13].value.isoformat()
        if status in ("PASS", "FAIL"):
            if not all((actual, evidence, tester, at, zone)):
                raise ValueError("PASS/FAIL requires actual, evidence, tester, time and timezone")
            # Native Excel stores naive date cells; explicit timezone column is mandatory.
            try:
                executed = datetime.fromisoformat(at.replace("Z", "+00:00"))
            except ValueError as error:
                raise ValueError("Execution timestamp must be ISO datetime") from error
            if executed.tzinfo is None:
                from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
                try:
                    executed = executed.replace(tzinfo=ZoneInfo(zone))
                except ZoneInfoNotFoundError as error:
                    raise ValueError("Execution timezone must be an available IANA zone") from error
            if executed > closed_at:
                raise ValueError("Execution is after run closure")
        if status in ("FAIL", "BLOCKED", "SKIPPED", "NOT RUN") or selected == "No":
            if not reason:
                raise ValueError("Outstanding/unexecuted/out-of-scope row requires reason and next action")
        links = tuple(evidence.splitlines()) if evidence else ()
        for link in links:
            shared_url(link, "recorded evidence")
            if link not in accesses:
                raise ValueError("Customer access is unverified for recorded evidence")
        if bug.startswith("https://"):
            shared_url(bug, "bug")
            if bug not in accesses:
                raise ValueError("Customer access is unverified for bug link")
        elif bug:
            public_text(bug, "bug")
        parts = locale["parts"]
        blocks = [(parts[k], public_text(case[k], k)) for k in ("configuration", "trigger", "observation", "actor")]
        if case["fixture"].startswith("local: "):
            inputs = case["fixture"][7:]
        else:
            inputs = "\n".join("; ".join(fixtures[name]) for name in case["fixture"].split(", "))
        if variant[0] != "base" or variant[1] != case["fixture"]:
            inputs += "\n" + variant[1]
        blocks.append((parts["inputs"], public_text(inputs, "inputs", dummy_input=True)))
        if case["steps"]:
            for number, action, expected, preservation in case["steps"]:
                blocks.append((f'{parts["steps"]} {number}', public_text(action, "step action")))
                blocks.append((f'{parts["steps"]} {number} — {parts["expected"]}', public_text(expected, "step expected")))
                blocks.append((f'{parts["steps"]} {number} — {parts["preservation"]}', public_text(preservation, "step preservation")))
            actions = "\n".join(f'{n}. {action}' for n, action, _, _ in case["steps"])
            expected = "\n".join(f'{n}. {value}\n{parts["preservation"]}: {preserve}' for n, _, value, preserve in case["steps"])
            if variant[2] != "@Steps":
                expected += "\n" + variant[2]
        else:
            actions = public_text(case["action"], "action")
            expected = public_text(case["expected"], "expected")
            if variant[2] != case["expected"]:
                expected += "\n" + variant[2]
            blocks.append((parts["action"], actions))
        expected += f'\n{parts["preservation"]}: ' + public_text(case["preservation"], "preservation")
        expected = public_text(expected, "resolved expected")
        blocks.append((parts["expected"], expected))
        conditions = "\n".join(f"{label}: {value}" for label, value in blocks[:5]) + f'\n{parts["action"]}: {actions}'
        valid = selected == "Yes" and case["basis"] == "Confirmed" and case["readiness"] == "Ready"
        published_actual = public_text(actual, "actual") if actual else locale["no_data"]
        if reason:
            reason = public_text(reason, "reason")
            published_actual += f'\n{locale["reason"]}: {reason}'
        if at:
            published_actual += f'\n{locale["executed"]}: {public_text(at + " " + zone, "execution time")}'
        if selected == "No":
            published_actual += "\n" + locale["out_of_scope"]
        if case["basis"] != "Confirmed" or case["readiness"] != "Ready":
            published_actual += "\n" + locale["observation"]
        shot_images = tuple(images.get(key, ()))
        title = public_text(case["group"] + " — " + case["function"] + "\n" + case["title"], "screen/function")
        detail = needs_detail(conditions, expected, published_actual, shot_images, screen_function=title,
                              link_count=len(links) + int(bug.startswith("https://")), bug=bug if not bug.startswith("https://") else "")
        result = CustomerResult(key[0], key[1], title, conditions, expected, published_actual, status, links, bug, detail)
        rows.append(result)
        if detail:
            details.append(CustomerDetail(*key, tuple(blocks) + ((parts["actual"], published_actual),), shot_images))
        states.append((case["id"], status, selected == "Yes", valid))
        if not valid or status != "PASS":
            open_items.append(result.identity + ": " + (reason or locale["observation"]))
    if all(status == "NOT RUN" for _, status, _, _ in states):
        raise ValueError("Entirely NOT RUN master is not a completed final")
    counts = Counter(status for _, status, selected, _ in states if selected)
    evaluated = sum(selected and valid and status in ("PASS", "FAIL") for _, status, selected, valid in states)
    confirmed = sum(selected and valid for _, _, selected, valid in states)
    passed = sum(selected and valid and status == "PASS" for _, status, selected, valid in states)
    case_states = {}
    for case in design["cases"]:
        group = [state for state in states if state[0] == case["id"]]
        case_states[case["id"]] = ("FAIL" if any(s == "FAIL" and ok for _, s, _, ok in group) else
                                  "PASS" if all(s == "PASS" and ok for _, s, _, ok in group) else
                                  "BLOCKED" if any(s == "BLOCKED" for _, s, _, _ in group) else "NOT RUN")
    metrics = {"cases": len(case_states), "variants": len(states), "selected_variants": sum(s for _, _, s, _ in states),
               "evaluated_variants": evaluated, "confirmed_variants": confirmed, "passed_variants": passed,
               "pass_rate": passed / evaluated if evaluated else None, "completion_rate": evaluated / confirmed if confirmed else None,
               "variant_statuses": dict(counts), "case_statuses": dict(Counter(case_states.values())), "case_results": case_states}
    scope_labels = working.SCOPE_LABELS[language]
    meta.update(feature=public_text(design["conventions"]["Feature"], "feature"),
                scope=public_text(design["scope"][1][scope_labels[1][0]], "scope"),
                excluded=public_text(design["scope"][1][scope_labels[1][1]], "excluded"),
                limitations=public_text(design["scope"][0][scope_labels[0][5]], "limitations"),
                conclusion=locale["evaluated"] if all(s == "PASS" for s in case_states.values()) and not design["gaps"] else locale["partial"],
                open_items="\n".join(open_items) or locale["no_data"])
    # Source gaps are never silently erased to improve the summary. Only the
    # business consequence/next action is published, not raw research locators.
    if design["gaps"]:
        meta["limitations"] += "\n" + "\n".join(public_text(row[4], "gap impact/next check") for row in design["gaps"])
    summary = tuple((locale["labels"][key], meta[key]) for key in
                    ("run_id", "feature", "build", "environment", "period", "scope", "excluded", "conclusion", "limitations", "open_items"))
    _unchanged(captures)
    return CustomerReportData(language, summary, metrics, tuple(rows), tuple(details)), captures


def export_customer_report(run_file, design_dir, output, language, closure_confirmation, evidence_access, screenshots):
    run_file, design_dir, output = Path(run_file), Path(design_dir), Path(output)
    if output.exists():
        raise FileExistsError("Existing customer/master artifact preserved; choose a new output")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Customer output must be .xlsx")
    data, captures = prepare_customer_report(run_file, design_dir, language, closure_confirmation, evidence_access, screenshots)
    if output.resolve() in {path.resolve() for path in captures}:
        raise ValueError("Output/input collision")
    from render_customer_report import render_customer_workbook
    # Renderer writes only an isolated temporary candidate. Final publication is
    # exclusive and occurs after input identity checks; master is never opened for writing.
    with tempfile.TemporaryDirectory(prefix="blend-customer-report-") as temporary:
        candidate = Path(temporary) / "customer.xlsx"
        receipt = render_customer_workbook(data, candidate)
        _unchanged(captures)
        raw = candidate.read_bytes()
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("xb") as handle:
            handle.write(raw)
    return {**receipt, "cases": data.metrics["cases"], "variants": data.metrics["variants"], "language": language,
            "family": "customer-test-report", "version": VERSION, "input_preserved": True}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-file", type=Path, required=True)
    parser.add_argument("--design-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--language", choices=("ja", "vi"), required=True)
    parser.add_argument("--delivery-stdin", action="store_true", required=True)
    args = parser.parse_args()
    try:
        payload = json.load(sys.stdin)
        if not isinstance(payload, dict) or set(payload) != {"closure_confirmation", "evidence_access", "screenshots"}:
            raise ValueError("Delivery JSON requires exactly closure_confirmation, evidence_access, screenshots")
        result = export_customer_report(args.run_file, args.design_dir, args.output, args.language, **payload)
    except (ValueError, TypeError, OSError, ImportError) as error:
        parser.exit(1, f"Customer export blocked: {error}\n")
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
