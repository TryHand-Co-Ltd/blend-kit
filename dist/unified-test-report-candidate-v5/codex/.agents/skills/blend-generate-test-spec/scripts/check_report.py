"""Inspect the existing editable report without writing, repairing or certifying it."""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass
import hashlib
import io
import json
import posixpath
from pathlib import Path
import re
import sys
from urllib.parse import unquote
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from report_model import (REPORT_LAYOUTS, FAMILY, VERSION, SUMMARY_FIELDS, INPUT_FIELDS,
    ClosureConfirmation, EvidenceAccess, _records, _capture, _unchanged, _read_text,
    public_text, shared_url, required, timestamp, prepare_report, display_lines, summary_formulas, detail_backlink_formula,
    recorded_decision_inputs)


@dataclass(frozen=True)
class ScreenshotReview:
    case_id: str
    variant: str
    sha256: str
    caption: str
    reviewed_by: str
    reviewed_at: str
    source: str


def _urls(value):
    return re.findall(r"https?://[^\s<>\])]+", value)


def _formula_reference_key(value, sheet_names):
    """Excel may unquote simple sheet names; never normalize formula string literals."""
    if not isinstance(value, str):
        return value
    parts = re.split(r'("(?:[^"]|"")*")', value)
    for index in range(0, len(parts), 2):
        for name in sheet_names:
            # Only these exact names can legally lose quotes in Excel. Spaces,
            # apostrophes and punctuation keep their original strict spelling.
            if re.fullmatch(r'[^\W\d]\w*', name):
                parts[index] = re.sub(r"(?<![\w.'])" + re.escape(name) + r'!',
                                      lambda _: "'" + name + "'!", parts[index])
    return ''.join(parts)


def _xml_value(value, field):
    """Validate metadata literals, including repeated URL escaping, without echoing values."""
    if not value:
        return
    decoded = value
    for _ in range(5):
        # These are public XML schema identifiers, not evidence links. The strict
        # lexical form excludes query strings, encoded locators and credentials.
        if re.fullmatch(r'https?://(?:schemas\.openxmlformats\.org|schemas\.microsoft\.com|www\.w3\.org|purl\.oclc\.org|purl\.org)/[A-Za-z0-9._/-]+', decoded):
            from report_model import PRIVATE
            if not PRIVATE.search(decoded):
                return
        public_text(decoded, field)  # No dummy-input exception for XML attributes.
        next_decoded = unquote(decoded)
        if next_decoded == decoded:
            return
        decoded = next_decoded
    raise ValueError(f'Excessive metadata encoding cannot be validated: {field}')


def _xml_attributes(root, name, urls):
    for element in root.iter():
        tag = element.tag.rsplit('}', 1)[-1]
        for key, value in element.attrib.items():
            attribute = key.rsplit('}', 1)[-1]
            field = f'{name}:{tag}@{attribute}'
            # ContentType is a MIME identifier. Do not mistake the registered
            # application/vnd... prefix for an application source-directory path.
            if name == '[Content_Types].xml' and attribute == 'ContentType':
                if re.fullmatch(r'(?:application/(?:xml|vnd\.openxmlformats-(?:officedocument|package)\.[A-Za-z0-9.-]+(?:\+xml)?)|image/(?:png|jpeg))', value):
                    continue
                raise ValueError(f'Unsupported content type: {field}')
            _xml_value(value, field)
            if not re.fullmatch(r'https?://(?:schemas\.openxmlformats\.org|schemas\.microsoft\.com|www\.w3\.org|purl\.oclc\.org|purl\.org)/[A-Za-z0-9._/-]+', value):
                urls.update(_urls(value))


def _package(raw, data, phase, payload, gaps):
    """Inspect saved relationships, metadata, text and image bytes, never strip them."""
    urls, media, captions = set(), {}, []
    locale = REPORT_LAYOUTS[data.language]
    allowed_formulas = {value.removeprefix('=') for value in summary_formulas(data).values()}
    allowed_formulas.update(detail_backlink_formula(row.identity, data).removeprefix('=') for row in data.rows)
    allowed_formulas.update('"' + label + '"' for label in locale['statuses'].values())
    allowed_formulas.add('"' + ','.join(locale['statuses'].values()) + '"')
    allowed_formulas = {_formula_reference_key(value, locale['sheets']) for value in allowed_formulas}
    with ZipFile(io.BytesIO(raw)) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise ValueError("Duplicate ZIP entries")
        for name in names:
            if name.startswith(('xl/externalLinks/', 'xl/embeddings/', 'customXml/', 'xl/activeX/')) or name.endswith(('vbaProject.bin', '.vml')):
                raise ValueError(f"Unsupported workbook component: {name}")
            content = archive.read(name)
            if name.startswith('xl/media/'):
                from PIL import Image
                with Image.open(io.BytesIO(content)) as image:
                    if image.format not in ('PNG', 'JPEG') or getattr(image, 'n_frames', 1) != 1:
                        raise ValueError('Only static PNG/JPEG evidence is supported')
                    if image.getexif() or any(k not in ('dpi', 'jfif', 'jfif_version', 'jfif_unit', 'jfif_density') for k in image.info):
                        raise ValueError(f'Image metadata must be reviewed and removed in the same workbook: {name}')
                with Image.open(io.BytesIO(content)) as image:
                    image.verify()
                media[name] = hashlib.sha256(content).hexdigest()
            elif name.endswith(('.xml', '.rels')):
                root = ET.fromstring(content)
                for _, (_, namespace_uri) in ET.iterparse(io.BytesIO(content), events=('start-ns',)):
                    _xml_value(namespace_uri, f'{name}:namespace')
                _xml_attributes(root, name, urls)
                if name == 'xl/workbook.xml':
                    privacy = root.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}workbookPr')
                    if privacy is None or privacy.get('filterPrivacy') not in ('1', 'true'):
                        raise ValueError('Workbook privacy setting is not enabled: workbookPr@filterPrivacy')
                if name.endswith('.rels'):
                    for rel in root:
                        target = rel.get('Target', '')
                        if rel.get('TargetMode') == 'External':
                            urls.add(shared_url(target, name))
                        elif re.match(r'(?i)(?:https?|file):', unquote(target)):
                            raise ValueError(f'Invalid relationship target: {name}')
                        else:
                            owner = name.replace('/_rels/', '/').removesuffix('.rels')
                            base = '' if name == '_rels/.rels' else posixpath.dirname(owner)
                            resolved = posixpath.normpath(posixpath.join(base, target)) if not target.startswith('/') else target[1:]
                            if resolved not in names:
                                raise ValueError(f'Invalid or missing internal relationship: {name}')
                else:
                    for element in root.iter():
                        tag = element.tag.rsplit('}', 1)[-1]
                        if tag in ('f', 'formula', 'formula1', 'formula2') and _formula_reference_key(element.text, locale['sheets']) not in allowed_formulas:
                            raise ValueError(f'Unrecognized formula in saved package: {name}')
                        # Formula text is compared against the canonical schema separately.
                        if element.text and tag not in ('f', 'formula', 'formula1', 'formula2'):
                            value = element.text
                            if value.strip():
                                dummy_input = '[dummy-input: ' in value and any(value in r.conditions for r in data.rows)
                                public_text(value, name, dummy_input=dummy_input)
                                scan = re.sub(r'\[dummy-input: [^\]\n]+\]', '', value) if dummy_input else value
                                urls.update(_urls(scan))
                        if tag == 'cNvPr':
                            captions.append(element.get('descr') or '')
            else:
                raise ValueError(f'Unsupported binary package entry: {name}')
    for url in urls:
        shared_url(url, 'saved workbook URL')
    if phase == 'complete':
        closure = _records(payload.get('closure_confirmation'), ClosureConfirmation)
        for key in ('closed_by', 'source', 'audience'):
            required(getattr(closure, key), f'closure.{key}')
        timestamp(closure.closed_at, 'closure.closed_at')
        accesses = {}
        for item in payload.get('evidence_access', []):
            access = _records(item, EvidenceAccess)
            shared_url(access.url, 'access.url')
            for key in ('audience', 'verified_by', 'source'):
                required(getattr(access, key), f'access.{key}')
            timestamp(access.verified_at, 'access.verified_at')
            if access.audience != closure.audience or access.method not in ('human-attestation', 'recipient-session') or access.url in accesses:
                raise ValueError('Evidence access audience, method or uniqueness mismatch')
            accesses[access.url] = access
        for url in sorted(urls - accesses.keys()):
            gaps.append(f'Access not verified for link: {url}')
        reviews = [_records(item, ScreenshotReview) for item in payload.get('screenshots', [])]
        if len(reviews) != len(media):
            gaps.append('Each embedded image requires one matching screenshot review')
        digests = Counter(media.values())
        reviewed = Counter()
        identities = {(r.case_id, r.variant) for r in data.rows}
        for review in reviews:
            if (review.case_id, review.variant) not in identities or not re.fullmatch('[0-9a-f]{64}', review.sha256):
                raise ValueError('Unknown screenshot identity or invalid digest')
            for key in ('caption', 'reviewed_by', 'source'):
                required(getattr(review, key), f'screenshot.{key}')
            public_text(review.caption, 'screenshot.caption')
            timestamp(review.reviewed_at, 'screenshot.reviewed_at')
            if review.caption not in captions:
                gaps.append('Screenshot caption must match saved image alternative text')
            reviewed[review.sha256] += 1
        if reviewed != digests:
            gaps.append('Screenshot reviews do not match the saved image bytes')
    elif media:
        gaps.append('Embedded images require review before completion')
    return urls


def check_report(report: Path, source_dir: Path, language='vi', phase='in-progress', *,
                 closure_confirmation=None, evidence_access=(), screenshots=()):
    from openpyxl import load_workbook
    from render_report import build_workbook
    if phase not in ('in-progress', 'complete'):
        raise ValueError('Unknown check phase')
    report = Path(report)
    data, captures = prepare_report(Path(source_dir), language)
    template = Path(__file__).resolve().parents[1] / 'assets' / REPORT_LAYOUTS[language]['template']
    captures[template] = _capture(template)
    captures[report] = _capture(report)
    raw = captures[report][0]
    book = load_workbook(io.BytesIO(raw), data_only=False)
    props = {p.name: p.value for p in book.custom_doc_props}
    if props != {'TemplateFamily': FAMILY, 'TemplateVersion': VERSION, 'Language': language}:
        raise ValueError('Unsupported report schema; legacy workbooks require separately authorized migration')
    expected = build_workbook(data)
    locale = REPORT_LAYOUTS[language]
    summary, results = book[locale['sheets'][0]], book[locale['sheets'][1]]
    if book.sheetnames not in (expected.sheetnames, list(locale['sheets'])):
        raise ValueError('Report sheets differ from matching design')
    if locale['sheets'][2] in book.sheetnames and locale['sheets'][2] not in expected.sheetnames:
        added = expected.create_sheet(locale['sheets'][2])
        for col, label in enumerate(locale['detail_headers'], 1):
            added.cell(locale['detail_header_row'], col, label)
    if any(sheet.sheet_state != 'visible' for sheet in book):
        raise ValueError('Hidden worksheets are unsupported')
    gaps, layout_notes = [], []
    for field in INPUT_FIELDS:
        value = _read_text(summary.cell(SUMMARY_FIELDS[field], 2), field, optional=True)
        if value:
            public_text(value, field)
            cell = summary.cell(SUMMARY_FIELDS[field], 2)
            height = summary.row_dimensions[cell.row].height
            height = summary.sheet_format.defaultRowHeight or 15 if height is None else height
            if height < (value.count('\n') + 1) * (cell.font.sz or 11):
                gaps.append(f'{summary.title}!{cell.coordinate}: row is too short for its explicit text lines')
            elif display_lines(value, 100) * 17 + 10 > height:
                layout_notes.append(f'{summary.title}!{cell.coordinate}: estimated wrapping exceeds row height; verify visibility in the spreadsheet application')
        else:
            gaps.append(f'{summary.title}!B{SUMMARY_FIELDS[field]}: missing {field}')
    if phase == 'complete' and summary.cell(SUMMARY_FIELDS['period'], 2).value:
        executed_at = timestamp(summary.cell(SUMMARY_FIELDS['period'], 2).value, 'execution timestamp')
        closure = _records(closure_confirmation, ClosureConfirmation)
        if executed_at > timestamp(closure.closed_at, 'closure.closed_at'):
            raise ValueError('Execution timestamp is after closure')
    # All source/formula cells are immutable. Result rows may be sorted together.
    start = locale['start_row']
    wanted = {r.identity: r for r in data.rows}
    image_caption_cells = set()
    image_bindings = []
    for sheet in book:
        for picture in sheet._images:
            if sheet.title != locale['sheets'][2]:
                raise ValueError('Evidence images belong in the detail sheet')
            anchor = picture.anchor._from
            caption_row = anchor.row  # image is one row below caption, OOXML is zero based
            if anchor.col != 2 or caption_row < locale['detail_start_row']:
                raise ValueError('Evidence image must be anchored in C immediately below its caption row')
            identity, label, caption = [sheet.cell(caption_row, col).value for col in (1, 2, 3)]
            if identity not in wanted or label != locale['evidence_link']:
                raise ValueError('Image caption row requires exact variant identity and evidence label')
            public_text(caption, 'image caption')
            image_bindings.append((identity, caption, hashlib.sha256(picture._data()).hexdigest()))
            image_caption_cells.update((sheet.title, sheet.cell(caption_row, col).coordinate) for col in (1, 2, 3))
    if phase == 'complete':
        reviews = [_records(item, ScreenshotReview) for item in screenshots]
        expected_bindings = Counter((f'{r.case_id} / {r.variant}', r.caption, r.sha256) for r in reviews)
        if expected_bindings != Counter(image_bindings):
            raise ValueError('Screenshot review identity/caption/digest differs from saved image placement')
    seen, seen_folded, states = set(), set(), []
    pristine_results = expected[results.title]
    pristine_rows = {pristine_results.cell(row, 1).value: row for row in range(start, start + len(data.rows))}
    for row in range(start, results.max_row + 1):
        if not any(results.cell(row, col).value is not None for col in range(1, 8)):
            continue
        identity = _read_text(results.cell(row, 1), f'{results.title}!A{row}')
        if identity not in wanted or identity.casefold() in seen_folded:
            raise ValueError(f'Duplicate or unknown variant: {results.title}!A{row}')
        seen.add(identity)
        seen_folded.add(identity.casefold())
        record = wanted[identity]
        original_row = pristine_rows[identity]
        for col in range(1, 5):
            actual, original = results.cell(row, col), pristine_results.cell(original_row, col)
            if actual.value != original.value or actual.data_type != original.data_type:
                raise ValueError(f'Changed design value/type: {results.title}!{actual.coordinate}')
            if original.hyperlink and (not actual.hyperlink or actual.hyperlink.location != original.hyperlink.location):
                raise ValueError(f'Changed detail link: {results.title}!{actual.coordinate}')
        actual, status, evidence = [_read_text(results.cell(row, col), f'{results.title}!{results.cell(row, col).coordinate}', optional=(col != 6)) for col in (5, 6, 7)]
        reverse = {v: k for k, v in locale['statuses'].items()}
        if status not in reverse:
            raise ValueError(f'Invalid pasted status: {results.title}!F{row}')
        token = reverse[status]
        for value, column in ((actual, 'E'), (evidence, 'G')):
            if value.strip():
                public_text(value, f'{results.title}!{column}{row}')
        links = _urls(evidence)
        if results.cell(row, 7).hyperlink and results.cell(row, 7).hyperlink.target:
            links.append(shared_url(results.cell(row, 7).hyperlink.target, f'{results.title}!G{row}'))
        reason = re.sub(r'https?://[^\s<>\])]+', '', evidence).strip()
        valid = token in ('PASS', 'FAIL') and record.eligible and recorded_decision_inputs(actual, evidence)
        if token in ('PASS', 'FAIL'):
            if not recorded_decision_inputs(actual, evidence):
                gaps.append(f'{identity}: nonblank actual and a standalone HTTPS URL on the first evidence line required')
            if not record.eligible:
                gaps.append(f'{identity}: expected result or preparation remains unconfirmed')
        if token != 'PASS' and not reason:
            gaps.append(f'{identity}: reason and next action required in evidence/issues')
        if token == 'NOT RUN':
            gaps.append(f'{identity}: not executed')
        required_height = max(display_lines(actual, 32), display_lines(evidence, 21)) * 17 + 10
        height = results.row_dimensions[row].height
        height = results.sheet_format.defaultRowHeight or 15 if height is None else height
        minimum_height = max(((value.count('\n') + 1) * (results.cell(row, col).font.sz or 11)
                              for value, col in ((actual, 5), (evidence, 7)) if value), default=0)
        if height < minimum_height:
            gaps.append(f'{identity}: row is too short for its explicit actual/evidence text lines')
        elif required_height > height:
            layout_notes.append(f'{identity}: estimated wrapping exceeds row height; verify visibility in the spreadsheet application')
        states.append((identity, token, valid))
    if seen != set(wanted):
        raise ValueError('Missing design variants in report')
    if results.max_row > start + len(data.rows) - 1 and any(results.cell(r, c).value is not None for r in range(start + len(data.rows), results.max_row + 1) for c in range(1, 8)):
        raise ValueError('Result table has rows outside its formula range')
    for sheet in book:
        pristine = expected[sheet.title]
        if sheet == results:
            if sheet.freeze_panes != pristine.freeze_panes or sheet.auto_filter.ref != pristine.auto_filter.ref:
                raise ValueError('Result freeze/filter range differs from report schema')
            expected_validation = [(v.type, v.formula1, str(v.sqref), v.allow_blank, v.showErrorMessage)
                                   for v in pristine.data_validations.dataValidation]
            actual_validation = [(v.type, v.formula1, str(v.sqref), v.allow_blank, v.showErrorMessage)
                                 for v in sheet.data_validations.dataValidation]
            if actual_validation != expected_validation:
                raise ValueError('Result status validation differs from report schema')
        for row in sheet.iter_rows(max_row=max(sheet.max_row, pristine.max_row),
                                   max_col=max(sheet.max_column, pristine.max_column)):
            for cell in row:
                editable = (sheet == results and start <= cell.row < start + len(data.rows) and 5 <= cell.column <= 7) or (sheet == summary and cell.column == 2 and cell.row in [SUMMARY_FIELDS[f] for f in INPUT_FIELDS])
                if editable or (sheet == results and start <= cell.row < start + len(data.rows) and cell.column <= 4):
                    continue
                original = pristine[cell.coordinate]
                if (sheet.title, cell.coordinate) in image_caption_cells and original.value is None:
                    continue
                same_value = (_formula_reference_key(cell.value, locale['sheets']) == _formula_reference_key(original.value, locale['sheets'])
                              if cell.data_type == original.data_type == 'f' else cell.value == original.value)
                if not same_value or cell.data_type != original.data_type:
                    raise ValueError(f'Changed immutable cell/formula: {sheet.title}!{cell.coordinate}')
                if cell.comment:
                    public_text(cell.comment.text, f'{sheet.title}!{cell.coordinate} comment')
        if any(d.hidden for d in sheet.row_dimensions.values()) or any(d.hidden for d in sheet.column_dimensions.values()):
            raise ValueError('Hidden rows/columns are unsupported')
    urls = _package(raw, data, phase, {'closure_confirmation': closure_confirmation,
                    'evidence_access': evidence_access, 'screenshots': screenshots}, gaps)
    _unchanged(captures)
    counts = Counter(token for _, token, _ in states)
    evaluated = sum(valid for _, _, valid in states)
    passed = sum(valid and token == 'PASS' for _, token, valid in states)
    case_results = {}
    for case_id in dict.fromkeys(row.case_id for row in data.rows):
        members = [(token, valid) for identity, token, valid in states if wanted[identity].case_id == case_id]
        case_results[case_id] = ('FAIL' if any(token == 'FAIL' and valid for token, valid in members) else
            'PASS' if all(token == 'PASS' and valid for token, valid in members) else
            'SKIPPED' if all(token == 'SKIPPED' for token, _ in members) else
            'BLOCKED' if any(token == 'BLOCKED' for token, _ in members) else
            'NOT RUN' if all(token == 'NOT RUN' for token, _ in members) else 'INCOMPLETE')
    return {'family': FAMILY, 'version': VERSION, 'language': language, 'phase': phase,
            'complete': phase == 'complete' and not gaps, 'read_only': True,
            'cases': len(set(r.case_id for r in data.rows)), 'variants': len(states),
            'states': dict(counts), 'case_results': case_results, 'evaluated': evaluated, 'passed': passed,
            'pass_rate': passed / evaluated if evaluated else None,
            'completion_rate': evaluated / len(states) if states else None,
            'gaps': gaps, 'layout_notes': layout_notes, 'links': len(urls)}


def main():
    # Keep encoding local to this CLI; business strings are never re-decoded.
    for stream in (sys.stdin, sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--report', type=Path, required=True)
    parser.add_argument('--source-dir', type=Path, required=True)
    parser.add_argument('--language', choices=('ja', 'vi'), default='vi')
    parser.add_argument('--phase', choices=('in-progress', 'complete'), default='in-progress')
    args = parser.parse_args()
    try:
        payload = json.load(sys.stdin) if args.phase == 'complete' else {}
        if args.phase == 'complete' and (not isinstance(payload, dict) or set(payload) != {'closure_confirmation', 'evidence_access', 'screenshots'}):
            raise ValueError('Completion stdin requires closure_confirmation, evidence_access and screenshots')
        result = check_report(args.report, args.source_dir, args.language, args.phase, **payload)
    except (ValueError, TypeError, OSError, ImportError, KeyError) as error:
        parser.exit(1, f'Report check blocked: {error}\n')
    print(json.dumps(result, ensure_ascii=False))
    return 1 if args.phase == 'complete' and not result['complete'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
