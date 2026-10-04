"""Render validated customer records with the versioned customer XLSX template.

This module does not read a working master, translate requirements, or attest access.
The customer_report entry point validates those boundaries before calling it.
"""
from __future__ import annotations

import io
import math
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.packaging.core import DocumentProperties

from customer_report import CUSTOMER_LAYOUTS, VERSION, CustomerReportData, display_lines as _lines

INK = '172033'
LINK = '1D4ED8'
SUMMARY_KEYS = ('run_id', 'feature', 'build', 'environment', 'period', 'scope',
                'excluded', 'conclusion', 'limitations', 'open_items')


def _text(cell, value):
    """All source text stays literal, including leading equals signs."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'
    cell.font = Font(name=cell.parent['A1'].font.name, size=11, color=INK)
    cell.alignment = Alignment(vertical='top', wrap_text=True)


def _chunks(value, width=95, max_lines=17):
    """Split only display rows; concatenating chunks preserves every character."""
    value = str(value)
    if _lines(value, width) <= max_lines:
        return [value]
    result, chunk = [], ''
    for char in value:
        if chunk and _lines(chunk + char, width) > max_lines:
            result.append(chunk)
            chunk = ''
        chunk += char
    if chunk:
        result.append(chunk)
    return result


def _height(sheet, row, values, widths):
    lines = max(_lines(value, width) for value, width in zip(values, widths))
    sheet.row_dimensions[row].height = min(409, max(30, lines * 17 + 10))


def _jump(cell, sheet, row, caption=None):
    if caption is not None:
        _text(cell, caption)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet.replace(chr(39), chr(39)*2)}'!A{row}", display=str(cell.value))
    cell.font = Font(name=cell.parent['A1'].font.name, size=11, color=LINK, underline='single')


def _finish(sheet, end_column, end_row, header=None):
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = False
    sheet.print_area = f'A1:{end_column}{end_row}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    if header:
        sheet.freeze_panes = f'B{header + 1}'
        sheet.print_title_rows = f'{header}:{header}'


def _image_descriptions(raw, captions):
    """Openpyxl lacks a public picture-alt-text setter; patch only cNvPr."""
    with ZipFile(io.BytesIO(raw)) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    descriptions = iter(captions)
    for name, content in list(entries.items()):
        if name.startswith('xl/drawings/drawing') and name.endswith('.xml'):
            drawing = ET.fromstring(content)
            for prop in drawing.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}cNvPr'):
                caption = next(descriptions)
                prop.set('descr', caption)
                prop.set('title', caption)
            entries[name] = ET.tostring(drawing, encoding='utf-8', xml_declaration=True)
    output = io.BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as target:
        for name, content in entries.items():
            target.writestr(name, content)
    return output.getvalue()


def render_customer_workbook(data: CustomerReportData, output: Path) -> dict:
    output = Path(output)
    if output.exists():
        raise FileExistsError('Existing output is preserved')
    locale = CUSTOMER_LAYOUTS[data.language]
    font = locale['font']
    asset = Path(__file__).resolve().parents[1] / 'assets' / locale['template']
    book = load_workbook(asset)
    props = {prop.name: prop.value for prop in book.custom_doc_props}
    if props != {'TemplateFamily': 'customer-test-report', 'TemplateVersion': VERSION, 'Language': data.language}:
        raise ValueError('Customer template family/version/language mismatch')
    if book.sheetnames != list(locale['sheets']):
        raise ValueError('Customer template sheet schema mismatch')
    summary, results, details = (book[name] for name in locale['sheets'])
    if tuple(results.cell(locale['header_row'], col).value for col in range(1, 8)) != locale['headers']:
        raise ValueError('Customer template results headers mismatch')
    if tuple(details.cell(locale['detail_header_row'], col).value for col in range(1, 4)) != locale['detail_headers']:
        raise ValueError('Customer template details headers mismatch')
    book.properties = DocumentProperties(creator='BLEND', lastModifiedBy='BLEND', title=locale['title'])
    summary.delete_rows(locale.get('summary_start_row', 4), max(summary.max_row, 1))
    row = locale.get('summary_start_row', 4)
    for label, value in data.summary:
        for chunk in _chunks(value):
            _text(summary.cell(row, 1), label)
            summary.cell(row, 1).font = Font(name=font, size=11, bold=True, color=INK)
            _text(summary.cell(row, 2), chunk)
            _height(summary, row, (label, chunk), (38, 100))
            row += 1
    row += 1
    metric_rows = {}
    for key in ('cases', 'variants', 'selected_variants', 'evaluated_variants', 'passed_variants', 'confirmed_variants', 'pass_rate', 'completion_rate'):
        metric_rows[key] = row
        _text(summary.cell(row, 1), locale['labels'][key])
        _text(summary.cell(row, 2), data.metrics[key] if data.metrics[key] is not None else locale['no_data'])
        summary.cell(row, 2).alignment = Alignment(horizontal='left', vertical='top')
        summary.cell(row, 2).number_format = '0.0%' if key.endswith('_rate') else '#,##0'
        _height(summary, row, (locale['labels'][key], ''), (38, 100))
        row += 1
    for rate, numerator, denominator in [('pass_rate','passed_variants','evaluated_variants'),('completion_rate','evaluated_variants','confirmed_variants')]:
        cell = summary.cell(metric_rows[rate], 2)
        cell.value = f'=IF(B{metric_rows[denominator]}=0,"{locale["no_data"]}",B{metric_rows[numerator]}/B{metric_rows[denominator]})'
    for key in ('case_statuses', 'variant_statuses'):
        row += 1
        _text(summary.cell(row, 1), locale['labels'][key])
        summary.cell(row, 1).font = Font(name=font, size=11, bold=True, color=INK)
        row += 1
        for status in ('PASS', 'FAIL', 'BLOCKED', 'SKIPPED', 'NOT RUN'):
            _text(summary.cell(row, 1), status)
            _text(summary.cell(row, 2), data.metrics[key].get(status, 0))
            summary.cell(row, 2).alignment = Alignment(horizontal='left', vertical='top')
            summary.row_dimensions[row].height = 25
            row += 1
    _finish(summary, 'B', row - 1)
    detail_rows, captions = {}, []
    row = locale['detail_start_row']
    records = {record.identity: record for record in data.rows}
    for detail in data.details:
        detail_rows[detail.identity] = row
        record = records[detail.identity]
        for label, value in ((locale['headers'][1], record.screen_function), *detail.blocks):
            for chunk in _chunks(value):
                for col, text in enumerate((detail.identity, label, chunk), 1):
                    _text(details.cell(row, col), text)
                _height(details, row, (detail.identity, label, chunk), (20, 28, 100))
                row += 1
        for index, url in enumerate(record.evidence, 1):
            for col, text in enumerate((detail.identity, locale['evidence_link'], f'{locale["evidence_link"]} {index} — {detail.identity}'), 1):
                _text(details.cell(row, col), text)
            details.cell(row, 3).hyperlink = url
            details.cell(row, 3).font = Font(name=font, size=11, color=LINK, underline='single')
            details.row_dimensions[row].height = 30
            row += 1
        if record.bug:
            for chunk in _chunks(f'{locale["bug_link"]} — {detail.identity}' if record.bug.startswith('https://') else record.bug):
                for col, text in enumerate((detail.identity, locale['bug_link'], chunk), 1):
                    _text(details.cell(row, col), text)
                if record.bug.startswith('https://'):
                    details.cell(row, 3).hyperlink = record.bug
                    details.cell(row, 3).font = Font(name=font, size=11, color=LINK, underline='single')
                _height(details, row, (detail.identity, locale['bug_link'], chunk), (20, 28, 100))
                row += 1
        for image_data in detail.images:
            for col, text in enumerate((detail.identity, locale['parts']['evidence'], image_data.caption), 1):
                _text(details.cell(row, col), text)
            _height(details, row, (detail.identity, locale['parts']['evidence'], image_data.caption), (20, 28, 100))
            row += 1
            picture = Image(io.BytesIO(image_data.content))
            scale = min(1, 720 / picture.width)
            picture.width, picture.height = picture.width * scale, picture.height * scale
            details.add_image(picture, f'C{row}')
            captions.append(detail.identity + ': ' + image_data.caption)
            image_rows = math.ceil(picture.height / 24)
            for offset in range(image_rows):
                details.row_dimensions[row + offset].height = 18
            row += image_rows + 1
    if data.details:
        _finish(details, 'C', row - 1, locale['detail_header_row'])
    else:
        book.remove(details)
    for row, record in enumerate(data.rows, locale['start_row']):
        if record.needs_detail and record.identity not in detail_rows:
            raise ValueError('Missing material testcase details')
        urls = (*record.evidence, *((record.bug,) if record.bug.startswith('https://') else ()))
        evidence = (f'{locale["evidence_link"]} — {record.identity}' if record.evidence else
                    f'{locale["bug_link"]} — {record.identity}' if urls else locale['no_data'])
        if record.bug and not record.bug.startswith('https://'):
            evidence += '\n' + record.bug
        values = [record.identity, record.screen_function, record.conditions, record.expected, record.actual, record.status, evidence]
        for col, value in enumerate(values, 1):
            _text(results.cell(row, col), value)
        if record.needs_detail:
            for col in (2, 3, 4, 5):
                if _lines(values[col - 1], 22 if col == 2 else 32) > 12:
                    values[col - 1] = locale['detail_link']
                    _jump(results.cell(row, col), details.title, detail_rows[record.identity], locale['detail_link'])
            _jump(results.cell(row, 1), details.title, detail_rows[record.identity])
            _jump(details.cell(detail_rows[record.identity], 1), results.title, row)
        if len(urls) > 1 or _lines(evidence, 21) > 12:
            if record.identity not in detail_rows:
                raise ValueError('Multiple evidence links require a details layout')
            values[6] = locale['detail_link'] + '\n' + locale['headers'][6]
            _jump(results.cell(row, 7), details.title, detail_rows[record.identity], values[6])
        elif urls:
            results.cell(row, 7).hyperlink = urls[0]
            results.cell(row, 7).font = Font(name=font, size=11, color=LINK, underline='single')
        _height(results, row, values, (16, 22, 32, 32, 32, 10, 21))
    last = max(locale['start_row'], locale['start_row'] + len(data.rows) - 1)
    results.auto_filter.ref = f'A{locale["header_row"]}:G{last}'
    _finish(results, 'G', last, locale['header_row'])
    for status, fill, color in [('FAIL','FEE2E2','991B1B'),('BLOCKED','FEF3C7','92400E'),('SKIPPED','F1F5F9','475569'),('NOT RUN','F1F5F9','475569')]:
        results.conditional_formatting.add(f'F{locale["start_row"]}:F{last}', CellIsRule(operator='equal', formula=[f'"{status}"'], fill=PatternFill('solid', fgColor=fill), font=Font(name=font, size=11, color=color, bold=True)))
    book.active = 0
    buffer = io.BytesIO()
    book.save(buffer)
    raw = _image_descriptions(buffer.getvalue(), captions) if captions else buffer.getvalue()
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('xb') as handle:
        handle.write(raw)
    return {'sheets': book.sheetnames, 'rows': len(data.rows), 'detail_cases': len(data.details), 'images': len(captions)}
