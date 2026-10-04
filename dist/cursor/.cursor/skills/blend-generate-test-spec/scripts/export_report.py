"""Project fixed-template BLEND Test Spec Markdown into the bundled VI workbook.

No application tests, translation, source approval or spreadsheet recalculation.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import io
import math
import re
import unicodedata
from pathlib import Path

VERSION = "1.0.0"
HEADINGS = {
    "vi": {"scope-and-approach": ["Mục tiêu và căn cứ", "Phạm vi", "Chuẩn bị và cách chạy", "Coverage và gaps"],
           "test-cases": ["Quy ước và context", "Các testcase"], "test-data": ["Quy ước dữ liệu", "Fixtures"]},
    "ja": {"scope-and-approach": ["目的・根拠", "対象範囲", "準備・実行方針", "Coverage・Gaps"],
           "test-cases": ["規約・コンテキスト", "テストケース"], "test-data": ["データ規約", "Fixtures"]},
}
CASE_KEYS = ["function", "source", "priority", "basis", "readiness", "gap", "configuration", "trigger",
             "observation", "actor", "fixture", "action", "expected", "preservation", "proof", "reset"]
CASE_LABELS = {
    "vi": ["Chức năng", "Căn cứ", "Priority", "Căn cứ kỳ vọng", "Readiness", "Gap", "Cấu hình", "Kích hoạt",
           "Quan sát", "Actor và quyền", "Fixture", "Thao tác", "Expected", "Bảo toàn", "Bằng chứng", "Reset"],
    "ja": ["機能", "根拠", "Priority", "期待根拠", "Readiness", "Gap", "設定", "起動", "観測", "Actor・権限",
           "Fixture", "操作", "Expected", "維持状態", "証拠", "Reset"],
}
FIXTURE_LABELS = {"vi": ["Vai trò", "Phạm vi", "Trạng thái", "Giá trị", "Target và non-target", "Tạo", "Kiểm tra", "Reset"],
                  "ja": ["役割", "範囲", "状態", "値", "Target・non-target", "作成", "確認", "Reset"]}
SCOPE_LABELS = {"vi": [["Revision", "Nguồn", "Căn cứ research", "Baseline code", "Mode", "Giới hạn"],
                        ["Trong scope", "Ngoài scope", "Vai trò"], ["Môi trường", "Chuẩn bị", "Chọn lượt chạy", "Điều kiện bắt đầu", "Điều kiện kết thúc", "Bằng chứng"]],
                "ja": [["Revision", "Source", "Research根拠", "Code baseline", "Mode", "調査限界"],
                        ["対象", "対象外", "役割"], ["環境", "準備", "実行選択", "開始条件", "終了条件", "証拠"]]}
RUN_HEADERS = ["Run ID", "Case ID", "Variant", "Nhóm nghiệp vụ", "Priority", "Căn cứ kỳ vọng", "Readiness",
               "Trong scope", "Status", "Actual", "Evidence", "Bug", "Tester", "Executed at", "Timezone",
               "Lý do / blocker", "Kiểm tra kết quả", "Case status", "Thiết kế"]
CASE_HEADERS = ["Case ID", "Nhóm nghiệp vụ", "Tiêu đề", *CASE_LABELS["vi"], "Steps", "Variants"]
DATA_HEADERS = ["Fixture ID", *FIXTURE_LABELS["vi"]]
RUN_START = 25
CASE_START = 8
DATA_START = 8


def cell_value(value: str) -> str:
    value = html.unescape(value.replace("<br>", "\n").replace("<br/>", "\n"))
    if len(value) > 32767:
        raise ValueError("Source exceeds Excel cell limit; split meaning using template steps/fixtures")
    if any(ord(ch) < 32 and ch not in "\n\t" for ch in value):
        raise ValueError("Unsupported control character")
    return value


def read_table(block: str, headers: list[str]) -> tuple[list[list[str]], str]:
    """Read one exact Markdown table; reject ragged/compound rows rather than guessing."""
    lines = block.strip().splitlines()
    if len(lines) < 2 or not lines[0].startswith("|"):
        raise ValueError(f"Missing table: {headers}")
    consumed = 0
    parsed = []
    for line in lines:
        if not line.strip().startswith("|"):
            break
        if not line.strip().endswith("|") or "\\|" in line:
            raise ValueError("Use &#124; for literal pipes and <br> for line breaks")
        row = [part.strip() for part in line.strip()[1:-1].split("|")]
        if len(row) != len(headers):
            raise ValueError("Ragged/compound Markdown table row")
        parsed.append(row)
        consumed += 1
    if parsed[0] != headers or not all(re.fullmatch(r":?-{3,}:?", value) for value in parsed[1]):
        raise ValueError(f"Table header/separator mismatch: {headers}")
    if any(not value for row in parsed[2:] for value in row):
        raise ValueError("Empty required table cell")
    return [[cell_value(value) for value in row] for row in parsed[2:]], "\n".join(lines[consumed:]).strip()


def fields(block: str, labels: list[str], *, exact: bool = True) -> tuple[dict[str, str], str]:
    rows, rest = read_table(block, ["Field", "Value"])
    actual = [row[0] for row in rows]
    if len(set(actual)) != len(actual) or (actual != labels if exact else any(item not in labels for item in actual)):
        raise ValueError(f"Required field order/membership mismatch: {labels}")
    return dict(rows), rest


def document(path: Path, family: str, language: str, *, raw: bytes | None = None) -> tuple[str, list[str]]:
    captured = path.read_bytes() if raw is None else raw
    text = captured.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
    marker = f"<!-- blend-template: {family}@{VERSION} -->"
    if re.findall(r"<!--\s*blend-template:[^>]*-->", text) != [marker]:
        raise ValueError(f"Missing/wrong template family/version: {path.name}")
    if "<!-- AUTHORING:" in text or re.search(r"\[(?:verified value|same design revision|Feature|source clause)\]", text):
        raise ValueError(f"Unfinished template: {path.name}")
    if "```" in text or "~~~" in text:
        raise ValueError("Fenced schema is unsupported; fields must be actual document content")
    titles = re.findall(r"^## (.+)$", text, re.M)
    if titles != HEADINGS[language][family]:
        raise ValueError(f"Required H2 order mismatch: {path.name}")
    sections = re.split(r"^## .+$", text, flags=re.M)[1:]
    return text, [section.strip() for section in sections]


def identifier(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_.-]*", value):
        raise ValueError(f"Invalid/compound identifier: {value}")
    return value


def references(value: str, allowed: set[str]) -> list[str]:
    """Canonical bare-ID list; no deduplication, annotations or implicit separators."""
    values = value.split(", ")
    if len(set(values)) != len(values) or any(item not in allowed for item in values):
        raise ValueError("Reference list must contain distinct defined bare IDs separated by comma-space")
    return values


def parse_sources(source: Path, language: str = "vi", *, captured: dict[str, bytes] | None = None) -> dict:
    texts = {}
    sections = {}
    for family in HEADINGS[language]:
        name = f"{family}.{language}.md"
        texts[family], sections[family] = document(source / name, family, language,
                                                  raw=captured[name] if captured is not None else None)
    scope = []
    for part, labels in zip(sections["scope-and-approach"][:3], SCOPE_LABELS[language]):
        data, rest = fields(part, labels)
        if rest:
            raise ValueError("Scope field sections cannot contain unprojected prose")
        scope.append(data)
    coverage_parts = re.split(r"^### (Coverage|Gaps)$", sections["scope-and-approach"][3], flags=re.M)
    if len(coverage_parts) != 5 or coverage_parts[0].strip() or coverage_parts[1::2] != ["Coverage", "Gaps"]:
        raise ValueError("Coverage/Gaps sections missing, duplicate or out of order")
    coverage, rest = read_table(coverage_parts[2], ["Clause / anchor", "Obligation / branch", "Case / variant / gap", "Rationale / proof limit"])
    if rest or not coverage:
        raise ValueError("Coverage must map source obligations")
    gaps, rest = read_table(coverage_parts[4], ["Gap ID", "Kind", "Known obligation / source", "Missing decision / proof", "Impact / next check"])
    if rest and (gaps or rest not in ("Không có gaps.", "Gapsなし。")):
        raise ValueError("Unsupported gap section tail")
    gap_ids = {identifier(row[0]) for row in gaps}
    if len(gap_ids) != len(gaps):
        raise ValueError("Duplicate gap IDs")
    for row in gaps:
        if row[1] not in ("business", "fact", "engineering", "preparation", "proof"):
            raise ValueError("Unknown gap kind")
    conventions, context_tail = fields(sections["test-cases"][0], ["Revision", "Feature", "Conventions"])
    contexts = {}
    if context_tail:
        chunks = re.split(r"^### Context: (.+)$", context_tail, flags=re.M)
        if chunks[0].strip() or len(chunks) % 2 != 1:
            raise ValueError("Unsupported context syntax")
        for context_id, block in zip(chunks[1::2], chunks[2::2]):
            identifier(context_id)
            data, rest = fields(block, CASE_LABELS[language], exact=False)
            if context_id in contexts or rest or not data or any(value.startswith("@") for value in data.values()):
                raise ValueError("Duplicate/recursive/invalid context")
            contexts[context_id] = data
    fixtures_meta, rest = fields(sections["test-data"][0], ["Revision", "Conventions"])
    if rest:
        raise ValueError("Unsupported data convention syntax")
    fixtures = []
    tail = sections["test-data"][1]
    if tail not in ("Không có fixture dùng chung.", "共通fixtureなし。"):
        chunks = re.split(r"^### Fixture: (.+)$", tail, flags=re.M)
        if chunks[0].strip() or len(chunks) == 1:
            raise ValueError("Missing fixture blocks or explicit empty-set statement")
        for fixture_id, block in zip(chunks[1::2], chunks[2::2]):
            data, rest = fields(block, FIXTURE_LABELS[language])
            if rest:
                raise ValueError("Unprojected fixture prose")
            fixtures.append({"id": identifier(fixture_id), "values": list(data.values())})
    fixture_ids = {item["id"] for item in fixtures}
    if len(fixture_ids) != len(fixtures):
        raise ValueError("Duplicate fixture IDs")
    revisions = [scope[0]["Revision"], conventions["Revision"], fixtures_meta["Revision"]]
    if len(set(revisions)) != 1:
        raise ValueError("Markdown design revisions differ")
    cases = []
    flows = re.split(r"^### Flow: (.+)$", sections["test-cases"][1], flags=re.M)
    if flows[0].strip() or len(flows) == 1:
        raise ValueError("Missing business-flow groups")
    for group, block in zip(flows[1::2], flows[2::2]):
        chunks = re.split(r"^#### (\S+) — (.+)$", block, flags=re.M)
        if chunks[0].strip() or len(chunks) == 1:
            raise ValueError("Missing case block/title")
        for offset in range(1, len(chunks), 3):
            case_id, title, body = chunks[offset:offset + 3]
            data, rest = fields(body, CASE_LABELS[language])
            for label, value in data.items():
                if value.startswith("@CTX"):
                    try:
                        data[label] = contexts[value[1:]][label]
                    except KeyError as error:
                        raise ValueError("Unresolved same-field shared context") from error
            data = dict(zip(CASE_KEYS, data.values()))
            if data["priority"] not in ("High", "Medium", "Low") or data["basis"] not in ("Confirmed", "Proposed", "Awaiting decision"):
                raise ValueError("Unsupported priority/expected basis")
            if data["readiness"] not in ("Ready", "Draft", "Blocked"):
                raise ValueError("Unsupported readiness")
            if data["readiness"] == "Ready" and (data["gap"] != "none" or data["basis"] == "Awaiting decision"):
                raise ValueError("Ready case has unknown oracle/gap")
            if data["readiness"] != "Ready":
                references(data["gap"], gap_ids)
            if not data["fixture"].startswith("local: "):
                fixture_refs = data["fixture"].split(", ")
                if any(value not in fixture_ids for value in fixture_refs) or len(set(fixture_refs)) != len(fixture_refs):
                    raise ValueError("Unknown/compound/duplicate fixture reference")
            steps, variants = [], []
            if rest.startswith("##### Steps\n"):
                steps, rest = read_table(rest[len("##### Steps\n"):], ["Step", "Action", "Expected", "Preservation"])
                if not steps or [row[0] for row in steps] != [str(i) for i in range(1, len(steps) + 1)] or data["action"] != "@Steps" or data["expected"] != "@Steps":
                    raise ValueError("Expanded steps must be ordered and uniquely replace inline action/expected")
            elif data["action"] == "@Steps" or data["expected"] == "@Steps":
                raise ValueError("Missing Steps table")
            if rest.startswith("##### Variants\n"):
                variants, rest = read_table(rest[len("##### Variants\n"):], ["Variant", "Inputs", "Expected"])
                if not variants or len({identifier(row[0]) for row in variants}) != len(variants):
                    raise ValueError("Empty/duplicate/compound variants")
            if rest:
                raise ValueError("Unsupported/unprojected case prose or table")
            cases.append({"id": identifier(case_id), "title": title, "group": group, **data, "steps": steps,
                          "variants": variants or [["base", data["fixture"], data["expected"]]]})
    case_ids = {case["id"] for case in cases}
    if len(case_ids) != len(cases):
        raise ValueError("Duplicate Case ID")
    variant_ids = {f'{case["id"]}:{row[0]}' for case in cases for row in case["variants"]}
    for row in coverage:
        references(row[2], case_ids | variant_ids | gap_ids)
    covered = {target.split(":")[0] for row in coverage for target in row[2].split(", ")}
    if case_ids - covered:
        raise ValueError("Case has no reverse coverage mapping")
    return {"revision": revisions[0], "texts": texts, "scope": scope, "coverage": coverage, "gaps": gaps,
            "conventions": conventions, "data_conventions": fixtures_meta["Conventions"], "fixtures": fixtures, "cases": cases}


def load_template(path: Path, *, raw: bytes | None = None):
    from openpyxl import load_workbook
    captured = path.read_bytes() if raw is None else raw
    workbook = load_workbook(io.BytesIO(captured))
    properties = {prop.name: prop.value for prop in workbook.custom_doc_props}
    if properties.get("TemplateFamily") != "test-case-report" or properties.get("TemplateVersion") != VERSION or properties.get("Language") != "vi":
        raise ValueError("Unsupported workbook template family/version/language")
    if workbook.sheetnames != ["Run", "Cases", "Data"]:
        raise ValueError("Workbook template sheets must be Run/Cases/Data")
    for sheet, cells in {"Run": ("B3", "B4", "B5", "B6", "B7", "B8", "B9", "B11", "B12"),
                         "Cases": ("B3", "B4", "B5"), "Data": ("B3", "B4")}.items():
        if any(workbook[sheet][cell].value is not None for cell in cells):
            raise ValueError("Workbook template contains prefilled source/run metadata")
    for sheet, row, headers in [("Run", RUN_START - 1, RUN_HEADERS), ("Cases", CASE_START - 1, CASE_HEADERS), ("Data", DATA_START - 1, DATA_HEADERS)]:
        if [workbook[sheet].cell(row, index).value for index in range(1, len(headers) + 1)] != headers:
            raise ValueError(f"Workbook template header mismatch: {sheet}")
        if any(cell.value is not None for cells in workbook[sheet].iter_rows(min_row=row + 1) for cell in cells):
            raise ValueError("Workbook template contains execution/design rows")
    return workbook


def write_text(sheet, row: int, column: int, value):
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell_value(value)
        cell.data_type = "s"  # untrusted Markdown must never become Excel formula
    return cell


def export(source: Path, output: Path, *, language: str = "vi", template: Path | None = None) -> dict:
    if language != "vi":
        raise ValueError("JA workbook capability unavailable; request VI workbook or Markdown-only subset")
    if output.exists():
        raise FileExistsError("Existing workbook preserved; request a new revision/path")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must be .xlsx")
    # Parse, project and fingerprint the same immutable bytes; never reread for identity.
    captured = {f"{family}.vi.md": (source / f"{family}.vi.md").read_bytes() for family in HEADINGS["vi"]}
    for family in HEADINGS["ja"]:
        path = source / f"{family}.ja.md"
        if path.is_file():
            captured[path.name] = path.read_bytes()
    fingerprints = {name: hashlib.sha256(raw).hexdigest() for name, raw in captured.items()}
    design = parse_sources(source, captured=captured)
    for lang in ("vi", "ja"):
        existing = [name for name in captured if name.endswith(f".{lang}.md")]
        if lang == "ja" and existing:
            if len(existing) != 3:
                raise ValueError("Partial JA trio: export supplied VI-only subset separately or complete counterparts")
            companion = parse_sources(source, "ja", captured=captured)
            if companion["revision"] != design["revision"] or [(case["id"], [row[0] for row in case["variants"]], case["priority"], case["basis"], case["readiness"], case["gap"]) for case in companion["cases"]] != [(case["id"], [row[0] for row in case["variants"]], case["priority"], case["basis"], case["readiness"], case["gap"]) for case in design["cases"]] or [item["id"] for item in companion["fixtures"]] != [item["id"] for item in design["fixtures"]]:
                raise ValueError("JA/VI identity/state parity mismatch; meaning still requires review")
    template = template or Path(__file__).resolve().parents[1] / "assets/test-case-report-template.xlsx"
    template_bytes = template.read_bytes()
    template_fingerprint = hashlib.sha256(template_bytes).hexdigest()
    workbook = load_template(template, raw=template_bytes)
    run, cases_sheet, data_sheet = [workbook[name] for name in workbook.sheetnames]
    write_text(run, 3, 2, design["revision"])
    write_text(run, 4, 2, design["scope"][0]["Nguồn"])
    write_text(run, 11, 2, "\n".join(f"{name}: sha256:{digest}" for name, digest in fingerprints.items()))
    write_text(run, 12, 2, "sha256:" + template_fingerprint)
    write_text(cases_sheet, 3, 2, design["revision"])
    write_text(data_sheet, 3, 2, design["revision"])
    write_text(cases_sheet, 4, 2, design["conventions"]["Feature"])
    write_text(cases_sheet, 5, 2, design["conventions"]["Conventions"])
    write_text(data_sheet, 4, 2, design["data_conventions"])
    result_rows = []
    for index, case in enumerate(design["cases"], CASE_START):
        values = [case["id"], case["group"], case["title"], *[case[key] for key in CASE_KEYS],
                  "\n".join(" | ".join(row) for row in case["steps"]), "\n".join(" | ".join(row) for row in case["variants"])]
        for column, value in enumerate(values, 1):
            write_text(cases_sheet, index, column, value)
        cases_sheet.cell(index, 1).hyperlink = f"#'Run'!B{RUN_START + len(result_rows)}"
        for variant in case["variants"]:
            result_rows.append([None, case["id"], variant[0], case["group"], case["priority"], case["basis"], case["readiness"], "Yes", "NOT RUN", None, None, None, None, None, None, None, None, None, f"Cases!A{index}"])
    for index, fixture in enumerate(design["fixtures"], DATA_START):
        for column, value in enumerate([fixture["id"], *fixture["values"]], 1):
            write_text(data_sheet, index, column, value)
    # Project scope/coverage/gaps without a second persisted source or silent prose loss.
    row = DATA_START + len(design["fixtures"]) + 2
    for family, text in design["texts"].items():
        if family == "scope-and-approach":
            write_text(data_sheet, row, 1, "Scope / research / coverage / gaps")
            # One paragraph/table row per record avoids Excel's cell length ceiling.
            for line in text.splitlines():
                if line.strip():
                    row += 1
                    write_text(data_sheet, row, 1, family)
                    write_text(data_sheet, row, 2, line)
    data_last = max(row, DATA_START)
    last = RUN_START + len(result_rows) - 1
    ranges = {col: f"${col}${RUN_START}:${col}${last}" for col in ("A", "B", "F", "G", "H", "I", "Q", "R")}
    for row, values in enumerate(result_rows, RUN_START):
        for column, value in enumerate(values, 1):
            if value is not None:
                write_text(run, row, column, value)
        run.cell(row, 1, '=$B$5')
        run.cell(row, 19).hyperlink = "#'" + values[18].replace("!", "'!")
        run.cell(row, 17, f'=IF(OR(AND(H{row}<>"Yes",H{row}<>"No"),AND(I{row}<>"NOT RUN",I{row}<>"PASS",I{row}<>"FAIL",I{row}<>"BLOCKED",I{row}<>"SKIPPED")),"Invalid scope/status",IF(AND(I{row}="SKIPPED",P{row}=""),"Missing reason",IF(AND(I{row}="BLOCKED",P{row}=""),"Missing blocker",IF(OR(I{row}="PASS",I{row}="FAIL"),IF(OR(A{row}="",J{row}="",K{row}="",M{row}="",N{row}="",O{row}=""),"Missing execution evidence",IF(OR(F{row}<>"Confirmed",G{row}<>"Ready"),"Observation only","OK")),"OK"))))')
        # All mandatory variants, including rows marked out of this run scope, prevent full case PASS.
        selectors = f'{ranges["A"]},A{row},{ranges["B"]},B{row}'
        run.cell(row, 18, f'=IF(COUNTIFS({selectors},{ranges["H"]},"Yes",{ranges["F"]},"Confirmed",{ranges["I"]},"FAIL",{ranges["Q"]},"OK")>0,"FAIL",IF(COUNTIFS({selectors},{ranges["H"]},"Yes",{ranges["F"]},"Confirmed",{ranges["G"]},"Ready",{ranges["I"]},"PASS",{ranges["Q"]},"OK")=COUNTIFS({selectors}),"PASS",IF(COUNTIFS({selectors},{ranges["I"]},"BLOCKED")>0,"BLOCKED","NOT RUN")))')
    run.cell(14, 2, f'=COUNTIF({ranges["H"]},"Yes")')
    run.cell(15, 2, len(design["cases"]))
    for label_row, status in [(16, "PASS"), (17, "FAIL"), (18, "BLOCKED"), (19, "NOT RUN"), (20, "SKIPPED")]:
        run.cell(label_row, 2, f'=COUNTIFS({ranges["H"]},"Yes",{ranges["F"]},"Confirmed",{ranges["Q"]},"OK",{ranges["I"]},"{status}")' if status in ("PASS", "FAIL") else f'=COUNTIFS({ranges["H"]},"Yes",{ranges["I"]},"{status}")')
    run.cell(21, 2, f'=IF(COUNTIFS({ranges["H"]},"Yes",{ranges["F"]},"Confirmed")=0,"No data",(B16+B17)/COUNTIFS({ranges["H"]},"Yes",{ranges["F"]},"Confirmed"))')
    run.cell(22, 2, '=IF(B16+B17=0,"No data",B16/(B16+B17))')
    for row in (21, 22):
        run.cell(row, 2).number_format = "0.0%"
    style_projection(workbook, last, CASE_START + len(design["cases"]) - 1, data_last)
    from openpyxl.workbook.properties import CalcProperties
    workbook.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    # Reparse bytes before claiming structural export. x mode also refuses a race collision.
    from openpyxl import load_workbook
    reopened = load_workbook(io.BytesIO(buffer.getvalue()))
    if [reopened["Run"].cell(i, 9).value for i in range(RUN_START, last + 1)] != ["NOT RUN"] * len(result_rows):
        raise ValueError("Saved workbook parity failed")
    if any(hashlib.sha256((source / name).read_bytes()).hexdigest() != digest for name, digest in fingerprints.items()):
        raise ValueError("Source changed during export; no workbook written")
    current_names = {f"{family}.{lang}.md" for lang in HEADINGS for family in HEADINGS[lang]
                     if (source / f"{family}.{lang}.md").is_file()}
    if current_names != set(captured) or hashlib.sha256(template.read_bytes()).hexdigest() != template_fingerprint:
        raise ValueError("Source inventory/template changed during export; no workbook written")
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("xb") as handle:
        handle.write(buffer.getvalue())
    return {"cases": len(design["cases"]), "variants": len(result_rows), "fixtures": len(design["fixtures"]), "revision": design["revision"]}


def wrapped_line_count(value: str, column_width: float) -> int:
    """Conservative Arial-10 estimate, retaining explicit lines and wide JA glyphs."""
    capacity = max(1, column_width - 3)
    def units(word):
        return sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1 for char in word)
    count = 0
    for paragraph in value.split("\n"):
        line_units = 0
        count += 1
        for word in paragraph.split(" "):
            width = units(word)
            if width > capacity:
                if line_units:
                    count += 1
                lines = math.ceil(width / capacity)
                count += lines - 1
                line_units = width % capacity or capacity
            elif line_units and line_units + 1 + width > capacity:
                count += 1
                line_units = width
            else:
                line_units += (1 if line_units else 0) + width
    return count


def projection_row_height(cells, sheet, *, formula_display: dict[int, str] | None = None) -> float:
    """Never size from formula code or silently clip above Excel's row-height limit."""
    formula_display = formula_display or {}
    lines = []
    for cell in cells:
        value = formula_display.get(cell.column, "") if cell.data_type == "f" else str(cell.value or "")
        lines.append(wrapped_line_count(value, sheet.column_dimensions[cell.column_letter].width))
    height = max(36, 15 * max(lines, default=1) + 9)
    if height > 409:
        # ponytail: Excel caps row height; an approved multirow template is the upgrade path.
        raise ValueError(f"Readable content exceeds Excel row height at {sheet.title}!{cells[0].row}; "
                         "multirow workbook layout requires an approved template/observer mapping; no output written")
    return height


def style_projection(workbook, run_last, case_last, data_last):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter
    # Steps/variants are reader-facing prose. Width plus actual wrapped lines determines height.
    for column in ("T", "U"):
        workbook["Cases"].column_dimensions[column].width = 50
    for name, start, last, width in [("Run", RUN_START, run_last, len(RUN_HEADERS)), ("Cases", CASE_START, case_last, len(CASE_HEADERS)), ("Data", DATA_START, data_last, len(DATA_HEADERS))]:
        sheet = workbook[name]
        for row in sheet.iter_rows(min_row=start, max_row=last, max_col=width):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[row[0].row].height = projection_row_height(
                row, sheet, formula_display={1: "Run ID", 17: "Missing execution evidence", 18: "NOT RUN"} if name == "Run" else None)
        sheet.auto_filter.ref = f"A{start - 1}:{get_column_letter(width)}{last}"
        sheet.print_title_rows = f"{start - 1}:{start - 1}"
        sheet.sheet_view.showGridLines = False
    run = workbook["Run"]
    # Keep actual run setup/results visible; provenance is expandable, never discarded.
    run.row_dimensions.group(10, 12, outline_level=1, hidden=True)
    run.row_dimensions[13].collapsed = True
    run.row_dimensions[13].height = 30
    run.row_dimensions[23].height = 6
    run.sheet_properties.outlinePr.summaryBelow = True
    write_text(run, 13, 1, "Chi tiết (+)")
    write_text(run, 13, 2, "Mở nhóm để xem quy tắc và source/template fingerprints.")
    for column, values in [("H", '"Yes,No"'), ("I", '"NOT RUN,PASS,FAIL,BLOCKED,SKIPPED"')]:
        validation = DataValidation(type="list", formula1=values, allow_blank=False)
        validation.errorTitle = "Giá trị không hợp lệ"
        validation.error = "Chọn trạng thái trong danh sách."
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        run.add_data_validation(validation)
        validation.add(f"{column}{RUN_START}:{column}{run_last}")
    for row in range(RUN_START, run_last + 1):
        for col in (8, 9, 10, 11, 12, 13, 14, 15, 16):
            run.cell(row, col).fill = PatternFill("solid", fgColor="FFF2CC")
        run.cell(row, 14).number_format = "yyyy-mm-dd hh:mm:ss"
    run.row_dimensions[11].height = min(409, 13 * sum(max(1, (len(line) + 43) // 44) for line in str(run["B11"].value).splitlines()) + 12)
    run.conditional_formatting.add(f"Q{RUN_START}:Q{run_last}", FormulaRule(formula=[f'Q{RUN_START}<>"OK"'], fill=PatternFill("solid", fgColor="FCE4D6")))
    run.conditional_formatting.add(f"I{RUN_START}:I{run_last}", FormulaRule(formula=[f'I{RUN_START}="FAIL"'], fill=PatternFill("solid", fgColor="FCE4D6")))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--language", default="vi")
    parser.add_argument("--template", type=Path)
    args = parser.parse_args()
    try:
        result = export(args.source_dir, args.output, language=args.language, template=args.template)
    except (ValueError, OSError, ImportError) as error:
        parser.exit(1, f"Export blocked: {error}\n")
    print(f"Exported {result} to {args.output}. Structural projection only; render/recalculation not performed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
