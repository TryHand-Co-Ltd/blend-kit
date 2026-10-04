"""Two-sheet TC blocks. Source and result cells have separate, explicit roles."""
from collections import Counter
from pathlib import Path
import hashlib
import io
import re

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

import report_model as model


def layout(data):
    """Pure source projection: no hidden dataset and no recorded results."""
    from render_report import _chunks
    vi = data.language == 'vi'
    labels = ('Điều kiện kiểm thử', 'Thao tác', 'Kết quả mong đợi', 'Kết quả thực tế') if vi else ('テスト条件', '操作', '期待結果', '実際の結果')
    cards, inputs, cursor = [], {}, 5
    for case in data.cases:
        members = [record for record in data.rows if record.case_id == case['id']]
        start = cursor
        rows = []
        def add(value='', kind='text', right=None, height=None, label=None):
            nonlocal cursor
            rows.append((cursor, value, kind, right, height, label))
            cursor += 1
            return cursor - 1
        def text(value):
            for piece in _chunks(value, 112, 13):
                add(piece)
        add(case['id'] + ' — ' + case['title'], 'title')
        header_status = add('Trạng thái' if vi else '状態', 'case_status')
        add(case['function'], 'screen', model.screen_path(case['screen_relative_path'], data.language))
        add(labels[0], 'section')
        for label, value in case['preparation_items']:
            add(value, 'key_value', label=label)
        add(labels[1], 'section')
        for item in case['action_items']:
            add(item['value'], 'step', label=item['label'])
        add(labels[2], 'section')
        for label, value in case['expected_items']:
            add(value, 'expected_item', label=label)
        add(labels[3], 'section')
        for record in members:
            if len(members) > 1:
                add(record.identity, 'variant')
                add(model.display_sentence(case['variants'][record.variant][0]), 'key_value', label='Dữ liệu biến thể' if vi else 'バリエーションデータ')
                extra = case['variants'][record.variant][1]
                if extra and extra != case['expected_source'] and extra != '@Steps':
                    add(model.display_sentence(extra), 'expected_item', label='Kết quả riêng' if vi else '個別の期待結果')
                status = add('Trạng thái' if vi else '状態', 'status')
            else:
                # The single editable status remains directly below the title.
                status = header_status
                if not (record.variant=='base' and case['variants'][record.variant][0]==case['fixture_source']):
                    add(model.display_sentence(case['variants'][record.variant][0]), 'key_value', label='Dữ liệu biến thể' if vi else 'バリエーションデータ')
                extra = case['variants'][record.variant][1]
                if extra and extra != case['expected_source'] and extra != '@Steps':
                    add(model.display_sentence(extra), 'expected_item', label='Kết quả riêng' if vi else '個別の期待結果')
            add('Kết quả thực tế' if vi else '実際の結果', 'label')
            actual = add('', 'input', height=64)
            add('Ảnh minh chứng' if vi else '証拠画像', 'label')
            picture = add('', 'picture', height=90)
            inputs[record.identity] = {'status': f'B{status}', 'actual': f'A{actual}', 'picture': picture}
            add('', height=12)
        add('', height=20)
        cards.append({'id': case['id'], 'start': start, 'status_row': header_status, 'rows': rows,
                      'members': members})
    return cards, inputs


def formulas(data, cards=None, inputs=None):
    """Case-level recorded counts feed Overview; never count prose as results."""
    if cards is None:
        cards, inputs = layout(data)
    sheet = "'" + model.REPORT_LAYOUTS[data.language]['sheets'][1] + "'!"
    states = model.REPORT_LAYOUTS[data.language]['statuses']
    output, case_statuses = {}, {}
    def counted(token, members, judged=False):
        if judged and not all(record.eligible for record in members):
            return '0'
        refs = [inputs[record.identity] for record in members]
        if len(refs) == 1:
            status = sheet + refs[0]['status']
            actual = sheet + refs[0]['actual']
        else:
            first, last = refs[0], refs[-1]
            status = sheet + first['status'] + ':' + last['status']
            # Multi-variant blocks use the same offsets after every status.
            start = int(first['status'][1:])
            end = int(last['status'][1:])
            actual = sheet + f'A{start+2}:A{end+2}'
        terms = [f'--EXACT({status},"{states[token]}")']
        if judged:
            terms += [f'--(LEN({model._excel_without_whitespace(actual)})>0)']
        return 'SUMPRODUCT(' + ','.join(terms) + ')'
    for card in cards:
        members = card['members']
        passed, failed = counted('PASS', members, True), counted('FAIL', members, True)
        skipped, blocked, unrun = (counted(token, members) for token in ('SKIPPED', 'BLOCKED', 'NOT RUN'))
        status = (f'=IF({failed}>0,"{states["FAIL"]}",IF({passed}={len(members)},"{states["PASS"]}",'
                  f'IF({skipped}={len(members)},"{states["SKIPPED"]}",IF({blocked}>0,"{states["BLOCKED"]}",'
                  f'IF({unrun}={len(members)},"{states["NOT RUN"]}","{model.REPORT_LAYOUTS[data.language]["case_pending"]}")))))')
        if len(members) > 1:
            output[(1, f'B{card["status_row"]}')] = status
            case_statuses[card['id']] = sheet + f'B{card["status_row"]}'
        else:
            # Overview derives validity; the header status is the actual input.
            case_statuses[card['id']] = re.sub(r'(?<![!A-Z0-9_])([BCD]\d+)', lambda match: sheet + match[1], status)
    for index, card in enumerate(cards, model.CASE_START_ROW):
        value = case_statuses[card['id']]
        output[(0, f'C{index}')] = value if value.startswith('=') else '=' + value
    output[(0, 'B16')] = f'={len(cards)}'
    output[(0, 'B17')] = f'={len(data.rows)}'
    last = cards[-1]['rows'][-1][0] if cards else 5
    status_range = sheet + f'B5:B{last}'
    status_rows = [int(inputs[record.identity]['status'][1:]) for record in data.rows]
    def row_mask(rows):
        return '--ISNUMBER(MATCH(ROW('+status_range+'),{'+','.join(map(str,rows))+'},0))' if rows else '0'
    def raw_count(token):
        terms = [f'--EXACT({status_range},"{states[token]}")', row_mask(status_rows)]
        return 'SUMPRODUCT(' + ','.join(terms) + ')'
    case_status_range = f'C{model.CASE_START_ROW}:C{model.CASE_START_ROW+len(cards)-1}'
    output[(0, 'B18')] = f'=COUNTIF({case_status_range},"{states["PASS"]}")'
    output[(0, 'B19')] = f'=COUNTIF({case_status_range},"{states["FAIL"]}")'
    for token in ('BLOCKED', 'SKIPPED', 'NOT RUN'):
        output[(0, f'B{model.METRIC_ROWS[token]}')] = '=' + raw_count(token)
    output[(0, 'B23')] = '=B18+B19'
    no_data = model.REPORT_LAYOUTS[data.language]['no_data']
    output[(0, 'B24')] = f'=IF(B23=0,"{no_data}",B18/B23)'
    output[(0, 'B25')] = f'=IF(B16=0,"{no_data}",B23/B16)'
    output[(0, 'B26')] = '=B16-B18'
    # Count only the exact designated status cells, including invalid pasted input.
    output[(0, 'B27')] = '=B17-SUM(B20:B22,' + raw_count('PASS') + ',' + raw_count('FAIL') + ')'
    if any(len(value) > 8192 for value in output.values()):
        raise ValueError('Report exceeds Excel formula-size limit; preserve source and narrow the explicitly requested run')
    return output


def build(data):
    from render_report import _text, _height, _jump, _finish
    locale = model.REPORT_LAYOUTS[data.language]
    path = Path(__file__).resolve().parents[1] / 'assets' / locale['template']
    book = load_workbook(path)
    if book.sheetnames != list(locale['sheets']) or {p.name:p.value for p in book.custom_doc_props} != {'TemplateFamily':model.FAMILY,'TemplateVersion':model.VERSION,'Language':data.language}:
        raise ValueError('Block template identity or sheets mismatch')
    summary, tests = book.worksheets
    # The source template contains a visible illustrative block for humans.
    # It is presentation-only and must never leak into a generated report.
    for merged in list(tests.merged_cells.ranges):
        if merged.min_row >= locale['start_row']:
            tests.unmerge_cells(str(merged))
    if tests.max_row >= locale['start_row']:
        tests.delete_rows(locale['start_row'], tests.max_row - locale['start_row'] + 1)
    tests.data_validations.dataValidation = []
    tests.conditional_formatting._cf_rules.clear()
    _text(tests['A1'], locale['block_title'])
    tests['A1'].font = Font(name=locale['font'], size=16, bold=True, color='172033')
    _text(tests['A2'], locale['instructions'])
    tests['A2'].font = Font(name=locale['font'], size=10, color='475569')
    tests.merge_cells('A2:E2')
    tests.row_dimensions[2].height = 42
    # Keep the validation cell compact. The remaining columns preserve the
    # total content width used by long conditions/actions and screen details.
    for column, width in (('A', 23), ('B', 24), ('C', 26), ('D', 14), ('E', 45)):
        tests.column_dimensions[column].width = width
    cards, inputs = layout(data)
    thin = Side(style='thin', color='D9E2EC')
    for key, row in model.SUMMARY_FIELDS.items():
        value = '' if key in model.INPUT_FIELDS else data.summary.get(key, '')
        _text(summary.cell(row,2), value)
        _height(summary,row,(summary.cell(row,1).value or '',value),(26,105))
        if key in model.INPUT_FIELDS:
            summary.cell(row,2).fill=PatternFill('solid',fgColor='F8FAFC')
            summary.cell(row,2).protection=Protection(locked=False)
    for index,card in enumerate(cards,model.CASE_START_ROW):
        case = next(case for case in data.cases if case['id']==card['id'])
        _text(summary.cell(index,1),card['id'])
        _text(summary.cell(index,2),model.display_markup(case['title']))
        _jump(summary.cell(index,2),tests.title,card['start'])
        _height(summary,index,(card['id'],case['title']),(26,83))
        for row,value,kind,right,height,label in card['rows']:
            _text(tests.cell(row,1),value)
            if kind in ('status','case_status'):
                # Status always lives in the compact B cell. It must stay
                # unmerged so Excel/Google Sheets anchor the popup beside it.
                pass
            elif kind=='screen':
                _text(tests.cell(row,1),'Màn hình / chức năng' if data.language=='vi' else '画面・機能')
                tests.cell(row,1).font=Font(name=locale['font'],size=11,bold=True,color='172033')
                _text(tests.cell(row,2),value)
                tests.cell(row,2).font=Font(name=locale['font'],size=11,bold=True,color='172033')
                tests.merge_cells(f'B{row}:C{row}')
                _text(tests.cell(row,4),'Đường dẫn' if data.language=='vi' else '相対URL')
                tests.cell(row,4).font=Font(name=locale['font'],size=11,bold=True,color='172033')
                _text(tests.cell(row,5),right)
            elif kind in ('key_value','expected_item'):
                _text(tests.cell(row,1),label or '•')
                tests.cell(row,1).font=Font(name=locale['font'],size=11,bold=True,color='172033')
                tests.merge_cells(f'B{row}:E{row}')
                _text(tests.cell(row,2),value)
            elif kind=='step':
                _text(tests.cell(row,1),label)
                tests.cell(row,1).font=Font(name=locale['font'],size=11,bold=True,color='172033')
                tests.merge_cells(f'B{row}:E{row}')
                _text(tests.cell(row,2),value)
            else:
                tests.merge_cells(f'A{row}:E{row}')
            if kind in ('title','section','label','variant'):
                tests.cell(row,1).font=Font(name=locale['font'],size=14 if kind=='title' else 11,bold=True,color='172033')
                tests.cell(row,1).alignment=Alignment(vertical='center',wrap_text=True)
            if kind in ('status','case_status'):
                tests.cell(row,1).font=Font(name=locale['font'],size=11,bold=True,color='172033')
                tests.cell(row,1).alignment=Alignment(vertical='center',wrap_text=True)
            if kind == 'title':
                for column in range(1,6):
                    tests.cell(row,column).fill=PatternFill('solid',fgColor='334155')
                tests.cell(row,1).font=Font(name=locale['font'],size=14,bold=True,color='FFFFFF')
                tests.row_dimensions[row].height=max(tests.row_dimensions[row].height or 0,34)
            if kind=='section':
                for column in range(1,6):
                    tests.cell(row,column).fill=PatternFill('solid',fgColor='E8EEF5')
                tests.row_dimensions[row].height=max(tests.row_dimensions[row].height or 0,28)
            if kind == 'variant':
                for column in range(1,6):
                    tests.cell(row,column).fill=PatternFill('solid',fgColor='F8FAFC')
            if kind in ('input','picture'):
                tests.cell(row,1).fill=PatternFill('solid',fgColor='F8FAFC')
                tests.cell(row,1).protection=Protection(locked=False)
                tests.cell(row,1).number_format='@'
                for column in range(1,6):
                    tests.cell(row,column).border=Border(top=thin,bottom=thin,left=thin if column==1 else Side(),right=thin if column==5 else Side())
            if kind in ('screen','key_value','expected_item','step','status','case_status'):
                for column in range(1,6):
                    tests.cell(row,column).border=Border(bottom=thin)
            for column in range(1,6):
                cell = tests.cell(row,column)
                if cell.value is not None:
                    cell.alignment=Alignment(horizontal=cell.alignment.horizontal,vertical='center',wrap_text=True)
            if kind == 'screen':
                _height(tests,row,(value,right or ''),(38,38))
            elif kind in ('key_value','expected_item','step'):
                _height(tests,row,(value,),(104,))
            else:
                _height(tests,row,(value,),(112,))
            if height:
                tests.row_dimensions[row].height=height
    validation=DataValidation(type='list',formula1='"'+','.join(locale['statuses'].values())+'"',allow_blank=False)
    validation.showErrorMessage=True
    tests.add_data_validation(validation)
    for fields in inputs.values():
        _text(tests[fields['status']],locale['statuses']['NOT RUN'])
        tests[fields['status']].fill=PatternFill('solid',fgColor='F8FAFC')
        tests[fields['status']].protection=Protection(locked=False)
        validation.add(fields['status'])
    for (index,address),formula in formulas(data,cards,inputs).items():
        book.worksheets[index][address]=formula
        book.worksheets[index][address].font=Font(name=locale['font'],size=11,color='172033')
        book.worksheets[index][address].alignment=Alignment(vertical='center',wrap_text=True)
    for sheet in (summary,tests):
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.alignment=Alignment(horizontal='left' if sheet is tests else cell.alignment.horizontal,vertical='center',wrap_text=True)
    _finish(summary,'C',model.CASE_START_ROW+len(cards)-1)
    _finish(tests,'E',cards[-1]['rows'][-1][0] if cards else 4)
    tests.freeze_panes='A5'
    tests.page_setup.orientation='portrait'
    tests.page_setup.paperSize=tests.PAPERSIZE_A3
    book.active=0
    return book


def check(report,source,language='vi',phase='in-progress',*,closure_confirmation=None,evidence_access=(),screenshots=()):
    from check_report import _package,_formula_reference_key
    if phase not in ('in-progress','complete'):
        raise ValueError('Unknown check phase')
    data,captures=model.prepare_report(Path(source),language)
    asset=Path(__file__).resolve().parents[1]/'assets'/model.REPORT_LAYOUTS[language]['template']
    captures[asset]=model._capture(asset)
    report=Path(report)
    captures[report]=model._capture(report)
    raw=captures[report][0]
    book=load_workbook(io.BytesIO(raw))
    pristine=build(data)
    locale=model.REPORT_LAYOUTS[language]
    if book.sheetnames!=pristine.sheetnames or {p.name:p.value for p in book.custom_doc_props}!={p.name:p.value for p in pristine.custom_doc_props}:
        raise ValueError('Unsupported block report schema')
    cards,inputs=layout(data)
    gaps,layout_notes,observations=[],[],[]
    editable={(locale['sheets'][0],f'B{model.SUMMARY_FIELDS[key]}') for key in model.INPUT_FIELDS}
    editable|={(locale['sheets'][1],fields[key]) for fields in inputs.values() for key in ('status','actual')}
    summary,tests=book.worksheets
    for key in model.INPUT_FIELDS:
        value=model._read_text(summary[f'B{model.SUMMARY_FIELDS[key]}'],key,optional=True)
        if value:
            model.public_text(value,key)
        else:
            gaps.append(f'{summary.title}: missing {key}')
    if phase=='complete' and summary[f'B{model.SUMMARY_FIELDS["period"]}'].value:
        closed=model._records(closure_confirmation,model.ClosureConfirmation)
        if model.timestamp(summary[f'B{model.SUMMARY_FIELDS["period"]}'].value,'execution time')>model.timestamp(closed.closed_at,'closed_at'):
            raise ValueError('Execution is after closure')
    states={value:key for key,value in locale['statuses'].items()}
    for record in data.rows:
        fields=inputs[record.identity]
        actual,status=[model._read_text(tests[fields[key]],record.identity+' '+key,optional=key!='status') for key in ('actual','status')]
        if model.markdown_leak(actual):
            raise ValueError('Actual result must use spreadsheet text such as [identifier], not Markdown syntax: '+record.identity)
        if status not in states:
            raise ValueError('Invalid pasted status: '+record.identity)
        token=states[status]
        for key,value in (('actual',actual),):
            if value.strip():
                model.public_text(value,record.identity+' '+key)
            if tests[fields[key]].hyperlink and tests[fields[key]].hyperlink.target:
                model.shared_url(tests[fields[key]].hyperlink.target,record.identity+' '+key)
            height=tests.row_dimensions[tests[fields[key]].row].height or 15
            if height<(value.count('\n')+1)*11:
                gaps.append(record.identity+': input row too short for explicit lines')
            elif model.display_lines(value,112)*17+10>height:
                layout_notes.append(record.identity+': inspect input wrapping in native engine')
        valid=token in ('PASS','FAIL') and record.eligible and bool(actual.strip())
        if token in ('PASS','FAIL') and not valid:
            gaps.append(record.identity+': confirmed/ready basis and actual result required')
        if token not in ('PASS','NOT RUN') and not actual.strip():
            gaps.append(record.identity+': actual result must include reason and next action')
        if token=='NOT RUN':
            gaps.append(record.identity+': not executed')
        observations.append((record,token,valid))
    for case in data.cases:
        if case['screen_relative_path']=='unknown':
            gaps.append(case['id']+': screen relative URL remains unverified')
    image_bindings=[]
    allowed_images={fields['picture']:(identity,fields) for identity,fields in inputs.items()}
    for sheet in book:
        if sheet.sheet_state!='visible' or any(d.hidden for d in sheet.row_dimensions.values()) or any(d.hidden for d in sheet.column_dimensions.values()):
            raise ValueError('Hidden report content unsupported')
        for image in sheet._images:
            anchor=image.anchor._from
            if sheet.title!=tests.title or anchor.col!=0 or anchor.row+1 not in allowed_images:
                raise ValueError('Evidence image must use its reserved TC/variant slot')
            identity,fields=allowed_images[anchor.row+1]
            case_id,variant=identity.split(' / ')
            image_bindings.append((case_id,variant,hashlib.sha256(image._data()).hexdigest()))
        original=pristine[sheet.title]
        if set(str(r) for r in sheet.merged_cells.ranges)!=set(str(r) for r in original.merged_cells.ranges):
            raise ValueError('Changed block merges')
        for row in sheet.iter_rows(max_row=max(sheet.max_row,original.max_row),max_col=max(sheet.max_column,original.max_column)):
            for cell in row:
                if (sheet.title,cell.coordinate) in editable:
                    continue
                old=original[cell.coordinate]
                equal=(_formula_reference_key(cell.value,locale['sheets'])==_formula_reference_key(old.value,locale['sheets']) if cell.data_type==old.data_type=='f' else cell.value==old.value)
                blank = cell.value in (None,'') and old.value in (None,'') and cell.data_type!='f'
                if not blank and (not equal or cell.data_type!=old.data_type):
                    raise ValueError('Changed source/formula: '+sheet.title+'!'+cell.coordinate)
                if old.hyperlink and (not cell.hyperlink or cell.hyperlink.location!=old.hyperlink.location):
                    raise ValueError('Changed navigation: '+cell.coordinate)
                if cell.comment:
                    model.public_text(cell.comment.text,'comment')
                if cell.value is not None:
                    height=sheet.row_dimensions[cell.row].height or 15
                    if height<(str(cell.value).count('\n')+1)*(cell.font.sz or 11):
                        raise ValueError('Clipped explicit lines: '+cell.coordinate)
    if phase=='complete':
        reviews=[model._records(value,model.ScreenshotReview) for value in screenshots]
        if Counter(image_bindings)!=Counter((r.case_id,r.variant,r.sha256) for r in reviews):
            raise ValueError('Screenshot review identity/digest mismatch')
    allowed={value.removeprefix('=') for sheet in pristine for row in sheet for cell in row if cell.data_type=='f' for value in (cell.value,)}
    urls=_package(raw,data,phase,{'closure_confirmation':closure_confirmation,'evidence_access':evidence_access,'screenshots':screenshots},gaps,allowed_formulas=allowed)
    model._unchanged(captures)
    counts=Counter(token for _,token,_ in observations)
    passed=sum(token=='PASS' and valid for _,token,valid in observations)
    evaluated=sum(valid for _,_,valid in observations)
    case_states={}
    for card in cards:
        members=[(token,valid) for record,token,valid in observations if record.case_id==card['id']]
        case_states[card['id']]=('FAIL' if any(token=='FAIL' and valid for token,valid in members) else 'PASS' if all(token=='PASS' and valid for token,valid in members) else 'SKIPPED' if all(token=='SKIPPED' for token,_ in members) else 'BLOCKED' if any(token=='BLOCKED' for token,_ in members) else 'NOT RUN' if all(token=='NOT RUN' for token,_ in members) else 'INCOMPLETE')
    return {'family':model.FAMILY,'version':model.VERSION,'language':language,'cases':len(cards),'variants':len(data.rows),'states':dict(counts),'case_states':case_states,'passed':passed,'evaluated':evaluated,'pass_rate':passed/evaluated if evaluated else None,'complete':phase=='complete' and not gaps,'read_only':True,'gaps':gaps,'layout_notes':layout_notes,'links':len(urls)}
