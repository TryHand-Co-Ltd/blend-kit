"""Populate the versioned editable report template without overwriting a run."""
from __future__ import annotations

import io
import unicodedata
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.hyperlink import Hyperlink

from report_model import ReportData, display_lines as _lines

INK = '172033'
LINK = '1D4ED8'
INPUT_FILL = 'FFF3CD'


def apply_reader_spacing(book):
    """Whole overview links, horizontal inset and middle/row-start alignment."""
    from copy import copy
    from openpyxl.cell.cell import MergedCell
    from openpyxl.cell.rich_text import CellRichText
    for index,sheet in enumerate(book.worksheets[:2]):
        for row in sheet.iter_rows(max_col=3 if index==0 else 5):
            number=row[0].row
            if sheet.row_dimensions[number].hidden or not any(c.value is not None for c in row):
                continue
            table_row=index==1 and ((not isinstance(sheet.cell(number,3),MergedCell) and sheet.cell(number,3).value is not None)
                or any(sheet.cell(number,c).protection.locked is False for c in (4,5)))
            for cell in row:
                if isinstance(cell,MergedCell):
                    continue
                alignment=copy(cell.alignment)
                alignment.vertical='top' if table_row else 'center'
                alignment.wrap_text=True
                alignment.horizontal='right' if index==0 and cell.column==2 and 16<=number<=27 else 'left'
                alignment.indent=1
                cell.alignment=alignment
                if index==0:
                    if isinstance(cell.value,CellRichText):
                        cell.value=str(cell.value)
                        cell.data_type='s'
                    if cell.hyperlink:
                        cell.font=Font(name=sheet['A1'].font.name,size=11,color=LINK,underline='single')


def _text(cell, value, *, emphasis=True):
    """All source text stays literal, including leading equals signs."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    import re
    # Emphasize source UI labels and decision values without changing literals.
    pattern = r'[^\s,;:.\n（）()]+(?: [^\s,;:.\n（）()]+){0,3}（[^）]+）|\b(?:AND|OR)\b|\b[A-Z](?:≥|≤|>=|<=|=|<|>)-?\d+(?:\.\d+)?%?'
    if cell.column == 3:
        pattern += r'|Không áp dụng|Không đỏ|Đỏ'
    if emphasis and isinstance(value, str) and not value.startswith('='):
        runs, offset = [], 0
        for match in re.finditer(pattern, value):
            if match.start() > offset:
                runs.append(value[offset:match.start()])
            runs.append(TextBlock(InlineFont(b=True), match.group()))
            offset = match.end()
        if runs:
            if offset < len(value):
                runs.append(value[offset:])
            value = CellRichText(runs)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'
    cell.font = Font(name=cell.parent['A1'].font.name, size=11, color=INK)
    cell.alignment = Alignment(vertical='top', wrap_text=True)


def _chunks(value, width=95, max_lines=17):
    """Split only display rows; concatenating chunks preserves every character."""
    value = str(value)
    if _lines(value, width) <= max_lines:
        return [value]
    result, chunk = [], []
    completed, units = 0, 0
    for char in value:
        char_units = 2 if unicodedata.east_asian_width(char) in ('W','F') else 1
        prospective = (completed + max(1,(units+width-1)//width) + 1 if char == '\n' else
            completed + max(1,(units+char_units+width-1)//width))
        if chunk and prospective > max_lines:
            result.append(''.join(chunk))
            chunk = []
            completed, units = 0, 0
        if char == '\n':
            completed += max(1,(units+width-1)//width)
            units = 0
        else:
            units += char_units
        chunk.append(char)
    if chunk:
        result.append(''.join(chunk))
    return result


def _height(sheet, row, values, widths, *, line_height=17, padding=10):
    lines = max(_lines(value, width) for value, width in zip(values, widths))
    sheet.row_dimensions[row].height = min(409, max(30, lines * line_height + padding))


def _jump(cell, sheet, row, caption=None):
    if caption is not None:
        _text(cell, caption,emphasis=False)
    else:
        _text(cell,str(cell.value),emphasis=False)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet.replace(chr(39), chr(39)*2)}'!A{row}", display=str(cell.value))
    cell.font = Font(name=cell.parent['A1'].font.name, size=11,
                     color=LINK,underline='single')


def _finish(sheet, end_column, end_row, header=None):
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = False
    sheet.print_area = f'A1:{end_column}{end_row}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    if header:
        sheet.freeze_panes = f'B{header + 1}'
        sheet.print_title_rows = f'{header}:{header}'


def build_workbook(data: ReportData):
    from block_report import build
    return build(data)


def serialize_report(book) -> bytes:
    """Keep Excel's privacy setting on newly serialized report bytes."""
    buffer = io.BytesIO()
    book.save(buffer)
    output = io.BytesIO()
    with ZipFile(buffer) as source, ZipFile(output, 'w') as target:
        for entry in source.infolist():
            raw = source.read(entry.filename)
            if entry.filename == 'xl/workbook.xml':
                root = ET.fromstring(raw)
                namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
                properties = root.find(f'{{{namespace}}}workbookPr')
                if properties is None:
                    properties = ET.Element(f'{{{namespace}}}workbookPr')
                    root.insert(1 if root.find(f'{{{namespace}}}fileVersion') is not None else 0, properties)
                properties.set('filterPrivacy', '1')
                raw = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            elif entry.filename.startswith('xl/worksheets/') and entry.filename.endswith('.xml'):
                root = ET.fromstring(raw)
                namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
                changed = False
                for text in root.iter(f'{{{namespace}}}t'):
                    # openpyxl 3.1.5 omits xml:space on whitespace-only rich
                    # runs. Excel rejects those string properties on normal open.
                    if text.text and text.text != text.text.strip():
                        text.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                        changed = True
                if changed:
                    raw = ET.tostring(root,encoding='utf-8',xml_declaration=True)
            target.writestr(entry, raw)
    return output.getvalue()


def render_report(data: ReportData, output: Path) -> dict:
    output = Path(output)
    if output.exists():
        raise FileExistsError('Existing output is preserved')
    book = build_workbook(data)
    raw = serialize_report(book)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('xb') as handle:
        handle.write(raw)
    return {'sheets': book.sheetnames, 'rows': len(data.rows),
            'detail_cases': len({book.worksheets[2].cell(row, 1).value for row in range(5, book.worksheets[2].max_row + 1)}) if len(book.worksheets) > 2 else 0,
            'images': 0}
