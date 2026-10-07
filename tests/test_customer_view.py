"""Customer layout parity and privacy regressions; no application execution."""
import hashlib
import io
from pathlib import Path
import sys
import tempfile
import unittest
from collections import Counter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'skills/blend-generate-test-spec/scripts'))
import block_report as blocks
import export_report
import report_model as model
from openpyxl import load_workbook
from render_report import serialize_report


class CustomerViewTests(unittest.TestCase):
    def test_reader_inset_is_idempotent_and_overview_links_plain(self):
        from render_report import apply_reader_spacing
        from openpyxl.cell.rich_text import CellRichText
        source=ROOT/'tests/fixtures/test-spec/inputs/valid'
        data,_=model.prepare_report(source)
        book=blocks.build(data)
        summary,tests=book.worksheets
        _,inputs=blocks.layout(data)
        fields=next(iter(inputs.values()))
        self.assertNotIn('\n',tests[fields['actual']].number_format)
        self.assertIn(tests[fields['actual']].value,(None,''))
        self.assertEqual(summary['B18'].alignment.horizontal,'right')
        self.assertEqual(summary['B18'].alignment.indent,1)
        self.assertEqual(summary['B18'].alignment.vertical,'center')
        self.assertEqual(tests[fields['actual']].alignment.vertical,'top')
        data_row=tests[fields['actual']].row
        self.assertEqual(tests.row_dimensions[data_row-1].height,3)
        self.assertFalse(tests.row_dimensions[data_row-1].hidden)
        self.assertFalse(any(tests.cell(data_row-1,c).value for c in range(1,6)))
        self.assertEqual(tests.cell(data_row-1,4).fill.fgColor.rgb[-6:],'F8FAFC')
        self.assertEqual(tests.cell(5,1).alignment.vertical,'center')
        self.assertTrue(all(not isinstance(c.value,CellRichText) for row in summary for c in row))
        for row in range(model.CASE_START_ROW,model.CASE_START_ROW+len(data.cases)):
            self.assertIsNotNone(summary.cell(row,2).hyperlink)
            self.assertEqual(summary.cell(row,2).font.color.rgb[-6:],'1D4ED8')
            self.assertEqual(summary.cell(row,2).font.underline,'single')
        before=[(s.title,c.coordinate,c.value,c.number_format) for s in book for row in s for c in row]
        heights=[(s.title,r,d.height) for s in book for r,d in s.row_dimensions.items()]
        apply_reader_spacing(book)
        self.assertEqual(before,[(s.title,c.coordinate,c.value,c.number_format) for s in book for row in s for c in row])
        self.assertEqual(heights,[(s.title,r,d.height) for s in book for r,d in s.row_dimensions.items()])

    def test_recorded_results_do_not_depend_on_readiness_or_actual(self):
        source = ROOT/'tests/fixtures/test-spec/inputs/valid'
        data,_ = model.prepare_report(source)
        cards,inputs=blocks.layout(data)
        case=next(card for card in cards if any(not r.eligible for r in card['members']))
        formula=blocks.formulas(data,cards,inputs)[(1,f'B{case["status_row"]}')]
        self.assertNotIn('LEN(',formula)
        self.assertIn('Đang thực hiện',formula)
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'report.xlsx'
            book=blocks.build(data)
            for record in case['members']:
                book.worksheets[1][inputs[record.identity]['status']]='Đạt'
            path.write_bytes(serialize_report(book))
            receipt=blocks.check(path,source)
            self.assertEqual(receipt['case_states'][case['id']],'PASS')
            self.assertEqual(receipt['passed'],len(case['members']))
            self.assertEqual(receipt['quality_qualified_passed'],0)
            self.assertTrue(any('actual result required' in gap for gap in receipt['gaps']))
            book.worksheets[1][inputs[case['members'][0].identity]['status']]='Không đạt'
            path.write_bytes(serialize_report(book))
            self.assertEqual(blocks.check(path,source)['case_states'][case['id']],'FAIL')
            first=inputs[case['members'][0].identity]['status']
            for label,expected in (('Bị chặn','BLOCKED'),('Chưa thực hiện','INCOMPLETE' if len(case['members'])>1 else 'NOT RUN')):
                book.worksheets[1][first]=label
                path.write_bytes(serialize_report(book))
                self.assertEqual(blocks.check(path,source)['case_states'][case['id']],expected)
            for record in case['members']:
                book.worksheets[1][inputs[record.identity]['status']]='Không thực hiện'
            path.write_bytes(serialize_report(book))
            self.assertEqual(blocks.check(path,source)['case_states'][case['id']],'SKIPPED')

    def test_rich_text_reopen_preserves_literals_and_checks_combined_privacy(self):
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from render_report import _text
        from openpyxl import Workbook
        from check_report import _package
        book=Workbook()
        literal='Chọn Hủy（キャンセル） rồi kiểm tra A=50 AND A<70.'
        _text(book.active['A1'],literal)
        self.assertIsInstance(book.active['A1'].value,CellRichText)
        _text(book.active['A3'],literal,emphasis=False)
        self.assertIsInstance(book.active['A3'].value,str)
        saved=load_workbook(io.BytesIO(serialize_report(book)),rich_text=True)
        self.assertEqual(str(saved.active['A1'].value),literal)
        self.assertTrue(any(isinstance(run,TextBlock) and run.font.b for run in saved.active['A1'].value))
        # Native Excel requires preservation on even whitespace-only runs.
        book.active['A2']=CellRichText([TextBlock(InlineFont(b=True),'First'),' ',TextBlock(InlineFont(b=True),'Second')])
        from zipfile import ZipFile
        import xml.etree.ElementTree as ET
        with ZipFile(io.BytesIO(serialize_report(book))) as package:
            root=ET.fromstring(package.read('xl/worksheets/sheet1.xml'))
            texts=list(root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t'))
            self.assertTrue(any(t.text==' ' and t.get('{http://www.w3.org/XML/1998/namespace}space')=='preserve' for t in texts))
        source=ROOT/'tests/fixtures/test-spec/inputs/valid'
        data,_=model.prepare_report(source)
        report=blocks.build(data)
        fields=next(iter(blocks.layout(data)[1].values()))
        report.worksheets[1][fields['actual']]=CellRichText(['secret',TextBlock(InlineFont(b=True),'=hidden')])
        with self.assertRaisesRegex(ValueError,'credential'):
            _package(serialize_report(report),data,'in-progress',{},[])

    def test_binding_sheet_matches_current_workbook(self):
        from update_report import report_bindings
        source = ROOT/'tests/fixtures/test-spec/inputs/valid'
        with tempfile.TemporaryDirectory() as directory:
            data, _ = model.prepare_report(source)
            path = Path(directory)/'current.xlsx'
            path.write_bytes(serialize_report(blocks.build(data)))
            binding = report_bindings(source,path)
            self.assertEqual(binding['sheet'],load_workbook(path).worksheets[1].title)
            self.assertEqual(binding['schema_version'],'2.5.0')

    def test_reader_removes_operational_filler_keeps_specific_oracles_and_actions(self):
        from copy import deepcopy
        source = ROOT/'tests/fixtures/test-spec/inputs/valid'
        design = deepcopy(model.working.parse_sources(source,'vi'))
        case = design['cases'][0]
        case['preservation'] = 'Giữ điểm, cấu hình và đối tượng không đích theo expected; chỉ thay phần thủ tục yêu cầu, không tự đổi oracle hoặc mở rộng phạm vi.'
        case['variants'] = [(name,inputs+'; Thao tác riêng: Run only branch '+name,expected) for name,inputs,expected in case['variants']]
        reader = model.customer_report(model.project_report(design,'vi'),design)
        card = reader.cases[0]
        self.assertEqual(card['action_items'],[])
        for name,branch in card['variants'].items():
            self.assertEqual(branch['action_delta'],'Run only branch '+name)
            self.assertNotIn('theo expected',branch['expected'])
            self.assertNotIn('Thao tác riêng:',branch['inputs'])
        for card in reader.cases:
            self.assertFalse(any(label in ('Lưu ý','Cần xác minh','Chuẩn bị') for label,_ in card['preparation_items']))
        self.assertNotIn('Chuẩn bị:',reader.summary['scope'])
        self.assertNotIn('Điều kiện bắt đầu:',reader.summary['scope'])
        self.assertNotIn('quan sát đã ghi nhận',reader.summary['limitations'])
        self.assertTrue(any('School B/year Z unchanged' in branch['expected'] for card in reader.cases for branch in card['variants'].values()))
        for caption in ('Picture','Image 1','screenshot','画像'):
            self.assertFalse(model.descriptive_image_title(caption))
        self.assertTrue(model.descriptive_image_title('Score 30 is not red after saving'))

    def test_native_deduplicated_payload_keeps_per_instance_reviews(self):
        from zipfile import ZipFile
        import xml.etree.ElementTree as ET
        from PIL import Image as Pixels
        from update_report import _add_readable_images
        from update_report import _refresh_evidence_headers
        source = ROOT/'tests/fixtures/test-spec/inputs/compact-scenario'
        data, _ = model.prepare_report(source,customer=True)
        book = blocks.build(data)
        _, inputs = blocks.layout(data)
        identity, fields = next(iter(inputs.items()))
        case, variant = identity.split(' / ')
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            pixels = directory/'image.png'
            Pixels.new('RGB',(120,72),'red').save(pixels)
            digest = hashlib.sha256(pixels.read_bytes()).hexdigest()
            reviews = []
            for index, cp in enumerate(fields['checkpoints'].values(),1):
                if cp['artifact'] != 'screenshot':
                    continue
                title = f'Synthetic view {index}'
                _add_readable_images(book.worksheets[1],cp,[({'case_id':case,'variant_id':variant,
                    'checkpoint_id':cp['id'],'annotated_sha256':digest,'observed_note':title},pixels)])
                reviews.append({'case_id':case,'variant':variant,'checkpoint_id':cp['id'],'sha256':digest,
                    'caption':title,'reviewed_by':'fixture','reviewed_at':'2026-10-05T12:00:00+07:00','source':'synthetic fixture'})
            self.assertEqual(len(reviews),2)
            _refresh_evidence_headers(book.worksheets[1],data,inputs)
            raw = serialize_report(book)
            output = io.BytesIO()
            with ZipFile(io.BytesIO(raw)) as old, ZipFile(output,'w') as deduplicated:
                media = [name for name in old.namelist() if name.startswith('xl/media/')]
                self.assertEqual(len(media),2)
                for entry in old.infolist():
                    if entry.filename == media[1]:
                        continue
                    content = old.read(entry.filename)
                    if entry.filename.startswith('xl/drawings/_rels/'):
                        relationships = ET.fromstring(content)
                        for relationship in relationships:
                            target = relationship.get('Target','')
                            if target.endswith('/'+media[1].split('/')[-1]):
                                relationship.set('Target', '/'+media[0])
                        content = ET.tostring(relationships)
                    deduplicated.writestr(entry,content)
            report = directory/'native-deduplicated.xlsx'
            report.write_bytes(output.getvalue())
            self.assertEqual(len(load_workbook(report).worksheets[1]._images),2)
            closure = {'closed_by':'fixture','closed_at':'2026-10-05T12:00:00+07:00','source':'synthetic fixture','audience':'fixture'}
            checked = blocks.check(report,source,phase='complete',closure_confirmation=closure,screenshots=reviews)
            self.assertFalse(any('Screenshot reviews' in gap or 'matching screenshot review' in gap for gap in checked['gaps']))
            with self.assertRaisesRegex(ValueError,'identity/digest mismatch'):
                blocks.check(report,source,phase='complete',closure_confirmation=closure,
                    screenshots=[{**review,'checkpoint_id':'wrong'} for review in reviews])
            Pixels.new('RGB',(120,72),'blue').save(pixels)
            with ZipFile(report,'a') as archive:
                archive.writestr('xl/media/unbound.png',pixels.read_bytes())
            with self.assertRaisesRegex(ValueError,'media payloads differ'):
                blocks.check(report,source)

    def test_linear_chunks_preserve_frozen_display_boundaries(self):
        import random
        from render_report import _chunks
        def legacy(value,width,max_lines):
            if model.display_lines(value,width) <= max_lines:
                return [value]
            result, chunk = [], ''
            for char in value:
                if chunk and model.display_lines(chunk+char,width) > max_lines:
                    result.append(chunk)
                    chunk = ''
                chunk += char
            if chunk:
                result.append(chunk)
            return result
        randomizer = random.Random(42)
        values = ['', '\n', '\n\n', '日本語\nĐiểm 30\n']
        values += [''.join(randomizer.choices('ABC日本語 30\n\t',k=300)) for _ in range(8)]
        for value in values:
            for width in (1,2,7,104):
                for lines in (1,3,23):
                    self.assertEqual(_chunks(value,width,lines),legacy(value,width,lines))

    def test_current_literal_inputs_and_ooxml_privacy(self):
        from zipfile import ZipFile
        from openpyxl.comments import Comment
        from update_report import report_bindings, update_report
        source = ROOT/'tests/fixtures/test-spec/inputs/compact-scenario'
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            report = directory/'current.xlsx'
            export_report.export(source,report)
            binding = report_bindings(source,report)
            fields = next(iter(binding['variants'].values()))
            identity = {key:binding[key] for key in ('design_revision','feature_id')}
            identity.update(case_id=fields['case_id'],variant_id=fields['variant_id'],run_id='2026-10-06-001')
            literal = '=Literal observation &lt;b&gt; <br>  exact whitespace; inspect failure.'
            update_report(source,report,{'identity':identity,'status':'FAIL','actual':literal})
            book = load_workbook(report)
            cell = book[binding['sheet']][fields['actual_cell']]
            self.assertEqual((cell.value,cell.data_type),(literal,'s'))
            blocks.check(report,source)
            original = report.read_bytes()
            for value in ('http://localhost:3000/private','file:///private','C:/private/report'):
                with self.assertRaises(ValueError):
                    update_report(source,report,{'identity':identity,'status':'FAIL','actual':value})
                self.assertEqual(report.read_bytes(),original)
            cell.value = '=1+1'
            rejected = directory/'formula.xlsx'
            rejected.write_bytes(serialize_report(book))
            with self.assertRaises(ValueError):
                blocks.check(rejected,source)
            book = load_workbook(report)
            book[binding['sheet']][fields['actual_cell']].comment = Comment('http://localhost/private','fixture')
            rejected.write_bytes(serialize_report(book))
            with self.assertRaises(ValueError):
                blocks.check(rejected,source)
            # Native Office XML attributes must receive the same privacy scan as cells.
            with ZipFile(io.BytesIO(original)) as archive:
                content = {entry.filename:archive.read(entry.filename) for entry in archive.infolist()}
            import xml.etree.ElementTree as ET
            worksheet = ET.fromstring(content['xl/worksheets/sheet2.xml'])
            worksheet.find('.//{*}c').set('privateLocator','C:/private/report')
            content['xl/worksheets/sheet2.xml'] = ET.tostring(worksheet)
            with ZipFile(rejected,'w') as archive:
                for name,raw in content.items():
                    archive.writestr(name,raw)
            with self.assertRaises(ValueError):
                blocks.check(rejected,source)
            self.assertEqual(report.read_bytes(),original)

    def test_old_workbook_versions_are_rejected_without_mutation(self):
        from update_report import report_bindings, update_report
        source = ROOT/'tests/fixtures/test-spec/inputs/valid'
        data, _ = model.prepare_report(source)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'old.xlsx'
            for version in ('1.0.0','1.1.0','1.2.0','2.0.0','2.1.0','2.2.0','2.3.0'):
                book = blocks.build(data)
                book.custom_doc_props['TemplateVersion'].value = version
                path.write_bytes(serialize_report(book))
                before = path.read_bytes()
                for call in (lambda: model.prepare_saved_report(source,path),
                             lambda: blocks.check(path,source),
                             lambda: report_bindings(source,path),
                             lambda: update_report(source,path,{})):
                    with self.assertRaises(ValueError):
                        call()
                    self.assertEqual(path.read_bytes(),before)

    def test_image_only_writeback_titles_collapse_identity_and_retry(self):
        import importlib.util
        from PIL import Image as Pixels
        from update_report import report_bindings, update_report
        spec = importlib.util.spec_from_file_location('readable_archive', ROOT/'skills/blend-automation-test/scripts/run_artifacts.py')
        archive = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(archive)
        source = ROOT/'tests/fixtures/test-spec/inputs/compact-scenario'
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            report = folder/'report.xlsx'
            export_report.export(source,report,customer=True)
            binding = report_bindings(source,report)
            name, fields = next(iter(binding['variants'].items()))
            checkpoint = next(cp for cp,value in fields['checkpoints'].items() if value['artifact']=='screenshot')
            identity = {key:binding[key] for key in ('design_revision','feature_id')}
            identity.update(case_id=fields['case_id'],variant_id=fields['variant_id'],run_id='2026-10-05-001')
            feature = folder/'context/features/SYN-SC-controls'
            feature.mkdir(parents=True)
            (feature/'README.md').write_text('Synthetic fixture',encoding='utf-8')
            (feature/'CONTEXT.md').write_text('Synthetic fixture',encoding='utf-8')
            metadata = {**{key:identity[key] for key in ('design_revision','feature_id')}, 'report':str(report),
                'inventory':[{'case_id':fields['case_id'],'variant_id':fields['variant_id']}]}
            run = archive.create_run(folder/'context','SYN-SC-controls',metadata,identity['run_id'])
            evidence = []
            for sequence,color in enumerate(('red','green','blue','orange'),1):
                pixels = folder/'pixels.png'
                Pixels.new('RGB',(1200,720),color).save(pixels)
                capture_identity = {**identity,'checkpoint_id':checkpoint}
                capture = {'full_page':False,'viewport':{'width':1200,'height':720},'observed_state':'Synthetic fixture',
                    'annotation_present':False,'captured_at':'2026-10-05T12:00:00+07:00','capture_reference':f'raw-{sequence}'}
                raw = archive.archive_capture(run,pixels,capture_identity,'raw',sequence,capture)
                annotated = archive.archive_capture(run,pixels,capture_identity,'annotated',sequence,
                    {**capture,'annotation_present':True,'capture_reference':f'annotated-{sequence}'})
                evidence.append({**capture_identity,'sequence':sequence,'raw_path':raw['path'],'annotated_path':annotated['path'],
                    'raw_sha256':raw['sha256'],'annotated_sha256':annotated['sha256'],'full_page':False,'viewport':capture['viewport'],
                    'assertion':'Synthetic assertion','focus':'Synthetic region','observed_note':f'Image {sequence}: observed {color} fixture.',
                    'reviewed_by':'fixture','reviewed_at':'2026-10-05T12:00:02+07:00','source':{'capture_tool':'fixture',
                        'raw_capture':f'raw-{sequence}','annotated_capture':f'annotated-{sequence}',
                        'annotation_method':'Synthetic fixture','pixel_review':{'tool':'fixture','reference':'synthetic-check'}}})
            evidence[0]['title']='Xác nhận thành tích（成績確認） — S01 hiển thị *29'
            evidence[1]['title']='成績確認 — 赤点は*29と表示'
            payload = {'identity':identity,'status':'FAIL','actual':'Synthetic mismatch observed; inspect the result.','evidence':evidence}
            result = update_report(source,report,payload,run_dir=run)
            self.assertEqual(result['evidence_added'],4)
            saved = load_workbook(report)
            sheet = saved.worksheets[1]
            self.assertEqual(saved.sheetnames,['Tổng quan','Kiểm thử'])
            heading=sheet.cell(fields['checkpoints'][checkpoint]['label_row'],1)
            self.assertEqual((heading.alignment.horizontal,heading.alignment.vertical,heading.alignment.indent),('left','center',1))
            for image in sheet._images:
                title = sheet.cell(image.anchor._from.row,1)
                self.assertEqual((title.alignment.horizontal,title.alignment.vertical,title.alignment.indent),('left','center',1))
                self.assertEqual(title.value,image.anchor.pic.nvPicPr.cNvPr.descr)
                self.assertTrue(sheet.row_dimensions[title.row].hidden)
                self.assertEqual(sheet.row_dimensions[title.row].outlineLevel,1)
                self.assertTrue(sheet.row_dimensions[image.anchor._from.row+1].hidden)
            self.assertEqual(update_report(source,report,payload,run_dir=run)['evidence_added'],0)
            before = report.read_bytes()
            with self.assertRaisesRegex(ValueError,'identity mismatch'):
                update_report(source,report,{**payload,'identity':{**identity,'design_revision':'wrong'}},run_dir=run)
            self.assertEqual(report.read_bytes(),before)
            self.assertEqual(blocks.check(report,source)['links'],0)

    def test_source_11_result_evidence_slot_writeback_and_retry(self):
        import importlib.util
        from PIL import Image as Pixels
        from update_report import report_bindings, update_report
        spec = importlib.util.spec_from_file_location('readable_archive', ROOT/'skills/blend-automation-test/scripts/run_artifacts.py')
        archive = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(archive)
        fixture = ROOT/'tests/fixtures/test-spec/inputs/valid'
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            import shutil
            source = folder/'source'
            shutil.copytree(fixture,source)
            cases = source/'test-cases.vi.md'
            cases.write_text(cases.read_text(encoding='utf-8').replace('test-cases@1.0.0','test-cases@1.1.0'),encoding='utf-8')
            cases_ja = source/'test-cases.ja.md'
            cases_ja.write_text(cases_ja.read_text(encoding='utf-8').replace('test-cases@1.0.0','test-cases@1.1.0'),encoding='utf-8')
            frozen_source = {path.name:path.read_bytes() for path in source.iterdir()}
            report = folder/'report.xlsx'
            export_report.export(source,report,customer=True)
            binding = report_bindings(source,report,feature_id='SYN-SC')
            name, fields = next(iter(binding['variants'].items()))
            self.assertEqual(fields['checkpoints'],{})
            checkpoint = next(iter(fields['evidence_slots']))
            self.assertEqual(checkpoint,'CP-result')
            self.assertIn('Eligible',fields['evidence_slots'][checkpoint]['expected'])
            self.assertEqual(fields['evidence_slots'][checkpoint]['source_anchor'],'variant Expected / existing Steps')
            others = {key:dict(value) for key,value in binding['variants'].items() if key != name}
            identity = {key:binding[key] for key in ('design_revision','feature_id')}
            identity.update(case_id=fields['case_id'],variant_id=fields['variant_id'],run_id='2026-10-05-001')
            feature = folder/'context/features/SYN-SC-controls'
            feature.mkdir(parents=True)
            (feature/'README.md').write_text('Synthetic fixture',encoding='utf-8')
            (feature/'CONTEXT.md').write_text('Synthetic fixture',encoding='utf-8')
            metadata = {**{key:identity[key] for key in ('design_revision','feature_id')}, 'report':str(report),
                'inventory':[{'case_id':fields['case_id'],'variant_id':fields['variant_id']}]}
            run = archive.create_run(folder/'context','SYN-SC-controls',metadata,identity['run_id'])
            evidence = []
            for sequence,color in enumerate(('red','green','blue','orange'),1):
                pixels = folder/'pixels.png'
                Pixels.new('RGB',(1200,720),color).save(pixels)
                capture_identity = {**identity,'checkpoint_id':checkpoint}
                capture = {'full_page':False,'viewport':{'width':1200,'height':720},'observed_state':'Synthetic fixture',
                    'annotation_present':False,'captured_at':'2026-10-05T12:00:00+07:00','capture_reference':f'raw-{sequence}'}
                raw = archive.archive_capture(run,pixels,capture_identity,'raw',sequence,capture)
                annotated = archive.archive_capture(run,pixels,capture_identity,'annotated',sequence,
                    {**capture,'annotation_present':True,'capture_reference':f'annotated-{sequence}'})
                evidence.append({**capture_identity,'sequence':sequence,'raw_path':raw['path'],'annotated_path':annotated['path'],
                    'raw_sha256':raw['sha256'],'annotated_sha256':annotated['sha256'],'full_page':False,'viewport':capture['viewport'],
                    'assertion':'Synthetic assertion','focus':'Synthetic region','observed_note':f'Image {sequence}: observed {color} fixture.',
                    'reviewed_by':'fixture','reviewed_at':'2026-10-05T12:00:02+07:00','source':{'capture_tool':'fixture',
                        'raw_capture':f'raw-{sequence}','annotated_capture':f'annotated-{sequence}',
                        'annotation_method':'Synthetic fixture','pixel_review':{'tool':'fixture','reference':'synthetic-check'}}})
            payload = {'identity':identity,'status':'FAIL','actual':'Synthetic mismatch observed; inspect the result.','evidence':evidence}
            result = update_report(source,report,payload,run_dir=run,feature_id='SYN-SC')
            self.assertEqual(result['evidence_added'],4)
            saved = load_workbook(report)
            sheet = saved.worksheets[1]
            self.assertEqual(saved.sheetnames,['Tổng quan','Kiểm thử'])
            for image in sheet._images:
                title = sheet.cell(image.anchor._from.row,1)
                self.assertEqual(title.value,image.anchor.pic.nvPicPr.cNvPr.descr)
                self.assertTrue(sheet.row_dimensions[title.row].hidden)
                self.assertEqual(sheet.row_dimensions[title.row].outlineLevel,1)
                self.assertTrue(sheet.row_dimensions[image.anchor._from.row+1].hidden)
            self.assertEqual(update_report(source,report,payload,run_dir=run,feature_id='SYN-SC')['evidence_added'],0)
            before = report.read_bytes()
            with self.assertRaisesRegex(ValueError,'identity mismatch'):
                update_report(source,report,{**payload,'identity':{**identity,'design_revision':'wrong'}},run_dir=run,feature_id='SYN-SC')
            self.assertEqual(report.read_bytes(),before)
            self.assertEqual(blocks.check(report,source)['links'],0)
            self.assertEqual(report.read_bytes(),before)
            self.assertEqual({path.name:path.read_bytes() for path in source.iterdir()},frozen_source)
            for other in others.values():
                self.assertEqual(saved[binding['sheet']][other['actual_cell']].value,None)
                self.assertEqual(saved[binding['sheet']][other['status_cell']].value,'Chưa thực hiện')
            for bad_record in ({**evidence[0],'checkpoint_id':'unknown'}, {**evidence[0],'variant_id':'wrong'}, {**evidence[0],'annotated_sha256':'0'*64}):
                with self.assertRaises(ValueError):
                    update_report(source,report,{**payload,'evidence':[bad_record]},run_dir=run,feature_id='SYN-SC')
                self.assertEqual(report.read_bytes(),before)
            old = load_workbook(report)
            old.custom_doc_props['TemplateVersion'].value = '2.3.0'
            unsupported = folder/'old-format.xlsx'
            unsupported.write_bytes(serialize_report(old))
            frozen_old = unsupported.read_bytes()
            for call in (lambda: report_bindings(source,unsupported,feature_id='SYN-SC'),
                         lambda: blocks.check(unsupported,source),
                         lambda: update_report(source,unsupported,payload,run_dir=run,feature_id='SYN-SC')):
                with self.assertRaises(ValueError):
                    call()
                self.assertEqual(unsupported.read_bytes(),frozen_old)

    def test_customer_writer_uses_compact_actual_width_and_checkpoint_capacity(self):
        from update_report import report_bindings, update_report
        source = ROOT/'tests/fixtures/test-spec/inputs/compact-scenario'
        with tempfile.TemporaryDirectory() as directory:
            report = Path(directory)/'customer.xlsx'
            export_report.export(source, report, customer=True)
            binding = report_bindings(source, report)
            name, fields = next(iter(binding['variants'].items()))
            self.assertEqual(fields['evidence_capacity']['rows_per_checkpoint'], blocks.READABLE_EVIDENCE_ROWS)
            identity = {key: binding[key] for key in ('design_revision', 'feature_id')}
            identity.update(case_id=fields['case_id'], variant_id=fields['variant_id'], run_id='2026-10-05-001')
            for actual in ('Observed mismatch; fix required.', 'Observed mismatch; next action: inspect result. '*9, 'Observed mismatch; fix required.'):
                update_report(source, report, {'identity': identity, 'status': 'FAIL', 'actual': actual, 'evidence': []})
                saved = load_workbook(report)
                cell = saved.worksheets[1][fields['actual_cell']]
                height = saved.worksheets[1].row_dimensions[cell.row].height
                self.assertEqual(cell.value, actual)
                self.assertGreaterEqual(height, model.display_lines(actual, 37)*14+6)
                for address in (fields['actual_cell'],fields['status_cell']):
                    alignment = saved.worksheets[1][address].alignment
                    self.assertEqual((alignment.horizontal,alignment.vertical,alignment.indent),('left','top',1))
                if len(actual)<50:
                    self.assertLess(height, 90)

    def test_localized_actual_and_checkpoint_ledger_are_independent(self):
        from update_report import report_bindings, update_report, _archive_helper
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        source=ROOT/'tests/fixtures/test-spec/inputs/compact-scenario'
        for language,actual in (
                ('vi','Xác nhận thành tích（成績確認）: ô S01 Đỏ hiển thị *29; điểm không thay đổi.'),
                ('ja','成績確認: S01の赤点は*29と表示され、点数は変更されない。')):
            with self.subTest(language=language), tempfile.TemporaryDirectory() as directory:
                folder=Path(directory)
                report=folder/'report.xlsx'
                export_report.export(source,report,language=language,customer=True)
                binding=report_bindings(source,report,language)
                target,fields=next((name,value) for name,value in binding['variants'].items()
                    if value['checkpoints'] and all(cp['artifact']=='inspection' for cp in value['checkpoints'].values()))
                identity={key:binding[key] for key in ('design_revision','feature_id')}
                identity.update(case_id=fields['case_id'],variant_id=fields['variant_id'],run_id='2026-10-05-001')
                feature=folder/'context/features/SYN-SC-controls'
                feature.mkdir(parents=True)
                for name in ('README.md','CONTEXT.md'):
                    (feature/name).write_text('Synthetic fixture',encoding='utf-8')
                archive=_archive_helper()
                run=archive.create_run(folder/'context','SYN-SC-controls',{
                    **{key:identity[key] for key in ('design_revision','feature_id')},'report':str(report.resolve())},identity['run_id'])
                observations={cp:'Synthetic inspected observation.' for cp in fields['checkpoints']}
                payload={'identity':identity,'status':'PASS','actual':actual,'observations':observations}
                update_report(source,report,payload,language=language,run_dir=run)
                book=load_workbook(report,rich_text=True)
                value=book.worksheets[1][fields['actual_cell']].value
                self.assertIsInstance(value,CellRichText)
                self.assertEqual(str(value),actual)
                self.assertIn('*29',[str(run.text) for run in value if isinstance(run,TextBlock) and run.font.b])
                self.assertEqual(archive.records(run)[-1]['observations'],observations)
                frozen=report.read_bytes()
                for invalid in ({**payload,'actual':'CP-result: raw log'},
                                {**payload,'observations':{'unknown':'Observation'}},
                                {**payload,'status':'PASS','observations':{}}):
                    with self.assertRaises(ValueError):
                        update_report(source,report,invalid,language=language,run_dir=run)
                    self.assertEqual(report.read_bytes(),frozen)
                # PASS needs proof for the exact Actual, independently of visible prose.
                receipt=blocks.check(report,source,language,run_dir=run)
                self.assertFalse(any(target+': missing checkpoint observations' in gap for gap in receipt['gaps']))
                self.assertTrue(any(target+': missing checkpoint observations' in gap for gap in blocks.check(report,source,language)['gaps']))
                book.worksheets[1][fields['actual_cell']]=actual+' Different observation.'
                report.write_bytes(serialize_report(book))
                self.assertTrue(any(target+': missing checkpoint observations' in gap for gap in blocks.check(report,source,language,run_dir=run)['gaps']))

    def test_inline_code_uses_native_bold_without_added_brackets(self):
        from openpyxl import Workbook
        from openpyxl.cell.rich_text import TextBlock
        from render_report import _text
        book=Workbook()
        raw='Form `incomplete-deleted-stale`: `(29)`, `*29`, `29*`, `**24`, `[1,2]`, `a[0]`, `[A-Z]+`.'
        expected='Form incomplete-deleted-stale: (29), *29, 29*, **24, [1,2], a[0], [A-Z]+.'
        tokens=('incomplete-deleted-stale','(29)','*29','29*','**24','[1,2]','a[0]','[A-Z]+')
        projected=model.display_markup(raw)
        self.assertEqual(model.display_markup(projected),projected)
        _text(book.active['D1'],projected)
        _text(book.active['A2'],'`[1,2]`',emphasis=False)
        _text(book.active['A3'],'`=1`')
        _text(book.active['A4'],'incomplete-deleted-stale')
        saved=load_workbook(io.BytesIO(serialize_report(book)),rich_text=True)
        value=saved.active['D1'].value
        self.assertEqual(str(value),expected)
        bold=[run.text for run in value if isinstance(run,TextBlock) and run.font.b]
        for token in tokens:
            self.assertIn(token,bold)
        self.assertEqual(saved.active['A2'].value,'[1,2]')
        self.assertEqual((saved.active['A3'].value,saved.active['A3'].data_type),('=1','s'))
        self.assertTrue(saved.active['A4'].value[0].font.b)
        self.assertFalse(model.markdown_leak(value))
        self.assertTrue(model.markdown_leak('**unexpected Markdown**'))

    def test_supported_sources_generate_only_current_standalone_reports_without_source_changes(self):
        for language in ('vi', 'ja'):
            for fixture in ('valid', 'scenario-grouping', 'compact-scenario'):
                with self.subTest(language=language, fixture=fixture), tempfile.TemporaryDirectory() as directory:
                    source = ROOT/'tests/fixtures/test-spec/inputs'/fixture
                    paths = list(source.glob('*.md'))
                    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
                    original, _ = model.prepare_report(source, language)
                    customer, _ = model.prepare_report(source, language, customer=True)
                    self.assertEqual(original.rows, customer.rows)
                    report = Path(directory)/'test-report.xlsx'
                    export_report.export(source, report, language=language, customer=True)
                    receipt = blocks.check(report, source, language)
                    self.assertEqual(receipt['version'], model.CUSTOMER_VERSION)
                    self.assertEqual(receipt['variants'], len(original.rows))
                    book = load_workbook(report)
                    self.assertEqual(book.worksheets[0].freeze_panes, 'A2')
                    self.assertEqual(book.worksheets[1].freeze_panes, 'A2')
                    for sheet in book:
                        self.assertEqual(len(sheet.sheet_view.selection),1)
                        self.assertEqual(sheet.sheet_view.selection[0].pane,'bottomLeft')
                        self.assertEqual(sheet.sheet_view.pane.activePane,'bottomLeft')
                    self.assertFalse(book.worksheets[1].sheet_properties.outlinePr.summaryBelow)
                    cards, inputs = blocks.layout(customer)
                    tests = book.worksheets[1]
                    self.assertTrue(all(tests.row_dimensions[row].hidden for row in blocks.reader_blank_rows(customer,cards)))
                    for fields in inputs.values():
                        areas = [cp['picture_rows'] for cp in fields['checkpoints'].values()]
                        if 'picture_rows' in fields:
                            areas.append(fields['picture_rows'])
                        for rows in areas:
                            self.assertTrue(all(tests.row_dimensions[row].hidden and tests.row_dimensions[row].outlineLevel==0 for row in rows))
                    self.assertEqual(len(book.sheetnames), 2)
                    self.assertEqual(book.worksheets[0].auto_filter.ref,
                                     f'A29:C{29+len(cards)}')
                    book.worksheets[0].row_dimensions[model.CASE_START_ROW].hidden = True
                    report.write_bytes(serialize_report(book))
                    blocks.check(report, source, language)
                    for fields in inputs.values():
                        self.assertIsNone(tests[fields['actual']].value)
                        self.assertEqual(tests[fields['status']].value, model.REPORT_LAYOUTS[language]['statuses']['NOT RUN'])
                        link = tests.cell(tests[fields['status']].row, 6).hyperlink
                        self.assertIsNone(link)
                        self.assertFalse(tests.row_dimensions[tests[fields['status']].row].hidden)
                        for column in (1,2,3,4,5):
                            alignment = tests.cell(tests[fields['status']].row,column).alignment
                            self.assertEqual((alignment.horizontal,alignment.vertical,alignment.indent),('left','top',1))
                    # Core hiding cannot be used to make a short-looking report pass.
                    core = tests[next(iter(inputs.values()))['actual']].row
                    tests.row_dimensions[core].hidden = True
                    report.write_bytes(serialize_report(book))
                    with self.assertRaisesRegex(ValueError, 'Hidden report content'):
                        blocks.check(report, source, language)
                    from copy import copy
                    book.worksheets[0].sheet_view.selection.append(copy(book.worksheets[0].sheet_view.selection[0]))
                    report.write_bytes(serialize_report(book))
                    with self.assertRaisesRegex(ValueError,'Invalid native pane selections'):
                        blocks.check(report, source, language)
                    self.assertEqual(before, {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths})


if __name__ == '__main__':
    unittest.main()
