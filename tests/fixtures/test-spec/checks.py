"""Focused source/parser and one-workbook generation controls; no application execution."""
from pathlib import Path
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
from test_kit import load_registry, mapping, validate_markdown


def scenario_semantics(design):
    """Independent oracle for the eight synthetic source obligations, not app execution."""
    cases = {case['id']: case for case in design['cases']}
    def oracle(case_id, variant, checkpoint='CP-result'):
        rows = [cp for cp in cases[case_id]['checkpoints'] if cp[0] == variant and cp[1] == checkpoint]
        assert len(rows) == 1
        return rows[0][4]
    assert oracle('TC-SC-01', 'lt') == '29 is red; 30 and 31 are not red; red count 1'
    assert oracle('TC-SC-01', 'le') == '29 and 30 are red; 31 is not red; red count 2'
    assert oracle('TC-SC-02', 'cancel') == 'R1 remains; stored score 29 remains red'
    assert oracle('TC-SC-02', 'confirm') == 'R1 is removed; stored score 29 remains red before recalculation'
    assert oracle('TC-SC-02', 'confirm', 'CP-after') == 'After recalculation score 29 is not red; numeric score 29 remains'
    assert oracle('TC-SC-03', 'base', 'CP-saved') == 'R1 is retained with empty threshold and marked incomplete'
    assert oracle('TC-SC-03', 'base', 'CP-judgment') == 'Incomplete R1 is excluded; score 29 is not red'
    assert oracle('TC-SC-04', 'base') == 'Stale save rejected; R1 remains absent; numeric score 29 unchanged'
    assert cases['TC-SC-03']['execution_lane'] == 'Browser'
    assert cases['TC-SC-04']['execution_lane'] == 'Integration'
    assert oracle('TC-SC-05', 'positive') == 'Eligible: enabled=yes and 29<30 are both true'
    assert oracle('TC-SC-05', 'disabled') == 'Not eligible: enabled is false while 29<30 remains true'
    assert oracle('TC-SC-05', 'boundary') == 'Not eligible: 30<30 is false while enabled remains true'
    variants = {v[0]: v for v in cases['TC-SC-01']['variants']}
    assert variants['lt'][1:3] == ['T=30; S=29,30,31; comparator <', 'Step 1: select comparator <']
    assert variants['le'][1:3] == ['T=30; S=29,30,31; comparator ≤', 'Step 1: select comparator ≤']
    assert all(v[3] == '@Checkpoints' for c in cases.values() for v in c['variants'])
    assert len(cases['TC-SC-02']['steps']) == 5  # one shared sequence, not one duplicated sequence per decision
    assert all(c['steps'][-1][1].startswith('Restore TD-SC') for c in cases.values())
    assert len(design['coverage']) == 8
    assert {r[0] for r in design['coverage']} == {f'SYN:O{i}' for i in range(1, 9)}
    old_ids = {'TC-OLD-LT:base', 'TC-OLD-LE:base', 'TC-OLD-CANCEL:base', 'TC-OLD-CONFIRM:base',
               'TC-OLD-INCOMPLETE:base', 'TC-OLD-STALE:base', 'TC-OLD-AND-POSITIVE:base',
               'TC-OLD-AND-NEGATIVE:disabled', 'TC-OLD-AND-NEGATIVE:boundary'}
    for old in old_ids:
        assert any('old ' + old + ' → ' in r[3] for r in design['coverage'])
    # Missing live seams are preserved as gaps, not fake Ready or removed obligations.
    assert all(c['readiness'] == 'Draft' and c['gap'] == 'G-PREP' for c in cases.values())


def compact_semantics(design):
    """Overall/final outcomes retain the independent synthetic branch obligations."""
    assert design['source_version'] == '1.3.0'
    cases = {case['id']: case for case in design['cases']}
    assert len(cases) == 5 and sum(len(c['variants']) for c in cases.values()) == 9
    for case in cases.values():
        assert case['expected'] and not case['expected'].startswith('@')
        assert all(len(step) == 4 and step[2:] == ['', ''] for step in case['steps'])
        assert not any(step[1].startswith('Restore TD-SC and verify') for step in case['steps'])
        assert case['reset'] == 'Restore TD-SC initial rule/scores independently and verify its defined baseline.'
        assert case['readiness'] == 'Draft' and case['gap'] == 'G-PREP'
    branches = {v[0]: v[3] for c in cases.values() for v in c['variants'] if v[0] != 'base'}
    assert branches['lt'] == '29 is red; 30 and 31 are not red; red count 1'
    assert branches['le'] == '29 and 30 are red; 31 is not red; red count 2'
    assert branches['cancel'] == 'After recalculation R1 remains; score 29 remains red; numeric score 29 unchanged'
    assert branches['confirm'] == 'R1 removed; old red result remains before recalculation; afterwards score 29 is not red and numeric score 29 remains'
    assert branches['positive'] == 'Eligible: enabled=yes and 29<30 are both true'
    assert branches['disabled'] == 'Not eligible: enabled is false while 29<30 remains true'
    assert branches['boundary'] == 'Not eligible: 30<30 is false while enabled remains true'
    temporal = next(cp for cp in cases['TC-SC-02']['checkpoints'] if cp[:2] == ['confirm', 'CP-result'])
    assert temporal[4] == 'R1 is removed; stored score 29 remains red before recalculation'
    assert len(design['coverage']) == 8


def run(root: Path):
    from openpyxl import load_workbook
    scripts = root / 'skills/blend-generate-test-spec/scripts'
    sys.path.insert(0, str(scripts))
    import export_report as exporter
    import report_model as model
    from check_report import check_report
    registry = load_registry(root)
    for family in ('scope-and-approach', 'test-cases', 'test-data'):
        for language in ('ja', 'vi'):
            row = mapping(registry, family, language)
            validate_markdown((root / row['Template']).read_text(encoding='utf-8'), row)
    inputs = root / 'tests/fixtures/test-spec/inputs/valid'
    with tempfile.TemporaryDirectory(prefix='blend-test-spec-') as temporary:
        work = Path(temporary)
        for language in ('ja', 'vi'):
            source = work / language
            shutil.copytree(inputs, source)
            output = work / f'report.{language}.xlsx'
            result = exporter.export(source, output, language=language)
            assert result['cases'] == 3 and result['variants'] == 5
            book = load_workbook(output)
            locale = model.REPORT_LAYOUTS[language]
            sheet = book[locale['sheets'][1]]
            assert book.sheetnames[:2] == list(locale['sheets'][:2])
            from block_report import layout, formulas
            report_data=model.prepare_report(source,language)[0]
            cards,field_map=layout(report_data)
            assert list(field_map) == ['TC-SYN-01 / a', 'TC-SYN-01 / b', 'TC-SYN-01 / c', 'TC-SYN-02 / base', 'TC-SYN-03 / base']
            assert all(sheet[fields['status']].value == locale['statuses']['NOT RUN'] for fields in field_map.values())
            assert all(sheet[fields['actual']].value is None for fields in field_map.values())
            assert sheet.freeze_panes
            assert all(sheet[fields["actual"]].column == 4 and sheet[fields["status"]].column == 5 for fields in field_map.values())
            assert sheet.data_validations.dataValidation
            assert all(book.worksheets[index][cell].data_type == 'f' for index,cell in formulas(report_data))
            assert all(book.worksheets[0].cell(model.SUMMARY_FIELDS[k], 2).value is None for k in model.INPUT_FIELDS)
            before = output.read_bytes()
            checked = check_report(output, source, language)
            assert checked['variants'] == 5 and checked['evaluated'] == 0 and checked['pass_rate'] is None and checked['gaps']
            assert output.read_bytes() == before
            # A Windows legacy codepage must not break JA/VI CLI JSON or errors.
            env = {**os.environ, 'PYTHONIOENCODING': 'cp1252', 'PYTHONUTF8': '0'}
            cli_output = work / f'cli-{language}.xlsx'
            command = [sys.executable, str(scripts / 'export_report.py'), '--source-dir', str(source),
                       '--output', str(cli_output), '--language', language]
            completed = subprocess.run(command, capture_output=True, env=env, timeout=30)
            assert completed.returncode == 0, completed.stderr.decode('utf-8')
            transport = json.loads(completed.stdout.decode('utf-8'))
            assert transport['language'] == language and transport['sheets'][:2] == list(locale['sheets'][:2])
            assert not completed.stderr
            missing = work / '存在しない-chưa-có'
            command[command.index('--source-dir') + 1] = str(missing)
            command[command.index('--output') + 1] = str(work / f'absent-{language}.xlsx')
            failed = subprocess.run(command, capture_output=True, env=env, timeout=30)
            error = failed.stderr.decode('utf-8')
            assert failed.returncode == 1 and missing.name in error and 'UnicodeEncodeError' not in error
            try:
                exporter.export(source, output, language=language)
            except FileExistsError:
                pass
            else:
                raise AssertionError('Existing run overwritten')
            assert output.read_bytes() == before
        def replace(folder, filename, old, new):
            path = folder / filename
            text = path.read_text(encoding='utf-8')
            assert old in text
            path.write_text(text.replace(old, new), encoding='utf-8')
        def blocked(name, mutate, language='vi'):
            source = work / name
            shutil.copytree(inputs, source)
            mutate(source)
            output = work / f'{name}.xlsx'
            try:
                exporter.export(source, output, language=language)
            except ValueError:
                assert not output.exists()
            else:
                raise AssertionError('Invalid source accepted: ' + name)
        blocked("wrong-version", lambda p: replace(p, "test-cases.vi.md", "test-cases@1.0.0", "test-cases@2.0.0"))
        blocked("different-revision", lambda p: replace(p, "test-data.vi.md", "syn-r2", "syn-r3"))
        blocked("raw-pipe", lambda p: replace(p, "test-cases.vi.md", "enabled=yes, score=59, max=100", "enabled=yes | score=59"))
        blocked("compound-variant", lambda p: replace(p, "test-cases.vi.md", "| a |", "| a/b |"))
        blocked("duplicate-case", lambda p: replace(p, "test-cases.vi.md", "#### TC-SYN-02", "#### TC-SYN-01"))
        blocked("casefold-variant-collision", lambda p: replace(p, "test-cases.vi.md", "| b |", "| A |"))
        blocked("casefold-case-collision", lambda p: replace(p, "test-cases.vi.md", "#### TC-SYN-02", "#### tc-syn-01"))
        blocked("unknown-fixture", lambda p: replace(p, "test-cases.vi.md", "| Fixture | TD-01 |", "| Fixture | TD-99 |"))
        blocked("missing-context", lambda p: replace(p, "test-cases.vi.md", "@CTX-1", "@CTX-99"))
        blocked("unknown-gap", lambda p: replace(p, "test-cases.vi.md", "| Gap | G-02 |", "| Gap | G-99 |"))
        blocked("unknown-oracle-ready", lambda p: replace(p, "test-cases.vi.md", "| Căn cứ kỳ vọng | Confirmed |", "| Căn cứ kỳ vọng | Awaiting decision |"))
        blocked("missing-step", lambda p: replace(p, "test-cases.vi.md", "##### Steps", "##### Unregistered Steps"))
        blocked("unprojected-case-prose", lambda p: replace(p, "test-cases.vi.md", "#### TC-SYN-02", "Important extra unparsed assertion\n\n#### TC-SYN-02"))
        blocked("ja-variant-drift", lambda p: replace(p, "test-cases.ja.md", "| b |", "| bb |"))
        blocked("missing-required-field", lambda p: replace(p, "test-cases.vi.md", "| Bảo toàn | Non-target B unchanged |\n", ""))
        blocked("duplicate-gap-ref", lambda p: replace(p, "test-cases.vi.md", "| Gap | G-02 |", "| Gap | G-02, G-02 |"))
        blocked("undefined-multiple-gap-ref", lambda p: replace(p, "test-cases.vi.md", "| Gap | G-02 |", "| Gap | G-02, G-99 |"))
        blocked("malformed-gap-separator", lambda p: replace(p, "test-cases.vi.md", "| Gap | G-02 |", "| Gap | G-02; G-01 |"))
        blocked("gap-prose-annotation", lambda p: replace(p, "test-cases.vi.md", "| Gap | G-02 |", "| Gap | G-02 (fixture) |"))
        blocked("gap-none-mixed", lambda p: replace(p, "test-cases.vi.md", "| Gap | G-02 |", "| Gap | none, G-02 |"))
        blocked("ready-with-gap-list", lambda p: replace(p, "test-cases.vi.md", "| Gap | none |", "| Gap | G-01, G-02 |"))
        blocked("gap-kind-list", lambda p: replace(p, "scope-and-approach.vi.md", "| G-02 | preparation |", "| G-02 | preparation, proof |"))
        blocked("duplicate-coverage-ref", lambda p: replace(p, "scope-and-approach.vi.md", "| TC-SYN-01:a, TC-SYN-01:b, TC-SYN-01:c |", "| TC-SYN-01:a, TC-SYN-01:a |"))
        blocked("coverage-semicolon", lambda p: replace(p, "scope-and-approach.vi.md", "| TC-SYN-03 |", "| TC-SYN-03; G-02 |"))
        blocked("coverage-prose", lambda p: replace(p, "scope-and-approach.vi.md", "| TC-SYN-03 |", "| TC-SYN-03 (Step 4), G-02 |"))
        # Large source remains readable via detail chunks; no row-height rejection/truncation.
        long_source = work / 'long'
        shutil.copytree(inputs, long_source)
        long_text = 'preserved state ' * 250
        for language in ('ja', 'vi'):
            replace(long_source, f'test-cases.{language}.md', '| a | enabled=yes, score=59, max=100 | Eligible |', '| a | enabled=yes, score=59, max=100 | ' + long_text + ' |')
        long_output = work / 'long.xlsx'
        exporter.export(long_source, long_output)
        long_book = load_workbook(long_output)
        assert len(long_book.sheetnames) == 2
        chunks = ''.join(str(row[1].value or '') for row in long_book.worksheets[1])
        long_data,_ = model.prepare_report(long_source)
        assert next(row for row in long_data.rows if row.identity == 'TC-SYN-01 / a').expected.count('preserved state') == 250
        assert chunks.casefold().count('preserved state') >= 250
        check_report(long_output, long_source)
        ja_source = root / 'tests/fixtures/test-spec/ja-quality-inputs'
        ja_output = work / 'ja-quality.xlsx'
        result = exporter.export(ja_source, ja_output, language='ja')
        assert result['cases'] == 1 and result['variants'] == 3
        check_report(ja_output, ja_source, 'ja')
        dummy_source = work / 'authorized-dummy-input'
        shutil.copytree(inputs, dummy_source)
        for language in ('ja', 'vi'):
            replace(dummy_source, f'test-cases.{language}.md', 'local: import preparation unavailable — G-02',
                    'local: [dummy-input: http://127.0.0.1/]')
        dummy_output = work / 'dummy.xlsx'
        exporter.export(dummy_source, dummy_output)
        check_report(dummy_output, dummy_source)
        # A realistic 40-case report retains readable formulas within Excel limits.
        from dataclasses import replace as replace_record
        data, _ = model.prepare_report(inputs, 'vi')
        rows = tuple(replace_record(data.rows[0], case_id=f'TC-SCALE-{i:03}', variant='base') for i in range(40))
        from copy import deepcopy
        scenario_source = root / 'tests/fixtures/test-spec/inputs/scenario-grouping'
        for language in ('vi', 'ja'):
            raw_rows = [line.strip('|').split('|') for line in
                        (scenario_source / f'source.{language}.md').read_text(encoding='utf-8').splitlines()
                        if line.startswith('| O')]
            assert {row[0].strip() for row in raw_rows} == {f'O{i}' for i in range(1, 9)}
            assert raw_rows[0][3].strip() == 'red/not red/not red'
            assert raw_rows[1][3].strip() == 'red/red/not red'
            assert raw_rows[-1][1].strip() == 'TC-OLD-AND-NEGATIVE:disabled, TC-OLD-AND-NEGATIVE:boundary'
            design = exporter.parse_sources(scenario_source, language)
            scenario_semantics(design)
            # The same schema can parse a wrong oracle; the independent behavioral
            # check must discriminate it rather than counting tables/keywords.
            wrong = deepcopy(design)
            cp = next(cp for cp in wrong['cases'][0]['checkpoints'] if cp[:2] == ['lt', 'CP-result'])
            cp[4] = '29 and 30 are red; 31 is not red; red count 2'
            try:
                scenario_semantics(wrong)
            except AssertionError:
                pass
            else:
                raise AssertionError('Less-than oracle inherited less-or-equal outcome')
            wrong = deepcopy(design)
            wrong['cases'][-1]['variants'] = [v for v in wrong['cases'][-1]['variants'] if v[0] != 'boundary']
            wrong['cases'][-1]['checkpoints'] = [cp for cp in wrong['cases'][-1]['checkpoints'] if cp[0] != 'boundary']
            try:
                scenario_semantics(wrong)
            except AssertionError:
                pass
            else:
                raise AssertionError('AND equality negative control dropped')
            projected, _ = model.prepare_report(scenario_source, language)
            assert projected.report_version == '2.4.0' and projected.feature_id == 'SYN-SC'
            assert len(projected.rows) == 9 and not any(row.eligible for row in projected.rows)
            lt = next(r for r in projected.rows if r.identity == 'TC-SC-01 / lt')
            le = next(r for r in projected.rows if r.identity == 'TC-SC-01 / le')
            assert 'red count 1' in lt.expected and 'red count 2' not in lt.expected
            assert 'red count 2' in le.expected and 'red count 1' not in le.expected
        # Parser failures must not silently discard checkpoint meaning or identity.
        for name, old, new in (
            ('missing-checkpoints', '##### Checkpoints', '##### Unregistered checkpoints'),
            ('invalid-checkpoint-stage', '| CP-result | 2 | result |', '| CP-result | 2 | unknown |'),
            ('invalid-checkpoint-variant', '| lt | CP-result |', '| other | CP-result |'),
            ('nonconcrete-checkpoint', '29 is red; 30 and 31 are not red; red count 1', '@Checkpoints'),
            ('casefold-checkpoint-collision', '| lt | CP-result |', '| lt | cp-SETUP |'),
        ):
            invalid = work / name
            shutil.copytree(scenario_source, invalid)
            replace(invalid, 'test-cases.vi.md', old, new)
            try:
                exporter.parse_sources(invalid, 'vi')
            except ValueError:
                pass
            else:
                raise AssertionError('Invalid scenario accepted: ' + name)
        compact_source = root / 'tests/fixtures/test-spec/inputs/compact-scenario'
        from test_kit import _gate
        check_output = _gate.check_output
        for language in ('vi', 'ja'):
            design = exporter.parse_sources(compact_source, language)
            compact_semantics(design)
            assert design['cases'][0]['context_refs']['actor'] == '@CTX-SC'
            assert len(design['contexts']['CTX-SC']) == 5
            projected, _ = model.prepare_report(compact_source, language)
            assert projected.report_version == '2.4.0' and projected.feature_id == 'SYN-SC'
            final_expected = {(case['id'], variant[0]): variant[3]
                              for case in design['cases'] for variant in case['variants']}
            assert all(row.expected.casefold() == final_expected[(row.case_id, row.variant)].casefold()
                       for row in projected.rows)
            assert not any(row.eligible for row in projected.rows)
            for family in ('scope-and-approach', 'test-cases', 'test-data'):
                filename = f'{family}.{language}.md'
                check_output((compact_source / filename).read_text(encoding='utf-8'),
                             filename, registry, family, language, root)
            legacy_name = f'test-cases.{language}.md'
            previous_source = (scenario_source / legacy_name).read_text(encoding='utf-8')
            previous_schema = _gate.source_mapping(previous_source, registry, 'test-cases', language)
            _gate.check_identity(previous_source, previous_schema)
        for name, old, new in (
            ('compact-missing-overall', '| Expected | The comparator selects the correct red scores at the 30 boundary; numeric scores and non-target 82 stay unchanged. |', '| Expected | @Checkpoints |'),
            ('compact-missing-final', '| lt | T=30; S=29,30,31; comparator < | Step 1: select comparator < | 29 is red; 30 and 31 are not red; red count 1 |', '| lt | T=30; S=29,30,31; comparator < | Step 1: select comparator < | @Checkpoints |'),
            ('compact-unresolved-final', '29 is red; 30 and 31 are not red; red count 1', '[concrete final outcome]'),
            ('compact-aliased-final', '29 is red; 30 and 31 are not red; red count 1', 'same as le'),
            ('compact-cross-branch', '| lt | T=30; S=29,30,31; comparator < | Step 1: select comparator < | 29 is red; 30 and 31 are not red; red count 1 |', '| lt | T=30; S=29,30,31; comparator < | Step 1: select comparator < | le: 29 and 30 are red; 31 is not red |'),
            ('compact-old-step-columns', '| Step | Action |', '| Step | Action | Expected | Preservation |'),
            ('compact-unknown-delta', 'Step 1: select comparator <', 'Step 99: select comparator <'),
            ('compact-missing-checkpoints', '##### Checkpoints', '##### Omitted checkpoint proof'),
        ):
            invalid = work / name
            shutil.copytree(compact_source, invalid)
            replace(invalid, 'test-cases.vi.md', old, new)
            try:
                exporter.parse_sources(invalid, 'vi')
            except ValueError:
                pass
            else:
                raise AssertionError('Invalid compact source accepted: ' + name)
        scale_cases=[]
        for record in rows:
            case=deepcopy(data.cases[0])
            case['id']=record.case_id
            case['variants']={'base':next(iter(case['variants'].values()))}
            scale_cases.append(case)
        scaled = model.ReportData('vi', data.summary, rows,tuple(scale_cases))
        from block_report import formulas as block_formulas
        scale_formulas = block_formulas(scaled)
        assert len([cell for (index,cell) in scale_formulas if index==0 and cell.startswith('C') and int(cell[1:]) >= 30]) == 40
        assert max(map(len, scale_formulas.values())) <= 8192
        from render_report import render_report
        scale_output = work / 'forty-cases.xlsx'
        assert render_report(scaled, scale_output)['rows'] == 40
        for colliding in ((replace_record(rows[0], case_id='Case', variant='a'), replace_record(rows[0], case_id='case', variant='b')),
                          (replace_record(rows[0], case_id='Case', variant='a'), replace_record(rows[0], case_id='Case', variant='A'))):
            try:
                model.ReportData('vi', data.summary, colliding)
            except ValueError as error:
                assert 'case-insensitive' in str(error)
            else:
                raise AssertionError('Projected identity collision accepted')
        # Source changes after projection cannot publish a workbook against stale bytes.
        import render_report
        race_source = work / 'source-race'
        shutil.copytree(inputs, race_source)
        original_render = render_report.render_report
        def changed_source(data, output):
            result = original_render(data, output)
            path = race_source / 'test-cases.vi.md'
            path.write_bytes(path.read_bytes() + b'\n')
            return result
        render_report.render_report = changed_source
        try:
            race_output = work / 'source-race.xlsx'
            try:
                exporter.export(race_source, race_output)
            except ValueError as error:
                assert 'changed during operation' in str(error)
            else:
                raise AssertionError('Source capture race accepted')
            assert not race_output.exists()
        finally:
            render_report.render_report = original_render
    return ['JA/VI one-workbook generation and blank editable inputs',
            'source parser malformed identity/context/variant/gap/coverage controls retained',
            'read-only same-file check, no overwrite, formulas and localized dropdowns',
            'long details and Japanese-only source preservation; no native execution proof',
            'scenario grouping retains eight source obligations, own branch oracles, staged checkpoints and old-ID mapping',
            'compact1.3 concrete overall/final outcomes and two-column steps; byte-preserved1.2 grammar/gates',
            'JA/VI CLI stdout/stderr are UTF-8 under forced Windows cp1252 without global settings']
