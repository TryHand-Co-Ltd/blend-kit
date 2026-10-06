"""Meaningful v2 same-file controls using frozen synthetic inputs, never feature QA."""
from pathlib import Path
import importlib
import tempfile
import shutil
import sys


def run(root):
    sys.path.insert(0,str(root/'skills/blend-generate-test-spec/scripts'))
    model=importlib.import_module('report_model')
    blocks=importlib.import_module('block_report')
    from render_report import serialize_report
    from openpyxl import load_workbook
    checks=[]
    def refused(call):
        try:
            call()
        except (ValueError,TypeError,OSError):
            return
        raise AssertionError('Invalid report accepted')
    with tempfile.TemporaryDirectory(prefix='blend-block-controls-') as directory:
        base=Path(directory)
        for lang in ('ja','vi'):
            folder=base/lang
            source=folder/'design'
            shutil.copytree(root/'tests/fixtures/test-spec/inputs/valid',source)
            # Both source languages declare the same synthetic URL, never a BLEND route.
            for language in ('ja','vi'):
                case_path=source/f'test-cases.{language}.md'
                text=case_path.read_text(encoding='utf-8')
                label='Chức năng' if language=='vi' else '機能'
                import re
                text=re.sub(r'^(\| '+label+r' \| .*? \|)$',r'\1\n| screen_relative_path | /synthetic/preview |',text,flags=re.M)
                case_path.write_text(text,encoding='utf-8')
            report=folder/'report.xlsx'
            receipt=model.working.export(source,report,language=lang)
            assert receipt['version']=='2.5.0'
            data,_=model.prepare_report(source,lang)
            cards,inputs=blocks.layout(data)
            report_formulas=blocks.formulas(data,cards,inputs)
            assert report_formulas[(0,'B18')].startswith('=COUNTIF(')
            eligible=next(record for record in data.rows if record.eligible)
            eligible_card=next(card for card in cards if card['id']==eligible.case_id)
            overview_row=model.CASE_START_ROW+cards.index(eligible_card)
            case_formula=report_formulas[(0,f'C{overview_row}')]
            if case_formula == f"='{model.REPORT_LAYOUTS[lang]['sheets'][1]}'!B{eligible_card['status_row']}":
                case_formula=report_formulas[(1,f'B{eligible_card["status_row"]}')]
            assert inputs[eligible.identity]['actual'] not in case_formula
            book=load_workbook(report)
            assert len(book.sheetnames)==2 and not any(sheet._images for sheet in book)
            visible='\n'.join(str(cell.value or '') for sheet in book for row in sheet for cell in row)
            assert '[TC-ID]' not in visible and '[テストケース名]' not in visible and 'không phải testcase thật' not in visible
            assert not model.markdown_leak(visible)
            assert not re.search(r'(?m)^•\s*[-+*•]\s+',visible)
            tests=book.worksheets[1]
            for fields in inputs.values():
                assert tests[fields['status']].value==model.REPORT_LAYOUTS[lang]['statuses']['NOT RUN']
                assert tests[fields['actual']].value is None
                assert tests[fields['status']].coordinate.startswith('E')
            first=report.read_bytes()
            result=blocks.check(report,source,lang)
            assert result['states']=={'NOT RUN':len(data.rows)} and report.read_bytes()==first
            refused(lambda:model.working.export(source,report,language=lang))
            delivery={'closure_confirmation':{'closed_by':'Synthetic QA','closed_at':'2026-10-04T10:00:00+00:00','source':'Synthetic closure','audience':'Synthetic recipient'},'evidence_access':[{'url':'https://evidence.example.com/synthetic/run','audience':'Synthetic recipient','verified_by':'Synthetic QA','verified_at':'2026-10-04T10:00:00+00:00','method':'human-attestation','source':'Synthetic access'}],'screenshots':[]}
            for key in model.INPUT_FIELDS:
                book.worksheets[0][f'B{model.SUMMARY_FIELDS[key]}']=('2026-10-04T09:00:00+00:00' if key=='period' else 'Synthetic '+key)
            states=model.REPORT_LAYOUTS[lang]['statuses']
            for record in data.rows:
                fields=inputs[record.identity]
                tests[fields['status']]=states['PASS'] if record.eligible else states['SKIPPED']
                tests[fields['actual']]='Synthetic observed result'
            report.write_bytes(serialize_report(book))
            frozen=report.read_bytes()
            complete=blocks.check(report,source,lang,'complete',**delivery)
            assert complete['complete'] and complete['passed']==sum(r.eligible for r in data.rows)
            assert report.read_bytes()==frozen
            record=next(r for r in data.rows if r.eligible)
            fields=inputs[record.identity]
            for field,value in (('status','Invalid pasted state'),('actual','=1'),('actual','\u00a0\u2003')):
                mutated=load_workbook(report)
                mutated.worksheets[1][fields[field]]=value
                if field=='actual' and value=='=1':
                    mutated.worksheets[1][fields[field]].data_type='f'
                path=folder/'negative.xlsx'
                path.write_bytes(serialize_report(mutated))
                if field=='status' or value=='=1':
                    refused(lambda:blocks.check(path,source,lang,'complete',**delivery))
                elif value!='\u00a0\u2003':
                    assert not blocks.check(path,source,lang,'complete',**delivery)['complete']
                else:
                    assert not blocks.check(path,source,lang,'complete',**delivery)['complete']
            mutated=load_workbook(report)
            mutated.worksheets[1].cell(cards[0]['start'],1,'Changed design title')
            path=folder/'changed-source.xlsx'
            path.write_bytes(serialize_report(mutated))
            refused(lambda:blocks.check(path,source,lang))
            mutated=load_workbook(report)
            mutated.worksheets[1].row_dimensions[cards[0]['start']].hidden=True
            path.write_bytes(serialize_report(mutated))
            refused(lambda:blocks.check(path,source,lang))
            for url in ('https://example.com/screen','//example.com/screen','/../private','C:/private'):
                refused(lambda:model.screen_path(url,lang))
            checks.append(lang+'-two-sheets-title-status-path-blank-inputs-readonly-nooverwrite-completion-negative-controls')
    return checks
