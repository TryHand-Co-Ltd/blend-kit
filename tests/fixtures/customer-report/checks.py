"""Focused same-file editable-report controls, with synthetic input only."""
from pathlib import Path
import importlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET


def load_domain(root):
    scripts = root / 'skills/blend-generate-test-spec/scripts'
    if str(scripts) not in sys.path:
        sys.path.insert(0, str(scripts))
    return importlib.import_module('report_model')


def run(root, scope='all'):
    spec = importlib.util.spec_from_file_location('block_controls',Path(__file__).with_name('block_checks.py'))
    controls = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(controls)
    return controls.run(root) + scenario_controls(root) + scenario_controls(root,compact=True)


def scenario_controls(root, sample_dir=None, *, compact=False):
    """Synthetic metadata/image layout controls; no assertion of a real pixel review."""
    model=load_domain(root)
    from update_report import update_report, report_bindings, _validator
    from render_report import serialize_report
    from openpyxl import load_workbook
    from PIL import Image as PILImage, ImageDraw
    import hashlib
    checks=[]
    helper_module=sys.modules.get('run_artifacts')
    if helper_module is None:
        helper_path=root/'skills/blend-automation-test/scripts/run_artifacts.py'
        spec=importlib.util.spec_from_file_location('run_artifacts',helper_path)
        helper_module=importlib.util.module_from_spec(spec)
        sys.modules['run_artifacts']=helper_module
        spec.loader.exec_module(helper_module)
    helper=helper_module
    def refused(call):
        try:
            call()
        except (ValueError,TypeError,OSError):
            return
        raise AssertionError('Unsafe scenario writeback accepted')
    with tempfile.TemporaryDirectory(prefix='blend-scenario-report-') as directory:
        base=Path(directory)
        for language in ('vi','ja'):
            work=base/language
            work.mkdir()
            source=work/'design'
            shutil.copytree(root/'tests/fixtures/test-spec/inputs'/('compact-scenario' if compact else 'scenario-grouping'),source)
            report=work/'report.xlsx'
            receipt=model.working.export(source,report,language=language)
            assert receipt['version']=='2.5.0'
            data,_=model.prepare_report(source,language)
            assert data.feature_id=='SYN-SC'
            first=next(r for r in data.rows if r.case_id=='TC-SC-01' and r.variant=='lt')
            assert 'red count 1' in first.expected and 'red count 2' not in first.expected
            binding=report_bindings(source,report,language)
            fields=binding['variants']['TC-SC-01 / lt']
            book=load_workbook(report)
            assert len(book.sheetnames)==2
            assert len(fields['checkpoints']['CP-result']['picture_rows'])==__import__('block_report').READABLE_EVIDENCE_ROWS
            immutable_formula=book.worksheets[0]['B18'].value
            ctx=work/'context'
            feature=ctx/'features/SYN-SC-synthetic'
            feature.mkdir(parents=True)
            (feature/'README.md').write_text('Synthetic fixture, not BLEND authority',encoding='utf-8')
            (feature/'CONTEXT.md').write_text('Synthetic fixture, not BLEND authority',encoding='utf-8')
            run=helper.create_run(ctx,'SYN-SC-synthetic',{'design_revision':data.summary['revision'],
                'feature_id':'SYN-SC','report':str(report)},'2026-10-05-001')
            identity={'design_revision':data.summary['revision'],'feature_id':'SYN-SC',
                'case_id':'TC-SC-01','variant_id':'lt','run_id':run.name}
            evidence=[]
            for sequence,height in ((1,5000),(2,1080),(3,1080)):
                cp='CP-result'
                item={**identity,'checkpoint_id':cp}
                note=f'Synthetic structural image {sequence}; no runtime or pixel-review proof.'
                image=PILImage.new('RGB',(1920,height),'white')
                ImageDraw.Draw(image).rectangle((20,60,1890,height-20),outline='red',width=5)
                ImageDraw.Draw(image).text((20,20),note,fill='black')
                path=work/f'image-{sequence}.png'
                image.save(path)
                receipts=[]
                for kind in ('raw','annotated'):
                    receipts.append(helper.archive_capture(run,path,item,kind,sequence,{'full_page':True,
                        'viewport':{'width':1920,'height':1080},'observed_state':'synthetic layout only',
                        'annotation_present':kind=='annotated','captured_at':'2026-10-05T10:00:00+07:00',
                        'capture_reference':'synthetic structural fixture '+kind+str(sequence)}))
                raw,annotated=receipts
                evidence.append({**item,'sequence':sequence,'raw_path':raw['path'],
                    'annotated_path':annotated['path'],'raw_sha256':raw['sha256'],
                    'annotated_sha256':annotated['sha256'],'viewport':{'width':1920,'height':1080},
                    'full_page':True,'assertion':'Synthetic layout control','focus':'Synthetic rectangle',
                    'observed_note':note,'reviewed_by':'Synthetic metadata control, not a human reviewer',
                    'reviewed_at':'2026-10-05T10:01:00+07:00','source':{'capture_tool':'synthetic fixture',
                        'raw_capture':raw['capture_reference'],'annotated_capture':annotated['capture_reference'],
                        'annotation_method':'synthetic image generation','pixel_review':{
                            'tool':'synthetic metadata test only','reference':'not an actual viewer call; runtime review remains separate'}}})
            payload={'identity':identity,'status':'FAIL','actual':'=Literal observed text; synthetic failure reason; next action: verify runtime.',
                'evidence':evidence}
            result=update_report(source,report,payload,language=language,run_dir=run)
            assert result['readback_verified'] and result['evidence_added']==3
            book=load_workbook(report)
            tests=book.worksheets[1]
            assert tests[fields['actual_cell']].data_type=='s' and tests[fields['actual_cell']].value.startswith('=')
            assert len(tests._images)==3 and book.worksheets[0]['B18'].value==immutable_formula
            assert max(d.height or 15 for d in tests.row_dimensions.values())<=409
            cards,_ = __import__('block_report').layout(data)
            following=next(card for card in cards if card['id']!=identity['case_id'])
            assert all(image.anchor._from.row+1<following['start'] for image in tests._images)
            assert all(image.anchor.pic.nvPicPr.cNvPr.descr.startswith('Synthetic structural') for image in tests._images)
            original=report.read_bytes()
            refused(lambda:update_report(source,report,{**payload,'identity':{**identity,'design_revision':'wrong'}},language=language,run_dir=run))
            refused(lambda:update_report(source,report,{**payload,'status':'PASS'},language=language,run_dir=run))
            refused(lambda:update_report(source,report,{**payload,'evidence':[{**evidence[0],'full_page':False}]},language=language,run_dir=run))
            refused(lambda:update_report(source,report,{**payload,'evidence':[{**evidence[0],'annotated_sha256':'0'*64}]},language=language,run_dir=run))
            assert report.read_bytes()==original
            ledger=run/'run-summary.md'
            ledger_before=ledger.read_text(encoding='utf-8')
            ledger.write_text(ledger_before.replace(json.dumps(str(report)),json.dumps(str(work/'another-authoritative-report.xlsx'))),encoding='utf-8')
            try:
                refused(lambda:update_report(source,report,payload,language=language,run_dir=run))
                assert report.read_bytes()==original
            finally:
                ledger.write_text(ledger_before,encoding='utf-8')
            second={**identity,'variant_id':'le'}
            update_report(source,report,{'identity':second,'status':'BLOCKED',
                'actual':'CP-setup: synthetic prerequisite unavailable; next action: prepare fixture.'},language=language)
            reread=load_workbook(report)
            assert len(reread.worksheets[1]._images)==3
            assert reread.worksheets[1][fields['actual_cell']].value==payload['actual']
            assert reread.worksheets[1][binding['variants']['TC-SC-01 / le']['status_cell']].value==model.REPORT_LAYOUTS[language]['statuses']['BLOCKED']
            mutations=work/'mutation.xlsx'
            moved=load_workbook(report)
            moved.worksheets[1]._images[0].anchor._from.row=fields['checkpoints']['CP-setup']['picture_rows'][0]-1
            mutations.write_bytes(serialize_report(moved))
            refused(lambda:__import__('block_report').check(mutations,source,language))
            overlapping=load_workbook(report)
            overlapping.worksheets[1]._images[1].anchor._from=__import__('copy').copy(overlapping.worksheets[1]._images[0].anchor._from)
            mutations.write_bytes(serialize_report(overlapping))
            refused(lambda:__import__('block_report').check(mutations,source,language))
            tamper=load_workbook(report)
            tamper.worksheets[1].data_validations.dataValidation[0].formula1='"Invalid"'
            mutations.write_bytes(serialize_report(tamper))
            refused(lambda:__import__('block_report').check(mutations,source,language))
            # Identical accepted evidence is idempotent, while checkpoint history survives.
            again=update_report(source,report,payload,language=language,run_dir=run)
            assert again['evidence_added']==0
            oversized=PILImage.new('RGB',(200,30000),'white')
            oversize_path=work/'oversized.png'
            oversized.save(oversize_path)
            cp_identity={**identity,'checkpoint_id':'CP-result'}
            tall_receipts=[]
            for kind in ('raw','annotated'):
                tall_receipts.append(helper.archive_capture(run,oversize_path,cp_identity,kind,4,
                    {'full_page':True,'viewport':{'width':200,'height':1080},'observed_state':'synthetic tall capacity control',
                     'annotation_present':kind=='annotated','captured_at':'2026-10-05T10:00:00+07:00',
                     'capture_reference':'synthetic tall '+kind}))
            tall={**evidence[0],'viewport':{'width':200,'height':1080},'sequence':4,'raw_path':tall_receipts[0]['path'],'annotated_path':tall_receipts[1]['path'],
                'raw_sha256':tall_receipts[0]['sha256'],'annotated_sha256':tall_receipts[1]['sha256'],
                'source':{**evidence[0]['source'],'raw_capture':'synthetic tall raw','annotated_capture':'synthetic tall annotated'}}
            original=report.read_bytes()
            refused(lambda:update_report(source,report,{**payload,'evidence':[tall]},language=language,run_dir=run))
            assert report.read_bytes()==original
            # Simulate a user edit between capture and atomic replacement.
            import update_report as writer
            original_serializer=writer.serialize_report
            def concurrent_change(book):
                result=original_serializer(book)
                report.write_bytes(report.read_bytes()+b'Synthetic concurrent user edit')
                return result
            writer.serialize_report=concurrent_change
            try:
                refused(lambda:update_report(source,report,{**payload,'evidence':[]},language=language))
                assert report.read_bytes().endswith(b'Synthetic concurrent user edit')
            finally:
                writer.serialize_report=original_serializer
                report.write_bytes(original)
            if sample_dir is not None:
                sample_dir=Path(sample_dir)
                sample_dir.mkdir(parents=True,exist_ok=True)
                with (sample_dir/f'synthetic-scenario-{receipt["version"][:3]}.{language}.xlsx').open('xb') as handle:
                    handle.write(report.read_bytes())
            checks.append(language+'-scenario-'+receipt['version']+'-branch-oracle-literal-write-three-tall-images-preservation-binding-overlap-validation-controls')
    return checks
