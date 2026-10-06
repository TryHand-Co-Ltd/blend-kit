"""Focused package safety/closure checks; no actual agent execution claim."""
from __future__ import annotations

import contextlib
import hashlib
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from test_kit import SKILLS, _gate, mapping, validate_output_filename


def run(root: Path) -> list[str]:
    path = root / "scripts/build-kit.py"
    spec = importlib.util.spec_from_file_location("blend_kit_builder", path)
    if spec is None or spec.loader is None:
        raise ValueError("Package builder could not be loaded")
    builder = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(builder)
    controls = json.loads((root / "tests/fixtures/package/inputs/controls.json").read_text())
    assert builder.SKILLS == SKILLS
    assert SKILLS[:-1] == tuple(controls["skills"]), "Historical five-skill controls changed"
    assert SKILLS[-1] == "blend-automation-test"
    current_resources = {f"{builder.REPORT_OWNER}/assets/test-report-block-template.{language}.xlsx":
                         (root / builder.REPORT_OWNER / "assets" / f"test-report-block-template.{language}.xlsx").read_bytes()
                         for language in ("ja", "vi")}
    assert builder.PRIVATE.search('host in ("localhost",)') is None
    for locator in ("http://localhost/private", "https://localhost/path", "localhost:3000"):
        assert builder.PRIVATE.search(locator), "Local locator passed package privacy check"

    def rejected(action, label: str) -> None:
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                action()
        except (ValueError, OSError):
            return
        raise AssertionError(f"Invalid package accepted: {label}")

    with tempfile.TemporaryDirectory(prefix="blend kit package ") as temporary:
        base = Path(temporary)
        source = base / controls["relocated_source"]
        source.mkdir(parents=True)
        for directory in ("skills", "shared", "assets"):
            shutil.copytree(root / directory, source / directory,
                            ignore=shutil.ignore_patterns("__pycache__", "*.pyc", "~$*"))
        source_script = source / "scripts/build-kit.py"
        source_script.parent.mkdir()
        shutil.copyfile(path, source_script)
        # The source may contain private/workflow material elsewhere; it must never ship.
        for relative in controls["excluded_source_files"]:
            forbidden = source / relative
            forbidden.parent.mkdir(parents=True, exist_ok=True)
            forbidden.write_text("excluded source material")
        # Office owner files are ephemeral/private: exclude before reading or copying.
        owner_file = source / "skills/blend-generate-test-spec/assets/~$test-report-block-template.vi.xlsx"
        owner_bytes = b"synthetic Office owner metadata; deliberately not an XLSX"
        owner_file.write_bytes(owner_bytes)
        output = base / controls["relocated_output"]
        completed = subprocess.run([sys.executable, str(source_script), "--source", str(source),
                                    "--output", str(output), "--profiles", *controls["profiles"], "--check"],
                                   cwd=base, capture_output=True, text=True)
        if completed.returncode:
            raise AssertionError(f"Relocated CLI failed: {completed.stdout}{completed.stderr}")
        rows = builder.read_registry((source / "shared/artifact-formats.md").read_text(encoding="utf-8"))
        for language in ("ja", "vi"):
            current = mapping(rows, "test-cases", language)
            assert current["Version"] == "1.3.0"
            for version in ("1.0.0", "1.1.0", "1.2.0", "1.3.0"):
                marker = f"<!-- blend-template: test-cases@{version} -->"
                dispatched = _gate.source_mapping(marker, rows, "test-cases", language)
                assert dispatched["Version"] == version
                _gate.check_identity(marker, dispatched)
            fenced = "```md\n<!-- blend-template: test-cases@1.1.0 -->\n```\n<!-- blend-template: test-cases@1.3.0 -->"
            assert _gate.source_mapping(fenced, rows, "test-cases", language) == current
            duplicate = "<!-- blend-template: test-cases@1.1.0 -->\n<!-- blend-template: test-cases@1.3.0 -->"
            rejected(lambda: _gate.check_identity(duplicate, _gate.source_mapping(duplicate, rows, "test-cases", language)),
                     "duplicate source identity dispatch")
        assert {(row["Output type"], row["Language"]) for row in rows} == {
            (row["Output type"], row["Language"]) for row in builder.read_registry(
                (root / "shared/artifact-formats.md").read_text(encoding="utf-8"))}
        ja_workbook = mapping(rows, "test-report", "ja")
        vi_workbook = mapping(rows, "test-report", "vi")
        rejected(lambda: builder.validate_template((source / ja_workbook["Template"]).read_bytes(), vi_workbook),
                 "JA workbook asset mapped as VI")
        rejected(lambda: builder.validate_template((source / vi_workbook["Template"]).read_bytes(), ja_workbook),
                 "VI workbook asset mapped as JA")
        registry_text = (source / "shared/artifact-formats.md").read_text(encoding="utf-8")
        rejected(lambda: builder.read_registry("\n".join(line for line in registry_text.splitlines()
                 if not (line.startswith("| test-report |") and "| ja |" in line))), "missing JA workbook pair")
        for language in ("ja", "vi"):
            report_row = mapping(rows, "test-report", language)
            report_raw = (source / report_row["Template"]).read_bytes()
            builder.validate_template(report_raw, report_row)
            for legacy_family in ("test-case-report", "customer-test-report"):
                rejected(lambda: builder.validate_template(report_raw, {**report_row, "Family": legacy_family}),
                         "legacy family accepted as active report")
            other = "vi" if language == "ja" else "ja"
            rejected(lambda: builder.validate_template(report_raw, mapping(rows, "test-report", other)),
                     "report language swapped")
        source_gate = source / "shared/scripts/artifact_gate.py"
        saved = base / "saved artifacts with spaces"
        saved.mkdir()
        plan_row = mapping(rows, "implementation-plan", "ja")
        plan = (source / plan_row["Template"]).read_text(encoding="utf-8")
        plan = re.sub(r"<!--(?!\s*blend-template:).*?-->", "", plan, flags=re.S)
        plan = re.sub(r"\{\{.*?\}\}", "Scoped synthetic basis; no approval/runtime claim", plan, flags=re.S)
        final_plan = saved / "actual saved plan.md"
        final_plan.write_text(plan, encoding="utf-8")
        raw_files = {"basis.md": b"\xef\xbb\xbfsource\r\nraw bytes\r\n",
                     "scope.txt": b"second explicitly selected input\n"}
        for filename, raw in raw_files.items():
            (saved / filename).write_bytes(raw)

        def gate_cli(helper: Path, arguments: list[str], expected: int = 0):
            result = subprocess.run([sys.executable, str(helper), *arguments], cwd=base,
                                    capture_output=True, text=True, encoding="utf-8",
                                    env={**os.environ, "PYTHONIOENCODING": "utf-8"})
            assert result.returncode == expected, f"Relocated gate returned {result.returncode}: {result.stdout}{result.stderr}"
            return result.stdout

        check_args = ["check", "--type", "implementation-plan", "--language", "ja", "--file", str(final_plan),
                      "--filename", "plans/score-save-implementation-plan.ja.md"]
        assert "PASS" in gate_cli(source_gate, check_args)
        capture_args = ["capture", "--root", str(saved), "--file", "basis.md", "--file", "scope.txt"]
        expected_receipts = [{"path": filename, "size": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}
                             for filename, raw in raw_files.items()]
        before_capture = {p.relative_to(saved).as_posix(): p.read_bytes() for p in saved.rglob("*") if p.is_file()}
        assert json.loads(gate_cli(source_gate, capture_args)) == {"kind": "observed-byte-capture", "files": expected_receipts}
        assert "Draft/Blocked" in gate_cli(source_gate, ["capture", "--root", str(saved), "--file", "../basis.md"], 1)
        for kind, names in controls["output_names"].items():
            for row in (row for row in rows if row["Output type"] == kind):
                for name in names:
                    validate_output_filename(name.format(language=row["Language"]), row)
        for kind, names in controls["invalid_output_names"].items():
            for name in names:
                rejected(lambda: validate_output_filename(name, mapping(rows, kind, "ja")), "output pattern")
        design = base / "synthetic design input"
        shutil.copytree(root / "tests/fixtures/test-spec/inputs/valid", design)
        design_before = {p.name: p.read_bytes() for p in design.iterdir() if p.is_file()}
        scenario_designs = {}
        for fixture, version in (("scenario-grouping", "2.5.0"), ("compact-scenario", "2.5.0")):
            relocated_design = base / fixture
            shutil.copytree(root / "tests/fixtures/test-spec/inputs" / fixture, relocated_design)
            scenario_designs[fixture] = (relocated_design, version)
        scenario_before = {p: p.read_bytes() for directory, _ in scenario_designs.values()
                           for p in directory.glob("*.md")}
        all_files = {p.relative_to(output).as_posix(): p.read_bytes() for p in output.rglob("*") if p.is_file()}
        for name in builder.SKILLS:
            skill_sets = []
            for profile, directory in builder.PROFILES.items():
                emitted = output / profile / directory / name
                resources = {p.relative_to(emitted).as_posix(): p.read_bytes() for p in emitted.rglob("*") if p.is_file()}
                skill_sets.append(resources)
                assert [p.name for p in emitted.rglob("SKILL.md")] == ["SKILL.md"]
                builder.frontmatter(resources["SKILL.md"], name)
                for common in ("workflow.md", "review-policy.md", "bug-hunter.md", "automation-testing.md"):
                    assert resources["_kit/shared/" + common] == builder.normalize((source / "shared" / common).read_bytes())
                assert not any(part.startswith("legacy") or part == "customer" for relative in resources for part in Path(relative).parts)
                license_bytes = (source / "shared/bug-hunter-LICENSE.txt").read_bytes()
                assert resources["_kit/shared/bug-hunter-LICENSE.txt"] == license_bytes
                assert resources["_kit/shared/scripts/artifact_gate.py"] == source_gate.read_bytes()
                assert len([p for p in resources if p.endswith("/bug-hunter.md")
                            and b"## Inline RoleDisposition" in resources[p]]) == 1
                if name == "blend-review-artifacts":
                    assert resources["references/bug-hunter-LICENSE.txt"] == license_bytes
                    assert b"compatibility pointer" in resources["references/bug-hunter.md"]
                if name == "blend-generate-test-spec":
                    for helper in builder.REPORT_RUNTIME_SCRIPTS:
                        assert resources["scripts/" + helper] == (source / "skills" / name / "scripts" / helper).read_bytes()
                    assert resources["scripts/run_artifacts.py"] == (source / builder.EVIDENCE_HELPER).read_bytes()
                    assert not any("customer_report" in relative or "customer-report.md" in relative
                                   for relative in resources)
                    active_assets = {Path(relative).name for relative in resources if relative.endswith(".xlsx")}
                    assert active_assets == {"test-report-block-template.ja.xlsx", "test-report-block-template.vi.xlsx"}
                    assert b"Pillow" in resources["requirements.txt"]
                    environment = {**os.environ, "PYTHONIOENCODING": "cp1252", "PYTHONUTF8": "0",
                                   "PYTHONDONTWRITEBYTECODE": "1"}
                    # Relocated report CLIs own UTF-8 transport even under a legacy Windows codepage.
                    for language in ("ja", "vi"):
                        report = base / f"{profile}-test-report.{language}.xlsx"
                        generate = [sys.executable, str(emitted / "scripts/export_report.py"),
                                    "--source-dir", str(design), "--output", str(report), "--language", language]
                        result = subprocess.run(generate, cwd=base, capture_output=True, text=True,
                                                encoding="utf-8", env=environment)
                        assert result.returncode == 0, result.stderr
                        receipt = json.loads(result.stdout)
                        assert receipt["family"] == "test-report"
                        before = report.read_bytes()
                        check = [sys.executable, str(emitted / "scripts/check_report.py"),
                                 "--report", str(report), "--source-dir", str(design), "--language", language]
                        result = subprocess.run([*check, "--phase", "in-progress"], cwd=base,
                                                capture_output=True, text=True, encoding="utf-8", env=environment)
                        assert result.returncode == 0, result.stderr
                        assert report.read_bytes() == before, "Read-only checker changed saved report"
                        incomplete = subprocess.run([*check, "--phase", "complete"], input="{}", cwd=base,
                                                    capture_output=True, text=True, encoding="utf-8", env=environment)
                        assert incomplete.returncode != 0, "Missing actual attestations accepted"
                        collision = subprocess.run(generate, cwd=base, capture_output=True, text=True,
                                                   encoding="utf-8", env=environment)
                        assert collision.returncode != 0, "Existing report overwritten"
                        assert report.read_bytes() == before
                        from openpyxl import load_workbook
                        final_book = load_workbook(report)
                        assert final_book.sheetnames[:2] == list(builder.report_layouts(source)[language]["sheets"][:2])
                        final_book.close()
                    # Exercise new compact and unchanged 1.2 designs from emitted
                    # runtime, away from the source checkout, on every profile/language.
                    for fixture, (scenario_design, report_version) in scenario_designs.items():
                        for language in ("ja", "vi"):
                            report = base / f"{profile}-{fixture}.{language}.xlsx"
                            result = subprocess.run([sys.executable, str(emitted / "scripts/export_report.py"),
                                "--source-dir", str(scenario_design), "--output", str(report), "--language", language],
                                cwd=base, capture_output=True, text=True, encoding="utf-8", env=environment)
                            assert result.returncode == 0, result.stderr
                            assert json.loads(result.stdout)["family"] == "test-report"
                            before = report.read_bytes()
                            result = subprocess.run([sys.executable, str(emitted / "scripts/check_report.py"),
                                "--report", str(report), "--source-dir", str(scenario_design),
                                "--language", language, "--phase", "in-progress"],
                                cwd=base, capture_output=True, text=True, encoding="utf-8", env=environment)
                            assert result.returncode == 0, result.stderr
                            book = load_workbook(report)
                            assert book.custom_doc_props["TemplateVersion"].value == report_version
                            assert len(book.sheetnames) == 2
                            book.close()
                            assert report.read_bytes() == before, "Scenario checker changed report bytes"
                if name == "blend-automation-test":
                    dependency = emitted / "_kit" / builder.REPORT_OWNER
                    for helper in builder.REPORT_RUNTIME_SCRIPTS:
                        assert resources[f"_kit/{builder.REPORT_OWNER}/scripts/{helper}"] == (source / builder.REPORT_OWNER / "scripts" / helper).read_bytes()
                    assert (dependency / "scripts/run_artifacts.py").read_bytes() == (source / builder.EVIDENCE_HELPER).read_bytes()
                    assert (dependency / "requirements.txt").read_bytes() == (source / builder.REPORT_OWNER / "requirements.txt").read_bytes()
                    # Execute the relocated writer import/CLI outside the source
                    # tree; help alone would not resolve the evidence validator.
                    result = subprocess.run([sys.executable, "-c",
                        "import sys; sys.path.insert(0, sys.argv[1]); import update_report; update_report._validator(); print('validator resolved')",
                        str(dependency / "scripts")], cwd=base, capture_output=True, text=True,
                        env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
                    assert result.returncode == 0, result.stderr
                    assert result.stdout.strip() == "validator resolved"
                    result = subprocess.run([sys.executable, str(dependency / "scripts/update_report.py"), "--help"],
                                            cwd=base, capture_output=True, text=True,
                                            env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
                    assert result.returncode == 0, result.stderr
                    saved_report = base / f"{profile}-test-report.vi.xlsx"
                    saved_bytes = saved_report.read_bytes()
                    result = subprocess.run([sys.executable, str(dependency / "scripts/update_report.py"),
                        "--source-dir", str(design), "--report", str(saved_report), "--language", "vi",
                        "--feature-id", "SYN-PACKAGE", "--bindings"], cwd=base,
                        capture_output=True, text=True, encoding="utf-8",
                        env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1", "PYTHONIOENCODING": "utf-8"})
                    assert result.returncode == 0, result.stderr
                    binding = json.loads(result.stdout)
                    assert binding["feature_id"] == "SYN-PACKAGE" and binding["schema_version"] == "2.5.0"
                    assert binding["variants"] and saved_report.read_bytes() == saved_bytes
                    for fixture, (scenario_design, report_version) in scenario_designs.items():
                        for language in ("ja", "vi"):
                            saved_report = base / f"{profile}-{fixture}.{language}.xlsx"
                            arguments = [sys.executable, str(dependency / "scripts/update_report.py"),
                                "--source-dir", str(scenario_design), "--report", str(saved_report), "--language", language]
                            saved_bytes = saved_report.read_bytes()
                            result = subprocess.run([*arguments, "--bindings"], cwd=base,
                                capture_output=True, text=True, encoding="utf-8", env=environment)
                            assert result.returncode == 0, result.stderr
                            binding = json.loads(result.stdout)
                            assert binding["schema_version"] == report_version and binding["feature_id"] == "SYN-SC"
                            assert len(binding["variants"]) == 9 and saved_report.read_bytes() == saved_bytes
                            fields = binding["variants"]["TC-SC-01 / lt"]
                            actual = "\n".join(f"{checkpoint}: Synthetic failure observation; no live application proof."
                                               for checkpoint in fields["checkpoints"])
                            payload = {"identity": {"design_revision": binding["design_revision"],
                                "feature_id": "SYN-SC", "case_id": "TC-SC-01", "variant_id": "lt", "run_id": "2026-10-05-001"},
                                "status": "FAIL", "actual": actual}
                            result = subprocess.run(arguments, input=json.dumps(payload), cwd=base,
                                capture_output=True, text=True, encoding="utf-8", env=environment)
                            assert result.returncode == 0, result.stderr
                            book = load_workbook(saved_report)
                            sheet = book[binding["sheet"]]
                            assert sheet[fields["actual_cell"]].value == actual
                            assert sheet[fields["status_cell"]].value == ("Không đạt" if language == "vi" else "不合格")
                            for identity, other in binding["variants"].items():
                                if identity != "TC-SC-01 / lt":
                                    assert sheet[other["actual_cell"]].value is None
                                    assert sheet[other["status_cell"]].value == ("Chưa thực hiện" if language == "vi" else "未実行")
                            if report_version == "2.5.0":
                                assert fields["actual_cell"].startswith("D") and fields["status_cell"].startswith("E")
                            book.close()
                runtime_registry = builder.read_registry(resources["_kit/shared/artifact-formats.md"].decode())
                for source_row, runtime_row in zip(rows, runtime_registry):
                    assert runtime_row["Template"] == "_kit/" + source_row["Template"]
                    assert resources[runtime_row["Template"]] == (source / source_row["Template"]).read_bytes()
                for relative in resources:
                    if relative.endswith(".md"):
                        builder.validate_links(resources, relative)
                    assert not any(part in builder.FORBIDDEN_PARTS or part.startswith("~$") for part in Path(relative).parts)
                gate = emitted / "_kit/shared/scripts/artifact_gate.py"
                assert "PASS" in gate_cli(gate, check_args)
                assert json.loads(gate_cli(gate, capture_args))["files"] == expected_receipts
            assert skill_sets[0] == skill_sets[1] == skill_sets[2]
        assert before_capture == {p.relative_to(saved).as_posix(): p.read_bytes() for p in saved.rglob("*") if p.is_file()}
        assert owner_file.read_bytes() == owner_bytes
        assert design_before == {p.name: p.read_bytes() for p in design.iterdir() if p.is_file()}
        assert scenario_before == {p: p.read_bytes() for p in scenario_before}, "Frozen scenario inputs changed"
        assert len(list(base.glob("*-test-report.*.xlsx"))) == 6, "Checking created a second report"
        # Inline-only drift, extended labels and a missing template fail at the shipped entrypoint.
        final_plan.write_text(plan.replace("情報源リビジョン: Scoped synthetic basis; no approval/runtime claim",
                                           "情報源リビジョン:\n\nDetails cannot replace an inline summary"), encoding="utf-8")
        assert "Draft/Blocked" in gate_cli(gate, check_args, 1)
        final_plan.write_text(plan.replace("情報源リビジョン:", "情報源リビジョン (selected):"), encoding="utf-8")
        assert "Draft/Blocked" in gate_cli(gate, check_args, 1)
        final_plan.write_text(plan, encoding="utf-8")
        bundled_asset = emitted / "_kit" / plan_row["Template"]
        asset_bytes = bundled_asset.read_bytes()
        bundled_asset.unlink()
        assert "Draft/Blocked" in gate_cli(gate, check_args, 1)
        bundled_asset.write_bytes(asset_bytes)
        with contextlib.redirect_stdout(io.StringIO()):
            builder.build(source, output, controls["profiles"], True)
        assert all_files == {p.relative_to(output).as_posix(): p.read_bytes() for p in output.rglob("*") if p.is_file()}

        sentinel = output / "unrelated.txt"
        sentinel.write_text("preserve me")
        rejected(lambda: builder.build(source, output, controls["profiles"], True), "unrelated existing file")
        assert sentinel.read_text() == "preserve me"
        # Test-owned temporary file; remove only this exact file, never a user directory.
        sentinel.unlink()
        conflict = next(p for p in output.rglob("SKILL.md"))
        conflict.write_text("user edited skill")
        rejected(lambda: builder.build(source, output, controls["profiles"], True), "existing changed resource")
        assert conflict.read_text() == "user edited skill"
        rejected(lambda: builder.build(source, source, ["codex"]), "output equals source")
        rejected(lambda: builder.build(source, source / "skills/generated", ["codex"]), "output inside skill source")
        rejected(lambda: builder.build(source, base / "duplicate", ["codex", "codex"]), "duplicate profile")
        rejected(lambda: builder.safe_path(source, controls["invalid_registry_path"]), "traversal")

        registry = source / "shared/artifact-formats.md"
        registry_bytes = registry.read_bytes()
        registry_text = registry_bytes.decode()
        first = rows[0]["Template"]
        registry.write_text(registry_text.replace("| " + first + " |", "| " + controls["invalid_registry_path"] + " |"), encoding="utf-8")
        rejected(lambda: builder.build(source, base / "bad registry", ["codex"]), "registry traversal")
        registry.write_bytes(registry_bytes)
        asset = source / first
        asset_bytes = asset.read_bytes()
        asset.write_bytes(asset_bytes.replace(("@" + rows[0]["Version"]).encode(), b"@9.0.0"))
        rejected(lambda: builder.build(source, base / "bad version", ["codex"]), "template version")
        asset.write_bytes(asset_bytes)
        registry.write_text(registry_text.replace("| research |", "| unsupported |", 1), encoding="utf-8")
        rejected(lambda: builder.build(source, base / "bad bilingual pair", ["codex"]), "bilingual mapping")
        registry.write_bytes(registry_bytes)
        skill = source / "skills/blend-generate-task/SKILL.md"
        skill_bytes = skill.read_bytes()
        skill.write_bytes(skill_bytes.replace(b"name: blend-generate-task", b"name: wrong-name", 1))
        rejected(lambda: builder.build(source, base / "bad name", ["codex"]), "SKILL directory/name mismatch")
        skill.write_bytes(skill_bytes)
        skill.write_bytes(skill_bytes + b"\n" + controls["private_path"].encode())
        rejected(lambda: builder.build(source, base / "private path", ["codex"]), "private machine path")
        skill.write_bytes(skill_bytes)
        for relative in ("skills/blend-generate-test-spec/scripts/render_report.py",
                         "skills/blend-generate-test-spec/scripts/block_report.py",
                         "skills/blend-generate-test-spec/scripts/report_model.py",
                         "skills/blend-generate-test-spec/scripts/check_report.py",
                         "skills/blend-generate-test-spec/scripts/update_report.py",
                         "skills/blend-automation-test/scripts/run_artifacts.py",
                         "skills/blend-generate-test-spec/assets/test-report-block-template.ja.xlsx",
                         "skills/blend-generate-test-spec/requirements.txt"):
            missing_resource = source / relative
            original = missing_resource.read_bytes()
            missing_resource.unlink()
            rejected(lambda: builder.build(source, base / "missing report dependency", ["codex"]),
                     "missing report dependency: " + relative)
            missing_resource.write_bytes(original)
        missing = source / "shared/bug-hunter-LICENSE.txt"
        missing.unlink()
        rejected(lambda: builder.build(source, base / "missing license", ["claude"]), "broken dependency closure")

    assert current_resources == {relative: (root / relative).read_bytes() for relative in current_resources}, "Package checks changed current assets"
    return ["one active report family/JA-VI assets; actual relocated UTF-8 generation and same-file read-only checks across all three profiles under cp1252; overwrite and missing completion attestations rejected",
            "canonical current workbook resources preserve bytes; Office owner files excluded and untouched",
            "six skills on three profiles build from relocated/spaced roots at another cwd without a test harness or a parent checkout; frozen five-skill controls retained",
            "all emitted resources/local links and rewritten registry paths resolve without source fallback; source1.0/1.1/1.2 legacy and current1.3 dispatch reject duplicate/fenced identity spoofing",
            f"all {len(rows)} registered template bytes, SKILL descriptions/bodies and resource inventories match across profiles",
            "source and all 18 shipped gates check JA Markdown and capture explicitly selected raw bytes from another cwd; automation writer resolves bundled evidence validator without sibling installation",
            "JA/VI source1.2/1.3 both generate only report2.5 export/check/bind/write through relocated canonical runtimes on all profiles; synthetic FAIL preserves eight untouched NOT RUN variants and frozen sources",
            "deployed inline/presentation/missing-asset drift and traversal captures remain Draft/Blocked",
            "topic/plan/code-review filename shapes accept multiple descriptive and exact-ID-prefixed names, reject literal/traversal/legacy output names",
            "one canonical behavioral protocol/common policy and exact license bytes bundled; legacy artifact pointers resolve",
            "private snapshots/configs/tests excluded; collision/rebuild preserves existing output",
            "current JA/VI report2.5 assets remain byte-identical; no legacy/customer copies are bundled",
            "invalid mapping/version/bilingual pair/name/private path/missing license rejected",
            "package checks do not prove installation, runtime discovery or actual agent semantic parity"]
